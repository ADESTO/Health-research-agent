from __future__ import annotations

import re

import duckdb
import pytest

from research_agent.config import settings
from research_agent.ingestion.fulltext import assess, clean_latex, select_for_reading
from research_agent.ingestion.health_filter import health_term_regex
from tests.fake_llm import FakeLLM


# ---------------------------------------------------------------- unit: health filter
@pytest.mark.parametrize("text,expected", [
    ("Deep learning on electronic health records", "electronic health records"),
    ("Sepsis prediction in the ICU", "sepsis"),
    ("COVID-19 spread modelling", "covid-19"),
    ("Hospitality industry demand forecasting", ""),
    ("Diseased trees in forests", ""),
])
def test_health_regex(text, expected):
    got = duckdb.execute("SELECT lower(regexp_extract(?, ?, 1))", [text, health_term_regex()]).fetchone()[0]
    assert got == expected


# ---------------------------------------------------------------- unit: LaTeX cleaning
def test_clean_latex_sections_and_comments():
    tex = (r"\documentclass{article}\begin{document}\section{Introduction} Hello world, this is the introduction % secret" "\n"
           r"\section{Methods} We used a \textbf{CNN}. \begin{equation} a=b \end{equation}"
           r"\begin{figure}\caption{Pipeline of the CNN}\end{figure}\section{Limitations} Small sample."
           r"\end{document}")
    clean, sections = clean_latex(tex)
    heads = [s["heading"] for s in sections]
    assert heads[:3] == ["Introduction", "Methods", "Limitations"]
    assert "secret" not in clean and "a=b" not in clean
    assert "CNN" in clean and "Pipeline of the CNN" in clean
    picked = select_for_reading(sections, budget=10_000)
    assert picked.index("Methods") < picked.index("Limitations")  # keeps paper order


def test_assess_flags_stubs_and_templates():
    assert assess("%auto-ignore", "", "Anything") == "stub"
    long_raw = "x" * 5000
    assert assess(long_raw, "## Intro\nInstructions for authors of AIP journals " * 50,
                  "Malaria forecasting in Kenya") == "template_suspect"


# ---------------------------------------------------------------- ingestion
def test_ingest_selects_health_subset(loaded_db):
    from research_agent.db import get_conn

    assert loaded_db["inserted"] > 50
    with get_conn() as pg:
        titles = [r["title"] for r in pg.execute("SELECT title FROM papers").fetchall()]
        reasons = {r["health_reason"].split(":")[0] for r in pg.execute("SELECT health_reason FROM papers")}
        n_stats = pg.execute("SELECT sum(n_papers) n FROM corpus_year_stats").fetchone()["n"]
        n_papers = pg.execute("SELECT count(*) n FROM papers").fetchone()["n"]
        n_emb = pg.execute("SELECT count(*) n FROM papers WHERE embedding IS NOT NULL").fetchone()["n"]
    assert not any(t.startswith(("Robust quadruped", "Next-to-leading", "Graph neural networks for traffic",
                                 "Efficient transformers")) for t in titles)
    assert reasons <= {"keyword", "category"}
    assert n_stats > n_papers  # denominators include non-health papers
    assert n_emb == n_papers


def test_ingest_is_resumable(loaded_db):
    from research_agent.db import connect
    from research_agent.ingestion.load import load_papers

    pg = connect()
    try:
        assert load_papers(pg, log=lambda *_: None) == 0
    finally:
        pg.close()


# ---------------------------------------------------------------- tools
@pytest.fixture
def ctx(loaded_db):
    from research_agent.runstate import RunContext

    c = RunContext.create("test question", llm_factory=lambda strong=False: FakeLLM())
    yield c
    c.close()


def test_hybrid_search_ranks_topic_first(ctx):
    from research_agent.tools.search import hybrid_search

    res = hybrid_search(ctx, "forecasting malaria incidence with rainfall", keywords="malaria", limit=10)
    assert res["n"] == 10
    assert all("malaria" in r["title"].lower() for r in res["results"][:5])
    res2 = hybrid_search(ctx, "sepsis", year_from=2024, year_to=2025)
    assert all(2024 <= r["year"] <= 2025 for r in res2["results"])


def test_shortlist_rejects_unknown_ids_and_caps(ctx):
    from research_agent.tools.search import add_to_shortlist, hybrid_search

    ids = [r["paper_id"] for r in hybrid_search(ctx, "malaria", limit=5)["results"]]
    out = add_to_shortlist(ctx, ids + ["nope.123"], reason="t")
    assert out["added"] == 5 and out["unknown_ids"] == ["nope.123"]


def test_trend_is_normalised_and_flags_partial_year(ctx):
    from research_agent.tools.trends import topic_trend

    t = topic_trend(ctx, '"foundation model"')
    last = t["series"][-1]
    assert last["year"] == 2026 and last["partial_year"] is True
    assert all(s["per_10k_arxiv"] is not None for s in t["series"])
    assert t["summary"]["ratio_late_over_early"] in ("inf",) or t["summary"]["ratio_late_over_early"] > 1


def test_claim_validation_rejects_untestable(ctx):
    from research_agent.tools.claims import propose_claim

    assert "error" in propose_claim(ctx, "Most use CNNs", "prevalence", {"field": "methods", "any_of": ["CNN"]})
    assert "error" in propose_claim(ctx, "x", "prevalence", {"field": "authors", "any_of": ["a"], "min_share": .1})


# ---------------------------------------------------------------- full multi-agent runs
@pytest.mark.parametrize("mode", ["orchestrated", "pipeline"])
def test_end_to_end_run(loaded_db, mode):
    from research_agent.agents.orchestrator import run_research
    from research_agent.db import get_conn

    events = []
    res = run_research("What ML methods are used for malaria forecasting and what is missing?",
                       mode=mode, llm_factory=lambda strong=False: FakeLLM(),
                       on_event=lambda a, k, p: events.append((a, k)))
    report = res["report"]
    assert report and "## Evidence table" in report and "## References" in report
    # hallucinated citation removed, unverified claim flagged
    assert "9999.99999" not in report and "citation removed" in report
    assert "not verified]" in report

    with get_conn() as pg:
        run = pg.execute("SELECT * FROM runs WHERE run_id=%s", (res["run_id"],)).fetchone()
        claims = {c["text"]: c for c in pg.execute("SELECT * FROM claims WHERE run_id=%s", (res["run_id"],))}
        n_ft = pg.execute("""SELECT count(*) n FROM run_papers rp JOIN extractions e USING (paper_id)
                             WHERE rp.run_id=%s AND e.source='fulltext'""", (res["run_id"],)).fetchone()["n"]
        stub = pg.execute("SELECT count(*) n FROM paper_fulltext WHERE status='stub'").fetchone()["n"]
        n_q = pg.execute("SELECT count(*) n FROM protocol_extractions WHERE run_id=%s", (res["run_id"],)).fetchone()["n"]
        n_ex = pg.execute("""SELECT count(*) n FROM run_papers rp JOIN extractions e USING (paper_id)
                             WHERE rp.run_id=%s""", (res["run_id"],)).fetchone()["n"]
    assert run["status"] == "done" and run["llm_calls"] > 10
    # ask runs get question-specific fields, extracted for every analysed paper after the literature step
    assert n_q == n_ex > 0 and ("protocol", "finish") in events
    assert "Numbers checked after writing" in report

    # Deterministic evidence verdicts
    assert claims["Most studies use data from East Africa."]["status"] == "supported"
    assert claims["No study uses MIMIC data."]["status"] == "supported"
    assert claims["Foundation-model papers have increased sharply."]["status"] == "supported"
    assert claims["Malaria research is declining."]["status"] == "unsupported"
    ext = claims["Few malaria forecasting studies use external validation."]
    assert ext["status"] == "supported" and ext["result"]["n_matching"] == 0

    # the extraction step's own allowance, plus whatever the pre-report resolution pass spent on
    # claims whose state rested on abstracts (tools/resolve.py); both are capped
    assert 0 < n_ft <= settings.max_fulltext + settings.resolve_max_papers
    assert stub <= 1
    agents_seen = {a for a, _ in events}
    assert {"discovery", "literature", "methods", "trends", "gaps", "evidence", "synthesis"} <= agents_seen
    if mode == "orchestrated":
        assert "orchestrator" in agents_seen
    # every paper in the reference list is from the shortlist
    assert len(re.findall(r"^- \*\*arXiv:", report, flags=re.M)) == \
        len(set(re.findall(r"^- \*\*arXiv:([^*]+)\*\*", report, flags=re.M)))


# ---------------------------------------------------------------- provider seam
def test_groq_message_translation_roundtrip():
    from research_agent.llm.groq_client import _to_openai_messages

    msgs = [
        {"role": "user", "content": [{"type": "text", "text": "hi"}]},
        {"role": "assistant", "content": [{"type": "text", "text": "searching"},
                                          {"type": "tool_use", "id": "t1", "name": "hybrid_search",
                                           "input": {"query": "malaria"}}]},
        {"role": "user", "content": [{"type": "tool_result", "tool_use_id": "t1", "content": "{}",
                                      "is_error": False}]},
    ]
    out = _to_openai_messages("sys", msgs)
    assert [m["role"] for m in out] == ["system", "user", "assistant", "tool"]
    assert out[2]["tool_calls"][0]["function"]["name"] == "hybrid_search"
    assert out[3]["tool_call_id"] == "t1"


def test_api_health_and_missing_run(loaded_db):
    from fastapi.testclient import TestClient

    from research_agent.api.main import app

    client = TestClient(app)
    assert client.get("/health").json()["papers"] == loaded_db["total_papers"]
    assert client.get("/runs/00000000-0000-0000-0000-000000000000").status_code == 404
    assert client.post("/runs", json={"question": "short"}).status_code == 422


def test_api_builds_a_map(loaded_db, monkeypatch):
    import time

    from fastapi.testclient import TestClient

    from research_agent.api import main as api
    from research_agent.opportunity import pipeline

    from research_agent import jobs
    monkeypatch.setattr(jobs, "_llm_factory", lambda strong=False: FakeLLM())
    api._hits.clear()
    client = TestClient(api.app)
    assert client.post("/runs", json={"question": "Can ML improve malaria forecasts?", "mode": "maps"}).status_code == 422
    res = client.post("/runs", json={"question": "Can machine learning improve 1-6 month malaria forecasting?",
                                     "mode": "map"}).json()
    assert res["mode"] == "map" and res["job_id"]
    assert jobs.Worker(name="test").run_once()          # a worker picks the job up
    for _ in range(120):
        run = client.get(f"/runs/{res['run_id']}").json()["run"]
        if run["status"] in ("done", "failed"):
            break
        time.sleep(0.5)
    assert run["status"] == "done" and run["is_map"] and run["has_report"]
    assert client.get(f"/runs/{res['run_id']}/report").json()["report_markdown"].startswith("# Research Opportunity Map")
    listed = {r["run_id"]: r["is_map"] for r in client.get("/runs").json()["runs"]}
    assert listed[res["run_id"]] is True
    # a map counts as two runs against the hourly limit: 2 used, 3 left, so one more map and one run fit
    assert api._hits and len(next(iter(api._hits.values()))) == 2
    api._hits.clear()


def test_anthropic_adapter_parses_sdk_response(monkeypatch):
    import anthropic
    from anthropic.types import Message

    from research_agent.llm.anthropic_client import AnthropicClient

    msg = Message.model_validate({
        "id": "m1", "type": "message", "role": "assistant", "model": "claude-haiku-4-5-20251001",
        "stop_reason": "tool_use", "stop_sequence": None,
        "usage": {"input_tokens": 12, "output_tokens": 5},
        "content": [{"type": "text", "text": "Let me search."},
                    {"type": "tool_use", "id": "tu1", "name": "hybrid_search", "input": {"query": "malaria"}}],
    })
    captured = {}

    def fake_create(**kw):
        captured.update(kw)
        return msg

    client = AnthropicClient("claude-haiku-4-5-20251001", api_key="test")
    monkeypatch.setattr(client._client.messages, "create", fake_create)
    resp = client.chat("sys", [{"role": "user", "content": [{"type": "text", "text": "q"}]}],
                       tools=[{"name": "hybrid_search", "description": "d",
                               "input_schema": {"type": "object", "properties": {}}}],
                       force_tool="hybrid_search")
    assert captured["tool_choice"] == {"type": "tool", "name": "hybrid_search"}
    assert resp.tool_calls[0].input == {"query": "malaria"} and resp.text == "Let me search."
    assert client.usage.input_tokens == 12
    # the stable prefix is always cached: the same system prompt and tool schema go out on every call
    assert captured["system"][0]["cache_control"] == {"type": "ephemeral"}
    assert captured["tools"][-1]["cache_control"] == {"type": "ephemeral"}
    # but a one-shot call (one paper, read once) must NOT pay a cache-write premium on its own text:
    # nothing will ever read that entry back
    assert "cache_control" not in captured["messages"][-1]["content"][-1]
    if not client._accepts_temperature:
        assert "temperature" not in captured

    # a loop in progress is different: the next turn resends everything up to here, so mark the end
    loop = [{"role": "user", "content": [{"type": "text", "text": "q"}]},
            {"role": "assistant", "content": [{"type": "text", "text": "a"}]},
            {"role": "user", "content": [{"type": "tool_result", "tool_use_id": "t", "content": "r"}]}]
    client.chat("sys", loop, tools=[{"name": "t", "description": "d",
                                     "input_schema": {"type": "object", "properties": {}}}])
    assert captured["messages"][-1]["content"][-1]["cache_control"] == {"type": "ephemeral"}
    assert "cache_control" not in loop[-1]["content"][-1], "the caller's own messages must not be mutated"

    # and nothing is marked at all when caching is off
    plain = AnthropicClient("claude-haiku-4-5-20251001", api_key="test", prompt_cache=False)
    monkeypatch.setattr(plain._client.messages, "create", fake_create)
    plain.chat("sys", loop, tools=[{"name": "t", "description": "d",
                                    "input_schema": {"type": "object", "properties": {}}}])
    assert captured["system"] == "sys" and "cache_control" not in captured["tools"][-1]
    _ = anthropic


def test_health_filter_rejects_passing_mentions(tmp_path):
    import datetime as dt

    import pyarrow as pa
    import pyarrow.parquet as pq

    from research_agent.ingestion.health_filter import health_subset_sql
    from tests.fixtures.make_fixtures import META_SCHEMA, _row

    rng = __import__("random").Random(1)
    filler = " We evaluate on several benchmarks and report consistent improvements over strong baselines."
    cases = {
        # (title, abstract, categories): expected
        ("Fair Influence Maximization: A Welfare Optimization Approach",
         "We study fairness in influence maximization with applications to public health campaigns.", "cs.SI"): False,
        ("Contrasting Effects of Strong Ties on SIR and SIS Processes in Temporal Networks",
         "We study epidemic spreading on temporal networks.", "physics.soc-ph"): False,
        ("Spreading processes on temporal contact networks",
         "We model epidemic spreading and infection risk for infectious disease control.", "physics.soc-ph"): True,
        ("Early sepsis prediction with gradient boosting",
         "We build a model on routinely collected variables.", "cs.LG"): True,
        ("Single-molecule tracking with a new estimator",
         "We propose an estimator for diffusion coefficients in living cells.", "q-bio.QM"): False,
        ("Dose calculation for proton therapy",
         "A fast Monte Carlo dose engine.", "physics.med-ph"): True,
    }
    rows = [_row(f"2001.{i:05d}", t, a + filler, c, 2020, rng) for i, (t, a, c) in enumerate(cases)]
    path = tmp_path / "m.parquet"
    pq.write_table(pa.Table.from_pylist(rows, schema=META_SCHEMA), path)
    got = {r[1] for r in duckdb.execute(health_subset_sql(str(path), 2010)).fetchall()}
    for (title, _, _), expected in cases.items():
        assert (title in got) == expected, title
    _ = dt


def test_run_with_nothing_found_ends_gracefully(loaded_db, monkeypatch):
    from research_agent.agents.orchestrator import run_research
    from tests import fake_llm

    def discovery_finds_nothing(self, step, last, messages):
        if step == 0:
            return [fake_llm._call("hybrid_search", query="quantum chromodynamics", year_from=2030)]
        return [fake_llm._call("finish", shortlist_size=0, coverage_notes="Nothing on this topic.")]

    monkeypatch.setattr(fake_llm.FakeLLM, "_discovery", discovery_finds_nothing)
    for mode in ("pipeline", "orchestrated"):
        res = run_research("Unrelated physics question?", mode=mode, llm_factory=lambda strong=False: FakeLLM())
        assert res["report"].startswith("# No relevant papers found")
        assert "Nothing on this topic." in res["report"]


def test_compaction_shortens_old_results_only():
    from research_agent.agents.base import KEEP_RECENT_MESSAGES, compact

    big = "x" * 5000
    msgs = [{"role": "user", "content": [{"type": "text", "text": "task"}]}]
    for i in range(6):
        msgs.append({"role": "assistant", "content": [{"type": "tool_use", "id": f"t{i}", "name": "s", "input": {}}]})
        msgs.append({"role": "user", "content": [{"type": "tool_result", "tool_use_id": f"t{i}", "content": big}]})
    out = compact(msgs, budget=12000)
    assert out[0] == msgs[0] and len(out) == len(msgs)
    assert len(out[2]["content"][0]["content"]) < 500                    # oldest shortened
    assert out[-1]["content"][0]["content"] == big                        # latest untouched
    assert all(m["content"][0].get("content") == big for m in out[-KEEP_RECENT_MESSAGES:] if m["role"] == "user")
    assert msgs[2]["content"][0]["content"] == big                        # input not mutated


def test_retry_waits_honour_retry_after(monkeypatch):
    from research_agent.llm import base

    class RateLimited(Exception):
        status_code = 429
        response = type("R", (), {"headers": {"retry-after": "12"}})()

    waits, calls = [], {"n": 0}
    monkeypatch.setattr(base.time, "sleep", lambda s: waits.append(s))

    def flaky():
        calls["n"] += 1
        if calls["n"] < 3:
            raise RateLimited()
        return "ok"

    assert base.with_retries(flaky) == "ok" and waits == [12.0, 12.0]


def test_agent_recovers_from_rejected_tool_call_and_forces_finish_with_finish_only(loaded_db):
    from research_agent.agents.base import Agent
    from research_agent.llm.base import LLMResponse, ToolCall, Usage
    from research_agent.runstate import RunContext
    from research_agent.tools.search import SEARCH_TOOLS

    class Stubborn:
        """Always searches; once rejects with a Groq-style error; checks what the forced finish offers."""
        provider, model = "stub", "stub"

        def __init__(self):
            self.usage, self.n, self.forced_tools = Usage(), 0, None

        def chat(self, system, messages, tools=None, force_tool=None, max_tokens=0, temperature=0):
            self.n += 1
            if force_tool == "finish":
                self.forced_tools = [t["name"] for t in tools]
                c = ToolCall("f", "finish", {"summary": "done"})
                return LLMResponse("", [c], [{"type": "tool_use", "id": "f", "name": "finish", "input": c.input}])
            if self.n == 2:
                raise RuntimeError("Error code: 400 - tool_use_failed: tool call validation failed")
            c = ToolCall(f"s{self.n}", "hybrid_search", {"query": "malaria"})  # identical every time
            return LLMResponse("", [c], [{"type": "tool_use", "id": c.id, "name": c.name, "input": c.input}])

    stub = Stubborn()
    ctx = RunContext.create("q", llm_factory=lambda strong=False: stub)
    events = []
    ctx.on_event = lambda a, k, p: events.append((k, p))
    agent = Agent("probe", "r", "You are a probe.", SEARCH_TOOLS, {"type": "object", "properties": {}}, max_turns=5)
    out = agent.run(ctx, "t")
    ctx.close()
    assert out["summary"] == "done" and stub.forced_tools == ["finish"]
    assert any(k == "error" and "invalid tool call" in p["error"] for k, p in events)
    previews = [p["preview"] for k, p in events if k == "tool_result"]
    assert any("already made this exact call" in pv for pv in previews)


def test_resume_skips_finished_agents_after_daily_limit(loaded_db):
    from research_agent.agents.orchestrator import run_research
    from research_agent.llm.base import DailyLimitError

    class QuotaHit(FakeLLM):
        def chat(self, system, messages, tools=None, force_tool=None, **kw):
            if "Literature Analyst" in system:
                raise DailyLimitError("The provider's DAILY token limit is used up")
            return super().chat(system, messages, tools=tools, force_tool=force_tool, **kw)

    with pytest.raises(RuntimeError) as err:
        run_research("What ML methods are used for malaria forecasting?", mode="pipeline",
                     llm_factory=lambda strong=False: QuotaHit())
    run_id = err.value.run_id

    events = []
    res = run_research(resume=run_id, mode="pipeline", llm_factory=lambda strong=False: FakeLLM(),
                       on_event=lambda a, k, p: events.append((a, k)))
    assert res["run_id"] == run_id and "## References" in res["report"]
    assert ("discovery", "skipped") in events and ("literature", "start") in events
    assert ("discovery", "start") not in events


def test_agent_ignores_invented_tool_arguments(ctx):
    from research_agent.agents.base import Agent
    from research_agent.tools.extraction import EXTRACTION_TOOLS

    agent = Agent("probe", "r", "s", EXTRACTION_TOOLS, {"type": "object", "properties": {}})
    tool = next(t for t in EXTRACTION_TOOLS if t.name == "extraction_coverage")
    out = agent._call_tool(ctx, tool, {"": ""})
    assert out["_ignored_arguments"] == [""] and "extracted" in out


def test_coerce_to_schema_only_fixes_mismatches():
    from research_agent.agents.base import coerce_to_schema

    schema = {"properties": {"observations": {"type": "array"}, "n": {"type": "integer"},
                             "summary": {"type": "string"}, "p": {"type": "object"}}}
    good = {"observations": ["a", "b"], "n": 3, "summary": "s", "p": {"x": 1}}
    assert coerce_to_schema(good, schema) == good  # well-formed calls are untouched
    fixed = coerce_to_schema({"observations": "one finding", "n": "4", "p": '{"x": 1}'}, schema)
    assert fixed == {"observations": ["one finding"], "n": 4, "p": {"x": 1}}


def test_groq_rejected_call_is_recovered_only_when_valid():
    from research_agent.llm.groq_client import _recover_rejected_call

    tools = [{"name": "finish", "description": "", "input_schema": {"type": "object"}}]

    class Err(Exception):
        def __init__(self, gen, code="tool_use_failed"):
            self.body = {"error": {"code": code, "failed_generation": gen}}

    ok = _recover_rejected_call(Err('{"name": "finish", "arguments": {"observations": "x"}}'), tools)
    assert ok.tool_calls[0].name == "finish" and ok.tool_calls[0].input == {"observations": "x"}
    assert _recover_rejected_call(Err('{"name": "finish", "arguments": {"obs'), tools) is None  # truncated
    assert _recover_rejected_call(Err('{"name": "nope", "arguments": {}}'), tools) is None     # unknown tool
    assert _recover_rejected_call(Err("{}", code="other"), tools) is None
    # open-weight models also mislabel the call itself, most often as "json", with the real payload inside.
    # The name carries no information when only one tool was asked for, so the payload is used.
    mis = Err('{"name": "json", "arguments": {"observations": "x"}}')
    assert _recover_rejected_call(mis, tools, "finish").tool_calls[0].name == "finish"
    assert _recover_rejected_call(mis, tools).tool_calls[0].name == "finish"          # only one tool offered
    bare = Err('{"name": "json", "observations": "x"}')                               # payload as the object
    assert _recover_rejected_call(bare, tools, "finish").tool_calls[0].input == {"observations": "x"}
    two = tools + [{"name": "other", "description": "", "input_schema": {"type": "object"}}]
    assert _recover_rejected_call(mis, two) is None             # a real choice of tool: never guess which
    empty = Err('{"name": "json", "arguments": {}}')
    assert _recover_rejected_call(empty, tools, "finish") is None    # nothing to act on: ask again instead


def test_groq_request_over_the_minute_limit_is_resent_with_a_smaller_reply_allowance():
    """Groq counts prompt + max_tokens against tokens-per-minute and rejects the whole request when the sum is over.
    The client shrinks max_tokens by exactly the excess and sends it again, without waiting out a minute."""
    from types import SimpleNamespace

    from research_agent.llm import groq_client as GC

    class TooLarge(Exception):
        status_code = 413

    seen = []

    def create(**kw):
        seen.append(kw["max_tokens"])
        prompt = 5900
        if prompt + kw["max_tokens"] > 8000:
            raise TooLarge("Error code: 413 - Request too large for model `openai/gpt-oss-120b` ... on tokens per "
                           f"minute (TPM): Limit 8000, Requested {prompt + kw['max_tokens']}, please reduce")
        msg = SimpleNamespace(content="ok", tool_calls=[])
        return SimpleNamespace(choices=[SimpleNamespace(message=msg, finish_reason="stop")],
                               usage=SimpleNamespace(prompt_tokens=prompt, completion_tokens=3))

    c = GC.GroqClient.__new__(GC.GroqClient)
    c._client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    c.model, c.usage = "openai/gpt-oss-120b", GC.Usage()
    assert c.chat("sys", [{"role": "user", "content": "hi"}], max_tokens=4096).text == "ok"
    assert seen[0] == 4096 and 5900 + seen[1] <= 8000 and len(seen) == 2
    # a prompt that alone nearly fills the minute: a clear message, not a silent cut to a useless reply
    seen.clear()

    def huge(**kw):
        seen.append(kw["max_tokens"])
        raise TooLarge(f"Request too large ... tokens per minute (TPM): Limit 8000, Requested {7800 + kw['max_tokens']}")
    c._client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=huge)))
    try:
        c.chat("sys", [{"role": "user", "content": "hi"}], max_tokens=4096)
        assert False, "should have raised"
    except RuntimeError as e:
        assert "PROMPT alone" in str(e) and len(seen) == 1


def test_report_normalises_brackets_and_flags_unbacked_numbers():
    from research_agent.agents.report import audit_numbers, normalise_citations

    body = normalise_citations("U-Net dominates【C1】 and diffusion is rare【C3】 ［arXiv:2401.00722］")
    assert body == "U-Net dominates[C1] and diffusion is rare[C3] [arXiv:2401.00722]"
    text, n = audit_numbers("12 of 16 papers use U-Net; 12 of 16 were read in full; 3/16 validate.",
                            {(12, 16)})
    assert n == 1 and "3/16 validate" not in text and "3/16 [unverified]" in text
    assert text.count("[unverified]") == 1


def test_claim_direction_and_strictness(ctx):
    from research_agent.tools.claims import direction_problem, propose_claim

    # the three mis-specified claims from the first real run
    assert direction_problem("Diffusion methods appear in less than 15% of the papers",
                             {"min_share": 0.15})
    assert direction_problem("External validation is reported in less than 20% of the papers",
                             {"min_share": 0.2})
    assert direction_problem("U-Net variants dominate, used in most papers", {"max_share": 0.5})
    assert direction_problem("U-Net variants are used in at least 50% of papers", {"min_share": 0.5}) is None

    bad = propose_claim(ctx, "Few papers use MIMIC.", "prevalence",
                        {"field": "datasets", "any_of": ["MIMIC"], "min_count": 1})
    assert "error" in bad
    good = propose_claim(ctx, "Fewer than 50% of papers use MIMIC.", "prevalence",
                         {"field": "datasets", "any_of": ["MIMIC"], "max_share": 0.5})
    again = propose_claim(ctx, "Fewer than 50% of papers use MIMIC.", "prevalence",
                          {"field": "datasets", "any_of": ["MIMIC"], "max_share": 0.5})
    assert again["status"] == "duplicate" and again["claim_id"] == good["claim_id"]
    row = ctx.pg.execute("SELECT predicate FROM claims WHERE id=%s", (good["claim_id"],)).fetchone()
    assert row["predicate"]["strict"] is True


def test_deepseek_adapter_request_and_response():
    import json as _json

    import httpx

    from research_agent.llm.deepseek_client import DeepSeekClient

    seen = []

    def handler(request):
        body = _json.loads(request.content)
        seen.append(body)
        if len(seen) == 1:
            return httpx.Response(200, json={
                "choices": [{"finish_reason": "tool_calls", "message": {
                    "content": "", "reasoning_content": "thinking about it",
                    "tool_calls": [{"id": "c1", "type": "function",
                                    "function": {"name": "finish", "arguments": '{"summary": "ok"}'}}]}}],
                "usage": {"prompt_tokens": 100, "prompt_cache_hit_tokens": 80, "completion_tokens": 7}})
        return httpx.Response(402, json={"error": {"message": "Insufficient Balance"}})

    client = DeepSeekClient("deepseek-flash", api_key="k", thinking="enabled",
                            transport=httpx.MockTransport(handler))
    tools = [{"name": "finish", "description": "d", "input_schema": {"type": "object", "properties": {}}}]
    resp = client.chat("sys", [{"role": "user", "content": [{"type": "text", "text": "hi"}]}],
                       tools=tools, force_tool="finish", max_tokens=500)
    assert seen[0]["tool_choice"] == {"type": "function", "function": {"name": "finish"}}
    assert seen[0]["thinking"] == {"type": "enabled"} and "temperature" not in seen[0]
    assert resp.tool_calls[0].input == {"summary": "ok"}
    assert client.usage.cached_input_tokens == 80 and client.usage.output_tokens == 7

    # reasoning must be sent back on the next request
    history = [{"role": "user", "content": [{"type": "text", "text": "hi"}]},
               {"role": "assistant", "content": resp.content},
               {"role": "user", "content": [{"type": "tool_result", "tool_use_id": "c1", "content": "{}"}]}]
    with pytest.raises(Exception) as err:
        client.chat("sys", history, tools=tools)
    assert seen[1]["messages"][2]["reasoning_content"] == "thinking about it"
    assert "402" in str(err.value) and "top up" in str(err.value)


def test_claim_scope_and_same_evidence_dedupe(ctx, monkeypatch):
    from research_agent.tools import claims as C

    assert C.scope_problem("Diffusion methods appear in at least 3% of 2024 arXiv health papers.")
    assert C.scope_problem("In the 16-paper shortlist, U-Net appears in most papers.") is None
    assert "error" in C.propose_claim(ctx, "CNNs appear in at least 6% of all health papers in the corpus.",
                                      "prevalence", {"field": "methods", "any_of": ["CNN"], "min_share": 0.06})

    rows = [{"paper_id": f"p{i}", "source": "fulltext",
             "data": {"methods": ["U-Net"] if i < 3 else ["random forest"]}} for i in range(5)]
    monkeypatch.setattr(C, "_rows", lambda _ctx: rows)
    first = C.propose_claim(ctx, "U-Net is used by most papers in the shortlist.", "prevalence",
                            {"field": "methods", "any_of": ["U-Net"], "min_share": 0.5})
    same = C.propose_claim(ctx, "U-Net variants remain widely used in the shortlist.", "prevalence",
                           {"field": "methods", "any_of": ["U-Net", "UNet", "nnU-Net"], "min_share": 0.4})
    assert same["status"] == "duplicate" and same["claim_id"] == first["claim_id"]
    rare = C.propose_claim(ctx, "Random forests are a minority in the shortlist.", "prevalence",
                           {"field": "methods", "any_of": ["random forest"], "max_share": 0.5})
    assert rare["status"] == "pending"


def test_truncated_finish_is_retried_then_salvaged(ctx):
    from research_agent.agents.base import Agent
    from research_agent.llm.base import LLMResponse, ToolCall

    raw = '{"summary": "short", "report_markdown": "## Summary\\n- U-Net dominates [C1]\\n- Transf'

    class CutOff:
        calls = 0

        def chat(self, system, messages, tools=None, force_tool=None, max_tokens=0, temperature=0.0):
            CutOff.calls += 1
            tc = ToolCall(f"f{CutOff.calls}", "finish", {"_raw_arguments": raw})
            return LLMResponse("", [tc], [{"type": "tool_use", "id": tc.id, "name": "finish", "input": tc.input}],
                               "length", {})

    schema = {"type": "object", "properties": {"summary": {"type": "string"},
                                               "report_markdown": {"type": "string"}}}
    agent = Agent("synthesis_probe", "r", "s", [], schema, max_turns=6)
    ctx.llm_factory = lambda strong=False: CutOff()
    out = agent.run(ctx, "write")
    assert CutOff.calls == 2  # asked once to shorten, then salvaged
    assert out["summary"] == "short"
    assert out["report_markdown"].startswith("## Summary\n- U-Net dominates [C1]")
    assert "cut off" in out["_note"]


def test_truncated_list_output_keeps_finished_items(ctx):
    from research_agent.agents.base import Agent, repair_truncated_json
    from research_agent.llm.base import LLMResponse, ToolCall
    from research_agent.opportunity.agents import DESIGN, GAP_REASONING
    from research_agent.opportunity.protocol import PROTOCOL

    raw = ('{"designs": [{"title": "A", "addresses": ["G1"], "why": "has \\"quotes\\", commas"}, '
           '{"title": "B", "addresses": ["G2", "G3"]}, {"title": "C half", "addre')
    seen_budgets = []

    class CutOff:
        def chat(self, system, messages, tools=None, force_tool=None, max_tokens=0, temperature=0.0):
            seen_budgets.append(max_tokens)
            tc = ToolCall(f"f{len(seen_budgets)}", "finish", {"_raw_arguments": raw})
            return LLMResponse("", [tc], [{"type": "tool_use", "id": tc.id, "name": "finish", "input": tc.input}],
                               "length", {})

    schema = {"type": "object", "properties": {"designs": {"type": "array", "items": {"type": "object"}}}}
    agent = Agent("design_probe", "r", "s", [], schema, max_turns=6, long_output=True)
    ctx.llm_factory = lambda strong=False: CutOff()
    out = agent.run(ctx, "design")
    # the two finished designs survive; the half-written third is dropped rather than kept half-empty
    assert [d["title"] for d in out["designs"]] == ["A", "B"] and out["designs"][0]["why"].startswith('has "quotes"')
    assert "cut off" in out["_note"] and len(seen_budgets) == 2
    assert seen_budgets[0] >= 8192 and all(a.long_output for a in (PROTOCOL, GAP_REASONING, DESIGN))
    assert repair_truncated_json('{"summary": "abc') is None and repair_truncated_json("not json") is None
    assert repair_truncated_json('{"a": 1, "b": [1, 2') == {"a": 1, "b": [1]}


def test_short_terms_match_whole_words(ctx, monkeypatch):
    from research_agent.tools import claims as C

    rows = [{"paper_id": "a", "source": "fulltext", "data": {"data_modalities": ["electron microscopy"]}},
            {"paper_id": "b", "source": "fulltext", "data": {"data_modalities": ["CT", "MRI scans"]}},
            {"paper_id": "c", "source": "fulltext", "data": {"data_modalities": ["contrast-enhanced CT"]}},
            {"paper_id": "d", "source": "fulltext", "data": {"data_modalities": ["CNNs", "U-Net++"]}}]
    monkeypatch.setattr(C, "_rows", lambda _ctx: rows)
    r = C.evaluate_prevalence(ctx, {"field": "data_modalities", "any_of": ["CT"], "min_count": 1})
    assert sorted(r["matched_paper_ids"]) == ["b", "c"]
    r = C.evaluate_prevalence(ctx, {"field": "data_modalities", "any_of": ["CNN", "U-Net"], "min_count": 1})
    assert r["matched_paper_ids"] == ["d"]


def test_wording_must_match_observed_share():
    from research_agent.tools.claims import wording_problem

    assert wording_problem("CT and MRI account for the large majority of papers", 0.625)
    assert wording_problem("Convolutions remain near-universal", 0.69)
    assert wording_problem("Convolutions remain near-universal", 0.81) is None
    assert wording_problem("Most papers use U-Net", 0.5)
    assert wording_problem("U-Net is the most common family", 0.4) is None
    assert wording_problem("Geography is almost never reported", 0.12) is None
    assert wording_problem("External validation is a small minority", 0.4)


def test_claim_citations_are_aligned_with_their_numbers():
    from research_agent.agents.report import _clean_author, align_claim_citations
    from research_agent.tools.claims import wording_problem

    def c(n, d, text, status="supported"):
        return {"status": status, "claim_type": "prevalence", "text": text,
                "result": {"n_matching": n, "denominator": d}}

    claims = {45: c(18, 60, "Transformer methods are mainstream"), 46: c(14, 60, "SAM foundation models"),
              47: c(16, 60, "Attention is near-ubiquitous", "unsupported"), 48: c(32, 60, "MRI and CT dominate"),
              49: c(18, 60, "Public benchmark datasets are concentrated"), 50: c(23, 60, "Transformer/attention")}
    body = ("SAM models appear in 14 of 60 papers [C48]. MRI and CT match 32 of 60 papers [C50].\n"
            "A benchmark query matches 18 of 60 papers [C50]. Attention appears in 16 of 60 [C46]. "
            "Transformers are common [C45].")
    out, fixes = align_claim_citations(body, claims)
    assert "14 of 60 papers [C46]" in out and "32 of 60 papers [C48]" in out
    assert "18 of 60 papers [C49]" in out                      # tie with C45 broken by wording
    assert "16 of 60 [C46 — number does not match this claim]" in out
    assert "Transformers are common [C45]" in out              # no number: left alone
    assert _clean_author(r"Ziya Ata Yaz{\i}c{\i}") == "Ziya Ata Yazıcı"
    assert _clean_author("Douwe J. Spaanderman (1)") == "Douwe J. Spaanderman"
    assert wording_problem("Diffusion is rare, at most 15% of papers", 0.05) is None


def test_brief_never_drops_claims_when_notes_are_long(ctx):
    import json as _json

    from research_agent.agents.base import _dump
    from research_agent.agents.report import REPORT_TOOLS, get_brief
    from research_agent.tools.claims import propose_claim

    ids = [propose_claim(ctx, f"In the shortlist, method {i} appears in fewer than half of papers.", "prevalence",
                         {"field": "methods", "any_of": [f"m{i}"], "max_share": 0.5})["claim_id"] for i in range(15)]
    for agent in ("discovery", "literature", "methods", "trends", "gaps", "evidence"):
        ctx.save_note(agent, {"summary": "x" * 5000, "old_ids": "C1 C2 C3"})
    tool = REPORT_TOOLS[0]
    text = _dump(get_brief(ctx), tool.max_chars)
    assert all(f'"C{i}"' in text for i in ids)
    assert "truncated" not in text[: text.index('"agent_outputs"')]
    _json.loads(text)  # still valid JSON


def test_claim_text_numbers_and_negated_wording():
    from research_agent.tools.claims import text_count_problem, wording_problem

    # a claim about what is missing is judged on the share that does not match
    assert wording_problem("Most papers are silent on validation level.", 0.30) is None
    assert wording_problem("Most papers report a validation level.", 0.30)
    # counts written into the claim text must be the counted ones; thresholds are not counts
    # 18 silent is the complement of 28 stating, so the number itself is fine
    assert text_count_problem("Validation is unreported in a minority (18/46 silent).", 28, 46) is None
    assert text_count_problem("Validation is unreported in a minority (12/46 silent).", 28, 46)
    assert text_count_problem("Named datasets are reported in fewer than half (19/46).", 14, 46)
    assert text_count_problem("Human judgment appears in at most 4 of 60 papers.", 3, 60) is None
    assert text_count_problem("Classical time series appear in 9+ papers.", 8, 46) is None
    assert text_count_problem("Code availability is stated in 9/46 papers.", 9, 46) is None
    # a count over a population the check did not measure is not evidence for the claim (C432 in a real run:
    # "15 of 58 (26%)" passed while the check counted 19 of 100)
    assert "out of 33" in text_count_problem("Benchmarks appear in 8 of the 33 dataset papers.", 10, 60)
    assert text_count_problem("In the 58 papers with methods, ML appears in 15 of 58 (26%).", 19, 100)
    assert text_count_problem("Used by 3 of 5 models in one paper.", 3, 60) is None   # not a paper population


def test_claim_percentages_must_match_the_measured_share():
    from research_agent.tools.claims import percent_problem

    assert percent_problem("In 58 shortlisted papers, 50 (86%) use climate covariates.", 0.80)
    assert percent_problem("Climate covariates appear in 80 of 100 papers (80%).", 0.80) is None
    assert percent_problem("Silent on code in 85% of papers.", 0.15) is None                 # complement
    assert percent_problem("Most models report accuracy above 95%.", 0.60) is None           # threshold
    assert percent_problem("17 of 100 report 95% credible intervals (17%).", 0.17) is None   # interval level


def test_text_thresholds_and_complements(ctx, monkeypatch):
    from research_agent.tools import claims as C
    from research_agent.agents.report import align_claim_citations

    assert C.text_count_problem("benchmarks appear in at least 8 of the 33 papers", 6, 33)
    assert C.text_count_problem("benchmarks appear in at least 8 of the 33 papers", 10, 33) is None
    assert C.text_count_problem("human judgment appears in at most 4 of 60 papers", 7, 60)
    assert C.text_count_problem("datasets are unnamed in 32 of 46 papers", 14, 46) is None  # complement

    claims = {86: {"status": "supported", "claim_type": "prevalence", "text": "Named datasets are rare",
                   "result": {"n_matching": 14, "denominator": 46}}}
    out, fixes = align_claim_citations("No dataset is named in 32 of 46 papers [C86].", claims)
    assert "number does not match" not in out and not fixes

    rows = [{"paper_id": f"p{i}", "source": "fulltext",
             "data": {"methods": ["ARIMA"] if i < 9 else (["Box-Jenkins"] if i == 9 else ["LSTM"])}}
            for i in range(12)]
    monkeypatch.setattr(C, "_rows", lambda _ctx: rows)
    first = C.propose_claim(ctx, "ARIMA is used by most shortlisted papers.", "prevalence",
                            {"field": "methods", "any_of": ["ARIMA"], "min_share": 0.5})
    near = C.propose_claim(ctx, "Classical time series appear in most shortlisted papers.", "prevalence",
                           {"field": "methods", "any_of": ["ARIMA", "Box-Jenkins"], "min_count": 5})
    assert near["status"] == "pending" and f"C{first['claim_id']}" in near["note"]


# ---------------------------------------------------------------- PubMed Central
def _pmc_article(pmcid, title, licence, year=2021, abstract="A study of malaria forecasting in Kenya "
                 "using routine surveillance data and gradient boosting models over five seasons."):
    return f"""<article><front>
      <journal-meta><journal-title>Journal of Test Medicine</journal-title></journal-meta>
      <article-meta>
        <article-id pub-id-type="pmcid">{pmcid}</article-id>
        <article-id pub-id-type="doi">10.1234/{pmcid}</article-id>
        <title-group><article-title>{title}</article-title></title-group>
        <contrib-group><contrib contrib-type="author"><name><surname>Otieno</surname>
          <given-names>A</given-names></name></contrib></contrib-group>
        <pub-date pub-type="epub"><year>{year}</year></pub-date>
        <abstract><p>{abstract}</p></abstract>
        <permissions><license license-type="{licence}">
          <ext-link xmlns:xlink="http://www.w3.org/1999/xlink"
            xlink:href="https://creativecommons.org/licenses/{licence}/4.0/">licence</ext-link>
        </license></permissions>
        <subj-group><subject>Public Health</subject></subj-group>
      </article-meta></front>
      <body><sec><title>Methods</title><p>We fitted a gradient boosting model.</p>
        <table-wrap><p>dropped table</p></table-wrap></sec>
        <sec><title>Limitations</title><p>Single district, no external validation.</p></sec></body>
    </article>"""


def _pmc_transport(articles, fail_first: int = 0):
    """A fake NCBI E-utilities server holding `articles`; optionally fails the first N efetch calls."""
    import re as _re
    from urllib.parse import parse_qs

    import httpx

    by_id = {_re.search(r"PMC(\d+)", a).group(1): a for a in articles}
    state = {"fails": fail_first}

    def handler(request):
        url = str(request.url)
        params = {k: v for k, v in request.url.params.items()}
        if request.content:
            params.update({k: v[0] for k, v in parse_qs(request.content.decode()).items()})
        if "/esearch.fcgi" in url and params.get("retmax") == "0":      # per-year denominators
            return httpx.Response(200, text="<eSearchResult><Count>1234</Count></eSearchResult>")
        if "/esearch.fcgi" in url:
            ids = "".join(f"<Id>{i}</Id>" for i in by_id)
            return httpx.Response(200, text=f"<eSearchResult><Count>{len(by_id)}</Count>"
                                            f"<IdList>{ids}</IdList></eSearchResult>")
        if state["fails"]:
            state["fails"] -= 1
            return httpx.Response(503, text="busy")
        wanted = [i.removeprefix("PMC") for i in params.get("id", "").split(",") if i]
        return httpx.Response(200, text="<pmc-articleset>" + "".join(by_id[i] for i in wanted if i in by_id)
                                        + "</pmc-articleset>")

    return httpx.MockTransport(handler)


def test_pmc_licence_filter_and_parsing():
    import xml.etree.ElementTree as ET

    from research_agent.ingestion import pmc

    assert pmc.licence_allows_reuse("cc-by https://creativecommons.org/licenses/by/4.0/")
    assert pmc.licence_allows_reuse("CC0 public domain")
    assert not pmc.licence_allows_reuse("cc-by-nc https://creativecommons.org/licenses/by-nc/4.0/")
    assert not pmc.licence_allows_reuse("")

    row = pmc.parse_article(ET.fromstring(_pmc_article("PMC111", "Malaria forecasting", "by")))
    assert row["paper_id"] == "PMC111" and row["year"] == 2021 and row["doi"] == "10.1234/PMC111"
    assert row["authors"] == "A Otieno" and row["categories"] == ["Public Health"]
    assert pmc.parse_article(ET.fromstring("<article><front/></article>")) is None  # no abstract


def test_pmc_ingest_writes_source_tagged_papers(loaded_db):
    from research_agent.db import get_conn
    from research_agent.ingestion import pmc

    client = pmc.PMCClient(transport=_pmc_transport([
        _pmc_article("PMC111", "Malaria forecasting with boosted trees", "by"),
        _pmc_article("PMC222", "Non-commercial study of malaria", "by-nc"),
    ]))
    res = pmc.ingest("malaria forecasting", from_year=2010, log=lambda *_: None, client=client)
    assert res["inserted"] == 1 and res["skipped_licence"] == 1

    with get_conn() as pg:
        row = pg.execute("SELECT source, title, journal_ref, health_reason FROM papers "
                         "WHERE paper_id='PMC111'").fetchone()
        n_arxiv = pg.execute("SELECT count(*) n FROM papers WHERE source='arxiv'").fetchone()["n"]
        stats = pg.execute("SELECT count(*) n FROM corpus_year_stats WHERE source='pmc'").fetchone()["n"]
    assert row["source"] == "pmc" and row["journal_ref"] == "Journal of Test Medicine"
    assert row["health_reason"].startswith("pmc:") and n_arxiv > 0   # arXiv papers untouched
    assert stats > 0                                                  # PMC denominators loaded

    again = pmc.ingest("malaria forecasting", from_year=2010, log=lambda *_: None, client=client)
    assert again["inserted"] == 0                                     # resumable


def test_pmc_fulltext_and_source_aware_tools(loaded_db, ctx):
    from research_agent.db import get_conn
    from research_agent.ingestion import pmc
    from research_agent.tools.search import corpus_count, hybrid_search
    from research_agent.tools.trends import corpus_sources, topic_series

    client = pmc.PMCClient(transport=_pmc_transport([
        _pmc_article("PMC111", "Malaria forecasting with boosted trees", "by")]))
    pmc.ingest("malaria forecasting", from_year=2010, log=lambda *_: None, client=client)

    with get_conn() as pg:
        status = pmc.fetch_fulltext(pg, ["PMC111"], client=client)
        text = pg.execute("SELECT clean_text, sections FROM paper_fulltext "
                          "WHERE paper_id='PMC111'").fetchone()
    assert status["PMC111"] in ("ok", "stub", "template_suspect")
    assert "gradient boosting" in text["clean_text"] and "dropped table" not in text["clean_text"]
    assert [s["heading"] for s in text["sections"]] == ["Methods", "Limitations"]

    assert corpus_sources(ctx)["sources"].keys() >= {"arxiv", "pmc"}
    assert corpus_count(ctx, "malaria")["by_source"].get("pmc", 0) == 1
    only_pmc = hybrid_search(ctx, "malaria forecasting", source="pmc", limit=10)["results"]
    assert only_pmc and all(r["source"] == "pmc" for r in only_pmc)
    assert all(r["source"] == "arxiv" for r in hybrid_search(ctx, "malaria", source="arxiv")["results"])
    assert topic_series(ctx, "malaria", source="pmc")["series"]     # normalised against PMC totals


def test_schema_upgrades_a_pre_pmc_database(loaded_db):
    """An existing database from before PMC support must migrate in place, keeping its papers."""
    from research_agent.db import get_conn, init_schema

    with get_conn() as pg:
        n_before = pg.execute("SELECT count(*) n FROM papers").fetchone()["n"]
        # put the database back into its old shape
        pg.execute("DELETE FROM corpus_year_stats WHERE source <> 'arxiv'")
        pg.execute("ALTER TABLE papers DROP COLUMN source")
        pg.execute("ALTER TABLE corpus_year_stats DROP COLUMN source")
        pg.execute("ALTER TABLE corpus_year_stats ADD PRIMARY KEY (year, primary_category)")

    init_schema()  # must not fail
    init_schema()  # and must be safe to run again

    with get_conn() as pg:
        assert pg.execute("SELECT count(*) n FROM papers").fetchone()["n"] == n_before
        assert pg.execute("SELECT count(*) n FROM papers WHERE source='arxiv'").fetchone()["n"] == n_before
        pk = pg.execute("""SELECT array_length(conkey, 1) k FROM pg_constraint
                           WHERE conname = 'corpus_year_stats_pkey'""").fetchone()["k"]
        idx = pg.execute("SELECT 1 FROM pg_indexes WHERE indexname = 'papers_source_idx'").fetchone()
    assert pk == 3 and idx


def test_pmc_noncommercial_flag_and_citations(loaded_db, ctx):
    from research_agent.agents.report import finalize_report
    from research_agent.db import get_conn
    from research_agent.ingestion import pmc
    from research_agent.tools.search import add_to_shortlist

    assert pmc.licence_is_noncommercial("by-nc https://creativecommons.org/licenses/by-nc/4.0/")
    assert not pmc.licence_is_noncommercial("by https://creativecommons.org/licenses/by/4.0/")
    assert not pmc.licence_is_noncommercial("")

    articles = [_pmc_article("PMC501", "Malaria forecasting with boosted trees", "by"),
                _pmc_article("PMC502", "Malaria early warning in Kenya", "by-nc")]
    client = pmc.PMCClient(transport=_pmc_transport(articles))
    first = pmc.ingest("malaria nc test", from_year=2010, log=lambda *_: None, client=client)
    assert first["inserted"] == 1 and first["skipped_licence"] == 1
    # rerun with the flag: only the skipped non-commercial article is added
    second = pmc.ingest("malaria nc test", from_year=2010, log=lambda *_: None, client=client,
                        include_noncommercial=True)
    assert second["inserted"] == 1 and second["skipped_licence"] == 0

    add_to_shortlist(ctx, ["PMC501", "PMC502"], reason="t")
    report, audit = finalize_report(
        ctx, "## Summary\n- Boosted trees [PMC501] and early warning [arXiv:PMC502]; bogus [PMC999].")
    assert "[PMC501]" in report and "[PMC502]" in report
    assert "PMC999" not in report.split("## Run facts")[0].replace("citation removed", "")
    assert audit["removed_paper_citations"] == ["PMC999"]
    refs = report.split("## References")[1]
    assert "https://pmc.ncbi.nlm.nih.gov/articles/PMC502/ *(licence: non-commercial)*" in refs
    assert "PMC501/" in refs and "PMC501/ *(licence" not in refs
    with get_conn() as pg:
        pg.execute("DELETE FROM run_papers WHERE paper_id IN ('PMC501','PMC502')")


def test_pmc_search_splits_large_results_and_retries(monkeypatch):
    from datetime import date

    import httpx

    from research_agent.ingestion import pmc

    # a fake search where every year holds 3 ids, and one request may return at most 4
    calls = []

    def handler(request):
        term = request.url.params["term"]
        calls.append(term)
        years = [int(y) for y in __import__("re").findall(r'"(\d{4})/\d\d/\d\d"', term)]
        span = range(years[0], years[1] + 1)
        ids = [f"{y}{k}" for y in span for k in range(3)]
        shown = "".join(f"<Id>{i}</Id>" for i in ids[:4])
        return httpx.Response(200, text=f"<eSearchResult><Count>{len(ids)}</Count>"
                                        f"<IdList>{shown}</IdList></eSearchResult>")

    monkeypatch.setattr(pmc, "MAX_IDS", 4)
    client = pmc.PMCClient(transport=httpx.MockTransport(handler))
    ids = pmc.search_ids(client, "malaria", 2018, 2021, log=lambda *_: None)
    assert len(ids) == 12 and len(set(ids)) == 12 and all(i.startswith("PMC") for i in ids)
    assert len(calls) > 1                       # the range was split to stay under the cap

    # One publication date over the cap: PubMed dates an article with no stated day to 1 January, so a single
    # "day" can hold a year of them. That day is split again by the date each record entered PubMed.
    seen = []

    def jan_first(request):
        term = request.url.params["term"]
        seen.append(term)
        days = __import__("re").findall(r'"(\d{4}/\d\d/\d\d)"\[PDAT\]', term)
        if "[CRDT]" in term:                      # the 1 January pile, re-split by creation date: small windows
            lo = int(__import__("re").findall(r'"(\d{4})/\d\d/\d\d"\[CRDT\]', term)[0])
            ids = [f"crdt{lo}{k}" for k in range(3)]
        elif days[0] <= "2012/01/01" <= days[1]:  # every article with no stated day, piled on one date
            return httpx.Response(200, text="<eSearchResult><Count>10110</Count><IdList>"
                                            "<Id>a</Id><Id>b</Id><Id>c</Id><Id>d</Id></IdList></eSearchResult>")
        else:
            ids = ["x"] if days[0] == days[1] else ["x", "y"]
        return httpx.Response(200, text=f"<eSearchResult><Count>{len(ids)}</Count><IdList>"
                                        + "".join(f"<Id>{i}</Id>" for i in ids) + "</IdList></eSearchResult>")

    warned = []
    jan = pmc.PMCClient(transport=httpx.MockTransport(jan_first))
    jan._min_gap = 0
    got = pmc.search_ids(jan, "malaria", 2012, 2012, log=warned.append)
    assert any("[CRDT]" in t and '"2012/01/01"[PDAT]' in t for t in seen)   # the day kept, split on creation date
    assert any("crdt" in i for i in got) and not warned                     # ids the cap would have hidden
    assert all(i.startswith("PMC") for i in got)

    # transient 503s are retried instead of killing a long ingest
    flaky = pmc.PMCClient(transport=_pmc_transport([_pmc_article("PMC777", "Malaria", "by")], fail_first=2))
    flaky._min_gap = 0
    monkeypatch.setattr(pmc.time, "sleep", lambda *_: None)
    assert [r["paper_id"] for r in pmc.fetch_ids(flaky, ["PMC777"])] == ["PMC777"]


def test_licence_rules_and_audit_of_loaded_papers(loaded_db):
    from research_agent.db import get_conn
    from research_agent.ingestion import pmc

    k = pmc.licence_kind
    assert k("© 2020 by the authors. All rights reserved.") == "unknown"   # a bare "by" is not CC BY
    assert k("open-access") == "unknown"
    assert k("open-access © 2019 by the authors. Licensee MDPI. " + "x" * 250 +
             " distributed under the Creative Commons Attribution (CC BY) license.") == "reusable"
    assert k("Creative Commons Attribution-NonCommercial 4.0 International License") == "noncommercial"
    assert k("https://creativecommons.org/publicdomain/zero/1.0/") == "reusable"

    with get_conn() as pg:
        base = dict(abstract="x" * 100, categories=["pmc"], year=2020, health_reason="pmc:test")
        for pid, lic in [("PMC9001", "cc-by https://creativecommons.org/licenses/by/4.0/"),
                         ("PMC9002", "© 2020 by the authors"),          # would have slipped in before
                         ("PMC9003", "by-nc https://creativecommons.org/licenses/by-nc/4.0/")]:
            pg.execute("INSERT INTO papers (paper_id, source, title, abstract, categories, year, license, "
                       "health_reason) VALUES (%s,'pmc',%s,%s,%s,%s,%s,%s)",
                       (pid, pid, base["abstract"], base["categories"], base["year"], lic, base["health_reason"]))
        dry = pmc.audit_licences(pg)
        assert dry["removed"] == 0 and dry["failing"] >= 2
        kept_nc = pmc.audit_licences(pg, remove=True, allow_noncommercial=True)
        left = {r["paper_id"] for r in pg.execute(
            "SELECT paper_id FROM papers WHERE paper_id IN ('PMC9001','PMC9002','PMC9003')").fetchall()}
        assert kept_nc["removed"] >= 1 and left == {"PMC9001", "PMC9003"}
        pg.execute("DELETE FROM papers WHERE paper_id IN ('PMC9001','PMC9003')")


def test_trend_claim_numbers_and_scope(ctx):
    from research_agent.tools.claims import propose_claim, trend_text_problem

    # real claims from the mixed-corpus malaria report
    assert trend_text_problem("ML/DL rose sharply, roughly a 9-fold increase", 5.32)            # C172
    assert trend_text_problem("validation grew only marginally (ratio ~1.1)", 2.38)             # C184
    assert trend_text_problem("did not increase (ratio <= 1.0)", 2.78)                          # C185
    assert trend_text_problem("grew in normalised share (ratio >= 1.5)", 4.83) is None         # C183
    assert trend_text_problem("grew several-fold", 6.6) is None                                 # no number
    assert trend_text_problem("(ratio ~3.2)", 3.17) is None and trend_text_problem("an ~1.8x rise", 1.8) is None
    # only the first multiplier is the claim's own; later ones describe a comparison topic
    assert trend_text_problem("rose ~2.1x while bias terms rose only ~1.5x", 2.07) is None

    bad = propose_claim(ctx, "ML use increased within the shortlist over time.", "trend",
                        {"keywords": "machine learning", "early": [2015, 2017], "late": [2023, 2025],
                         "direction": "increase"})
    assert "error" in bad
    ok = propose_claim(ctx, "In PMC, malaria ML papers grew.", "trend",
                       {"keywords": "malaria machine learning", "early": [2015, 2017], "late": [2023, 2025],
                        "direction": "increase"})
    row = ctx.pg.execute("SELECT predicate FROM claims WHERE id=%s", (ok["claim_id"],)).fetchone()
    assert row["predicate"]["source"] == "pmc"


def test_preprint_and_published_version_are_not_both_shortlisted(loaded_db, ctx):
    from research_agent.db import get_conn
    from research_agent.tools.search import add_to_shortlist

    with get_conn() as pg:
        for pid, src, title, doi in [
            ("2412.99999", "arxiv", "Forecasting Malaria in Indian States: A Time Series Approach with R Shiny "
                                    "Integration", None),
            ("PMC8888888", "pmc", "Forecasting Malaria in Indian States: A Time Series Approach with R Shiny "
                                  "Integration", "10.1/x"),
            ("1906.99999", "arxiv", "Data-Driven Malaria Prevalence Prediction in Large Densely-Populated Urban "
                                    "Holoendemic sub-Saharan West Africa: Harnessing Machine Learning", None),
            ("PMC7777777", "pmc", "Data-driven malaria prevalence prediction in large densely populated urban "
                                  "holoendemic sub-Saharan West Africa", None)]:
            pg.execute("INSERT INTO papers (paper_id, source, title, abstract, authors, categories, year, doi, "
                       "health_reason) VALUES (%s,%s,%s,%s,%s,%s,2024,%s,'t') ON CONFLICT DO NOTHING",
                       (pid, src, title, "x" * 100, "Sujit K. Ghosh, A. B" if "Indian" in title
                        else "Biobele J. Brown, C. D", ["t"], doi))

    first = add_to_shortlist(ctx, ["2412.99999", "1906.99999"], reason="t")
    assert first["added"] == 2
    # published versions replace the preprints already on the list
    second = add_to_shortlist(ctx, ["PMC8888888", "PMC7777777"], reason="t")
    assert len(second["replaced_preprints"]) == 2
    ids = set(ctx.shortlist_ids())
    assert {"PMC8888888", "PMC7777777"} <= ids and not {"2412.99999", "1906.99999"} & ids
    # and a preprint is not added back when its published version is there
    third = add_to_shortlist(ctx, ["2412.99999"], reason="t")
    assert third["added"] == 0 and third["skipped_duplicates"][0]["same_work_as"] == "PMC8888888"
    with get_conn() as pg:
        pg.execute("DELETE FROM run_papers WHERE run_id=%s", (ctx.run_id,))
        pg.execute("DELETE FROM papers WHERE paper_id IN ('2412.99999','PMC8888888','1906.99999','PMC7777777')")


def test_unsupported_counts_may_be_quoted_but_not_rejected_ones():
    from research_agent.agents.report import align_claim_citations, audit_numbers

    def c(n, d, status, text="t"):
        return {"status": status, "claim_type": "prevalence", "text": text,
                "result": {"n_matching": n, "denominator": d}}

    claims = {166: c(24, 56, "unsupported", "classical methods majority"),
              167: c(27, 56, "supported", "ML minority"), 109: c(10, 60, "rejected")}
    body = "Classical methods reach 24 of 56, short of half (claim C166, unsupported), and ML is a minority [C167]."
    out, fixes = align_claim_citations(body, claims)
    assert "number does not match" not in out and not fixes
    allowed = {(24, 56), (32, 56), (27, 56), (29, 56)}   # what finalize_report derives for C166 and C167
    assert audit_numbers("24 of 56 and 10 of 60", allowed)[1] == 1   # the rejected claim's count is flagged


def test_pmc_year_is_publication_not_collection():
    import xml.etree.ElementTree as ET

    from research_agent.ingestion import pmc

    xml = _pmc_article("PMC123", "Malaria forecasting", "by", year=2026).replace(
        '<pub-date pub-type="epub"><year>2026</year></pub-date>',
        '<pub-date pub-type="collection"><year>2027</year></pub-date>'
        '<pub-date pub-type="epub"><year>2026</year></pub-date>')
    assert pmc.parse_article(ET.fromstring(xml))["year"] == 2026


def test_extraction_evidence_is_checked_against_the_text():
    from research_agent.tools.extraction import check_evidence, normalise

    text = ("Title: Forecasting malaria in Kenya\n\nAbstract: We trained a gradient boosting model on "
            "weekly surveillance counts from 42 sites in western Kenya, with external validation on two "
            "held-out counties.")
    raw = {"methods": ["gradient boosting"], "geography": ["Kenya"], "datasets": ["MIMIC-IV"],
           "data_modalities": ["surveillance counts"], "validation_level": "external_site",
           "code_or_data_available": "yes"}
    evidence = {
        "methods": ["We trained a gradient boosting model on weekly"],          # verbatim
        "geography": ["from 42 sites in western Kenya,"],                      # verbatim, punctuation differs
        "data_modalities": ["weekly surveillance counts from 42 sites"],
        "validation_level": ["with external validation on two held-out counties"],
        "datasets": ["the MIMIC-IV database was used"],                        # invented: not in the text
        # code_or_data_available: no quote at all
    }
    out = check_evidence(normalise(raw), evidence, text)
    assert out["methods"] == ["gradient boosting"] and out["geography"] == ["Kenya"]
    assert out["validation_level"] == "external_site"
    assert out["datasets"] == [] and out["code_or_data_available"] == "not_stated"   # dropped
    assert out["_unverified"] == {"datasets": ["MIMIC-IV"], "code_or_data_available": "yes"}
    assert set(out["evidence"]) == {"methods", "geography", "data_modalities", "validation_level"}
    # a paraphrase is not a quote
    para = check_evidence(normalise({"methods": ["gradient boosting"]}),
                          {"methods": ["the authors used a boosted tree ensemble"]}, text)
    assert para["methods"] == [] and "methods" in para["_unverified"]


def test_cut_off_extraction_fails_loudly():
    from research_agent.llm.base import LLMResponse, ToolCall
    from research_agent.tools.reading import _read_one

    class CutOff:
        model = "x"

        def chat(self, *a, **k):
            tc = ToolCall("1", "record_reading", {"_raw_arguments": '{"problem": "malaria", "meth'})
            return LLMResponse("", [tc], [], "length", {})

    class Ctx:
        question = "q"

    import pytest as _pytest
    with _pytest.raises(ValueError, match="cut off"):
        _read_one(CutOff(), Ctx(), {"paper_id": "p", "title": "t", "abstract": "a"}, None, True)


def test_claims_are_measured_before_they_are_written(ctx, monkeypatch):
    from research_agent.tools import claims as C

    rows = [{"paper_id": f"p{i}", "source": "fulltext", "year": 2020, "title": "t",
             "data": {"methods": ["ARIMA"] if i < 9 else ["LSTM"]}} for i in range(46)]
    monkeypatch.setattr(C, "_rows", lambda _ctx: rows)

    dry = C.test_claim(ctx, "prevalence", {"field": "methods", "any_of": ["ARIMA"], "min_count": 5})
    assert dry["n_matching"] == 9 and dry["denominator"] == 46 and dry["bounds_pass"]
    assert not ctx.pg.execute("SELECT 1 FROM claims WHERE run_id=%s", (ctx.run_id,)).fetchone()  # nothing stored

    # text written before measuring, contradicting the count: refused, with the real number returned
    bad = C.propose_claim(ctx, "Classical methods appear in 18 of 46 papers.", "prevalence",
                          {"field": "methods", "any_of": ["ARIMA"], "min_count": 5})
    assert "error" in bad and "9 of 46" in bad["measured"]
    good = C.propose_claim(ctx, "Classical methods appear in 9 of 46 shortlisted papers.", "prevalence",
                           {"field": "methods", "any_of": ["ARIMA"], "min_count": 5})
    assert good["status"] == "pending" and good["measured_now"].startswith("9 of 46")


# ---------------------------------------------------------------- Research Opportunity Map
def test_protocol_validation_cleans_fields():
    from research_agent.opportunity.protocol import validate_protocol

    clean, problems = validate_protocol({"fields": [
        {"name": "Forecast Horizon", "type": "enum", "values": ["1 month", "3-6 months", "not stated"],
         "definition": "how far ahead", "desirable": ["3-6 months", "nonsense"],
         "search": {"3-6 months": '"six months"'}},
        {"name": "methods", "type": "list", "definition": "already a base field"},
        {"name": "lonely", "type": "enum", "values": ["only"], "definition": "x"},
        {"name": "forecast_horizon", "type": "enum", "values": ["a", "b"], "definition": "duplicate"},
        {"name": "covariates", "type": "list", "values": ["rainfall"], "definition": "inputs",
         "desirable": ["Intervention coverage"]}]})
    f = {x["name"]: x for x in clean["fields"]}
    assert set(f) == {"q_forecast_horizon", "q_covariates"} and len(problems) == 3
    assert f["q_forecast_horizon"]["values"] == ["1_month", "3_6_months", "not_stated"]
    assert f["q_forecast_horizon"]["desirable"] == ["3_6_months"]
    assert f["q_forecast_horizon"]["search"] == {"3_6_months": '"six months"'}
    assert f["q_covariates"]["desirable"] == ["intervention coverage"]


def test_subgroup_counts_with_where(ctx, monkeypatch):
    from research_agent.tools import claims as C

    rows = []
    for i in range(20):
        rows.append({"paper_id": f"p{i}", "source": "fulltext" if i < 8 else "abstract",
                     "corpus": "pmc" if i % 2 else "arxiv", "year": 2012 + i % 12, "title": "t",
                     "data": {"geography": ["Kenya"] if i < 10 else ["Brazil"],
                              "data_modalities": ["satellite imagery"] if i in (0, 1, 2, 11) else ["survey"],
                              "validation_level": "external" if i in (0, 3) else "internal"}})
    monkeypatch.setattr(C, "_rows", lambda _ctx: rows)
    sat_kenya = C.test_claim(ctx, "prevalence", {"field": "data_modalities", "any_of": ["satellite"],
                                                 "min_count": 1, "where": {"field": "geography", "any_of": ["Kenya"]}})
    assert (sat_kenya["n_matching"], sat_kenya["denominator"]) == (3, 10)
    pmc_ext = C.evaluate_prevalence(ctx, {"field": "validation_level", "any_of": ["external"], "min_count": 1,
                                          "where": [{"field": "source", "any_of": ["pmc"]}]})
    assert (pmc_ext["n_matching"], pmc_ext["denominator"]) == (1, 10) and pmc_ext["subgroup"]["size"] == 10
    not_kenya_recent = C.evaluate_prevalence(ctx, {
        "field": "data_modalities", "any_of": ["satellite"], "min_count": 1,
        "where": [{"field": "geography", "none_of": ["Kenya"]}, {"field": "year", "min": 2020},
                  {"field": "read", "any_of": ["abstract"]}]})
    assert not_kenya_recent["subgroup"]["size"] < 5 and "too few" in not_kenya_recent["caveat"]
    assert "error" in C.test_claim(ctx, "prevalence", {"field": "methods", "any_of": ["x"], "min_count": 1,
                                                       "where": [{"field": "nonsense", "any_of": ["y"]}]})


def test_research_opportunity_map_end_to_end(loaded_db):
    from research_agent.db import get_conn
    from research_agent.opportunity.pipeline import run_map

    with get_conn() as pg:   # a database from before the map existed has no protocol table
        pg.execute("DROP TABLE IF EXISTS protocol_extractions")
    events = []
    res = run_map("Can machine learning improve 1-6 month malaria forecasting?",
                  llm_factory=lambda strong=False: FakeLLM(), on_event=lambda a, k, p: events.append((a, k)))
    md, m = res["report"], res["map"]
    for section in ("## What is established", "## What is emerging", "## What is missing",
                    "## Evidence strength at a glance", "## Potential novelty", "## Candidate research designs",
                    "## Protocol", "## Method", "## References"):
        assert section in md, section

    labels = {g["label"] for g in m["gaps"]}
    # desirable practices no fixture paper has: 3-6 month horizon, spatial holdout, probabilistic output
    assert {"forecast horizon: three to six months", "validation split: tests on unseen places",
            "probabilistic: yes"} <= labels
    # grouped values make one gap (the group), never one gap per member
    assert not labels & {"validation split: spatial holdout", "validation split: spatiotemporal holdout"}
    group = next(g for g in m["gaps"] if g["label"] == "validation split: tests on unseen places")
    assert group["members"] == ["spatial holdout", "spatiotemporal holdout"] and group["role"] == "evaluation"
    assert "counts any of: spatial holdout, spatiotemporal holdout" in md
    assert all(0 <= g["ci95"][0] <= g["share"] <= g["ci95"][1] <= 1 for g in m["gaps"])
    assert all(g["confidence"]["grade"] in ("high", "moderate", "low") and g["confidence"]["reasons"]
               for g in m["gaps"])
    est = {e["label"] for e in m["established"]}
    assert "forecast horizon: one month" in est and "uses satellite: yes" in est
    assert all(e["strength"]["reasons"] for e in m["established"])

    with get_conn() as pg:
        notes = {r["agent"]: r["content"] for r in pg.execute(
            "SELECT agent, content FROM run_notes WHERE run_id=%s", (res["run_id"],)).fetchall()}
        n_proto = pg.execute("SELECT count(*) n FROM protocol_extractions WHERE run_id=%s",
                             (res["run_id"],)).fetchone()["n"]
    assert n_proto == m["N"] > 0
    # the invented probabilistic quote on non-malaria papers was discarded, so nobody counts as probabilistic
    assert notes["protocol_extraction"]["fields_dropped_without_evidence"].get("q_probabilistic", 0) > 0

    hyps = notes["hypotheses"]["items"]
    verdicts = {h["id"]: h["verdict"] for h in hyps}
    assert verdicts["H1"] == "supported" and verdicts["H2"] == "untestable"
    # the hypothesis whose text contradicted its count was refused, and repeating H1's test reused H1
    assert len(hyps) == 2
    gr = notes["gap_reasoning"]["checked"]["gaps"][0]
    assert [h["id"] for h in gr["hypotheses"]] == ["H1", "H2"]
    assert [ln["hypothesis_id"] for ln in gr["linked_gaps"]] == ["H1"]
    assert all(nm["paper_id"] != "0000.00000" for nm in gr["near_misses"])
    assert "supported" in md and "untestable" in md

    designs = notes["design"]["designs"]
    assert [d["title"] for d in designs] == ["District-holdout 3-6 month forecasts"]
    assert designs[0]["addresses"] == [m["gaps"][0]["id"]] and not designs[0]["challenging"]
    # a hypothesis listed under 'addresses' is kept as a reason, with its code-decided verdict
    assert [(h["id"], h["verdict"]) for h in designs[0]["rests_on"]] == [("H1", "supported")]
    assert "**Rests on:** H1 (supported)" in md
    assert designs[0]["supporting"][0].get("title")
    dropped = notes["design"]["dropped"]
    assert any("G99" in x for x in dropped) and any("9999.99999" in x for x in dropped)
    assert any("H77" in x for x in dropped) and not any("reference H1" in x for x in dropped)
    # invented ids are never cited; they only appear in the list of what the checks removed
    assert "[arXiv:9999.99999]" not in md and "[arXiv:0000.00000]" not in md
    assert "**Removed during checks:**" in md and "9999.99999" in md.split("**Removed during checks:**")[1]
    assert {"protocol", "gap_reasoning", "design"} <= {a for a, _ in events}


# ---------------------------------------------------------------- query parsing and map statistics
def test_query_parser_structure():
    from research_agent.tools.query import tsquery_sql

    sql, params = tsquery_sql('malaria (forecast OR prediction) -vaccine "early warning" climat*')
    assert sql.count("||") == 1 and "!!" in sql and "phraseto_tsquery" in sql and "'climat:*'" in str(params)
    assert list(params.values())[:3] == ["malaria", "forecast", "prediction"]
    # malformed input never raises and never widens the query
    for q in ("malaria (forecast", "malaria forecast)", ")(", "NOT", '"', "", "-", "a OR", "(( OR ))"):
        sql, params = tsquery_sql(q)
        assert isinstance(sql, str) and all(isinstance(v, str) for v in params.values())
    assert tsquery_sql("")[1] == {} and tsquery_sql("OR AND")[1] == {}


def test_query_parser_matches_postgres(loaded_db):
    from research_agent.db import get_conn
    from research_agent.tools.query import tsquery_sql

    docs = {"a": "malaria forecasting with rainfall", "b": "heart failure prediction", "c": "malaria vaccine trial",
            "d": "early warning system for dengue", "e": "climate-driven malaria prediction"}
    cases = {"malaria (forecast OR prediction)": {"a", "e"},
             "malaria -vaccine": {"a", "e"},
             '"early warning"': {"d"},
             "malaria forecast) OR heart": {"a", "b"},     # a stray ')' cannot split the query
             "(malaria OR dengue) AND NOT vaccine": {"a", "d", "e"},
             "climate-driven": {"e"},
             "predict*": {"b", "e"},
             "": set()}
    with get_conn() as pg:
        for q, want in cases.items():
            sql, params = tsquery_sql(q)
            got = {k for k, text in docs.items() if pg.execute(
                f"SELECT to_tsvector('english', %(doc)s) @@ ({sql}) AS m", {**params, "doc": text}).fetchone()["m"]}
            assert got == want, (q, got)


def test_corpus_count_respects_parentheses(ctx):
    from research_agent.tools.search import corpus_count

    malaria = corpus_count(ctx, "malaria")["matching_papers"]
    grouped = corpus_count(ctx, "malaria (forecast OR prediction)")["matching_papers"]
    prediction = corpus_count(ctx, "prediction")["matching_papers"]
    assert 0 < grouped <= malaria and prediction > 0
    # before the fix this read as 'malaria & forecast | prediction' and counted every sepsis 'prediction' paper
    assert corpus_count(ctx, "malaria forecast OR prediction")["matching_papers"] > grouped


def test_statistics_helpers():
    from research_agent.opportunity.compute import hyper_ge, hyper_le, wilson

    lo, hi = wilson(0, 40)
    assert lo == 0 and 0.08 < hi < 0.09            # zero of 40 still allows a true rate up to ~9%
    lo, hi = wilson(20, 40)
    assert lo < 0.5 < hi and abs((lo + hi) / 2 - 0.5) < 1e-9
    assert wilson(0, 0) == (0.0, 1.0)
    # hypergeometric tails are proper probabilities and complementary
    assert abs(hyper_le(3, 50, 20, 10) + hyper_ge(4, 50, 20, 10) - 1) < 1e-9
    assert hyper_le(0, 60, 30, 30) < 1e-15 and hyper_ge(0, 60, 30, 30) == 1


def test_clip_and_protocol_groups():
    from research_agent.opportunity.protocol import _clip, validate_protocol

    assert _clip("one two three four", 11) == "one two …" and _clip("short", 10) == "short"
    clean, problems = validate_protocol({"fields": [
        {"name": "split", "type": "enum", "role": "evaluation",
         "values": ["random", "spatial_holdout", "spatiotemporal_holdout", "temporal"],
         "definition": "x", "desirable": ["spatial_holdout", "spatiotemporal_holdout"],
         "groups": [{"label": "unseen places", "values": ["spatial_holdout", "spatiotemporal_holdout"]},
                    {"label": "not desirable", "values": ["random", "temporal"]},
                    {"label": "reuses a member", "values": ["spatial_holdout", "random"]}]},
        {"name": "horizon", "type": "enum", "values": ["a", "b"], "definition": "y", "role": "made_up"}]})
    f = {x["name"]: x for x in clean["fields"]}
    assert f["q_split"]["groups"] == [{"label": "unseen places",
                                       "values": ["spatial_holdout", "spatiotemporal_holdout"]}]
    assert f["q_split"]["role"] == "evaluation" and f["q_horizon"]["role"] != "made_up"
    assert sum("group" in p for p in problems) == 2


def _synthetic_rows():
    rows = []
    for i in range(60):
        data = {"methods": ["random forest" if i < 30 else "pearson correlation"],
                "data_modalities": (["satellite imagery"] if i % 2 else ["remote sensing"]) if i >= 30
                else ["case counts"],
                # deliberately NOT aligned with the method/data split: a corpus whose every field agrees
                # with every other is two separate literatures, and then nothing in it is an untried
                # combination (see compute.stratified_expected)
                "q_model_class": "machine_learning" if i % 2 else "statistical",
                "q_interventions": "none" if i % 3 == 0 else "itn_only"}
        if i >= 50:
            data["data_modalities"].append("mobile phone data")
        if i in (2, 40, 50):          # a weak rise: 1 of 35 early papers, 2 of 25 recent ones
            data["methods"].append("support vector machine")
        rows.append({"paper_id": f"p{i}", "source": "fulltext" if i % 3 == 0 else "abstract",
                     "corpus": "pmc" if i % 2 else "arxiv", "year": 2010 + i // 5, "title": "t", "data": data})
    return rows


def test_map_novelty_is_cross_dimension_and_tested(ctx, monkeypatch):
    from research_agent.opportunity import compute as C

    monkeypatch.setattr(C, "_rows", lambda _ctx: _synthetic_rows())
    monkeypatch.setattr(C, "known_fields", lambda _ctx: (["methods", "data_modalities"], {
        "q_model_class": ["machine_learning", "statistical"], "q_interventions": ["none", "itn_only"]}))
    monkeypatch.setattr(C, "_corpus_check", lambda *a: None)
    protocol = {"fields": [{"name": "q_model_class", "type": "enum", "role": "method",
                            "values": ["machine_learning", "statistical"], "desirable": [], "definition": "x"},
                           {"name": "q_interventions", "type": "enum", "role": "data",
                            "values": ["none", "itn_only"], "desirable": [], "definition": "y"},
                           {"name": "q_code_shared", "type": "enum", "role": "reporting",
                            "values": ["yes", "no"], "desirable": ["yes"], "definition": "z"}]}
    m = C.compute_map(ctx, protocol)
    labels = {e["label"] for sec in ("established", "emerging", "gaps") for e in m[sec]}
    assert not any("pearson" in x for x in labels)          # generic statistics are not research components
    pairs = [(n["a"], n["b"]) for n in m["novelty"]]
    assert pairs, m["novelty"]
    for a, b in pairs:
        assert a["field"] != b["field"] and a["role"] != b["role"]
    # random forest (method) never meets satellite imagery (data): a cross-dimension, significant pair
    rf_sat = next(n for n in m["novelty"] if {n["a"]["label"], n["b"]["label"]} ==
                  {"methods: random forest", "data modalities: satellite / remote sensing"})
    assert rf_sat["observed"] == 0 and rf_sat["p"] < 0.001
    # model class and methods are the same dimension, so they are never paired however disjoint
    assert not any({a["role"], b["role"]} == {"method"} for a, b in pairs)
    assert all(n["p"] <= C.THRESHOLDS["novelty_max_p"] for n in m["novelty"])
    for e in m["established"]:
        assert e["ci95"][0] <= e["share"] <= e["ci95"][1]
    rise = next(e for e in m["emerging"] if e["label"] == "data modalities: mobile phone data")
    assert rise["trend"]["p"] < 0.05 and any("p" in r for r in rise["strength"]["reasons"])
    # synonyms are one item: satellite imagery and remote sensing count together, 30 papers
    sat = next(e for e in m["established"] if e["label"] == "data modalities: satellite / remote sensing")
    assert sat["n"] == 30 and set(sat["members"]) == {"satellite imagery", "remote sensing"}
    assert not any(e["label"] in ("data modalities: satellite imagery", "data modalities: remote sensing")
                   for sec in ("established", "emerging", "watch") for e in m[sec])
    # an absence value ('none') is never half of a novel combination
    assert not any("interventions: none" in (n["a"]["label"], n["b"]["label"]) for n in m["novelty"])
    # the protocol's own code-sharing field replaces the general one: one code gap, not two
    labels = {g["label"] for g in m["gaps"]}
    assert "code shared: yes" in labels and "code or data shared" not in labels
    for e in m["emerging"]:
        assert e["trend"]["p"] <= C.THRESHOLDS["emerging_max_p"]
        assert e["strength"]["grade"] != "high" or e["trend"]["p"] < 0.05
    assert all(w["trend"]["p"] > C.THRESHOLDS["emerging_max_p"] for w in m["watch"])
    assert "methods: support vector machine" in {w["label"] for w in m["watch"]}
    assert "methods: support vector machine" not in {e["label"] for e in m["emerging"]}


def test_hypotheses_dedupe_and_flag_restatements(ctx, monkeypatch):
    from research_agent.opportunity import agents as A
    from research_agent.tools import claims as Cl

    rows = _synthetic_rows()
    monkeypatch.setattr(Cl, "_rows", lambda _ctx: rows)
    gap = {"id": "G1", "field": "data_modalities", "value": "mobile phone data|call records", "label": "x"}
    monkeypatch.setattr(A, "_map", lambda _ctx: {"gaps": [gap], "novelty": [], "emerging": [], "established": []})
    pred = {"field": "methods", "any_of": ["random forest"], "min_share": 0.4}
    first = A.test_hypothesis(ctx, "G1", "Random forest appears in 30 of 60 papers.", "prevalence", pred)
    again = A.test_hypothesis(ctx, "G1", "Random forests are used in 30 of 60 papers.", "prevalence", dict(pred))
    assert first["hypothesis_id"] == again["hypothesis_id"] == "H1" and "already" in again["note"]
    same = A.test_hypothesis(ctx, "G1", "Mobile phone data appears in 10 of 60 papers.", "prevalence",
                             {"field": "data_modalities", "any_of": ["mobile phone data"], "max_share": 0.2})
    assert same["hypothesis_id"] == "H2" and "re-measures" in same["note"]
    items = ctx.notes()["hypotheses"]["items"]
    assert [h["restates_item"] for h in items] == [False, True]


def test_topic_scoped_trends_and_absence_checks(ctx, monkeypatch):
    from research_agent.tools import claims as C
    from research_agent.tools.trends import topic_series

    scoped = topic_series(ctx, "rainfall", within="malaria")
    plain = topic_series(ctx, "malaria rainfall")
    assert "within" in scoped and all("pct_of_topic" in y and "per_10k_arxiv" not in y for y in scoped["series"])
    assert sum(y["matching"] for y in scoped["series"]) == sum(y["matching"] for y in plain["series"]) > 0
    assert all(y["matching"] <= y["topic_total"] for y in scoped["series"])
    # a claim about "malaria papers" must be measured within malaria papers
    pred = {"keywords": "rainfall", "early": [2015, 2017], "late": [2021, 2023], "direction": "increase"}
    refused = C.propose_claim(ctx, "Rainfall covariates rose in malaria papers.", "trend", dict(pred))
    assert "within" in refused["error"]
    assert "error" not in C.test_claim(ctx, "trend", {**pred, "within": "malaria"})

    # an absence claim on a free-text field is checked against the papers' own abstracts
    ids = [r["paper_id"] for r in ctx.pg.execute(
        "SELECT paper_id FROM papers WHERE tsv @@ plainto_tsquery('english', 'malaria rainfall') LIMIT 12").fetchall()]
    assert len(ids) >= 5
    rows = [{"paper_id": pid, "source": "abstract", "corpus": "arxiv", "year": 2020,
             "data": {"data_modalities": ["surveillance counts"]}} for pid in ids]
    monkeypatch.setattr(C, "_rows", lambda _ctx: rows)
    absent = C.evaluate_prevalence(ctx, {"field": "data_modalities", "any_of": ["rainfall"], "max_count": 0})
    assert absent["n_matching"] == 0 and not absent["supported"] and "not established" in absent["absence_problem"]
    # with re-checks off, proposing the absence is refused outright
    import dataclasses
    monkeypatch.setattr(C, "settings", dataclasses.replace(C.settings, recheck=False))
    refused = C.propose_claim(ctx, "No paper uses rainfall data (0 of %d)." % len(ids), "prevalence",
                              {"field": "data_modalities", "any_of": ["rainfall"], "max_count": 0})
    assert "not established" in refused["error"]
    # a term the abstracts do not mention stays a valid rarity claim
    fine = C.evaluate_prevalence(ctx, {"field": "data_modalities", "any_of": ["mobile phone"], "max_count": 0})
    assert fine["supported"] and "absence_problem" not in fine


def test_number_check_backs_corrects_and_drops(ctx, monkeypatch):
    from research_agent.agents import number_check as NC
    from research_agent.agents.report import allowed_counts
    from research_agent.tools import claims as C

    rows = [{"paper_id": f"p{i}", "source": "fulltext" if i < 10 else "abstract", "corpus": "pmc", "year": 2020,
             "title": "t", "data": {"methods": ["random forest"] if i < 7 else ["ARIMA"],
                                    "code_or_data_available": "yes" if i < 3 else "no"}} for i in range(40)]
    monkeypatch.setattr(C, "_rows", lambda _ctx: rows)
    body = ("## Summary\n"
            "- Random forests appear in 7 of 40 papers (18%) [PMC1].\n"
            "- Code is shared by 5 of 40 papers (13%).\n"
            "- Method families cover 34 of 40 papers across several groups.\n"
            "- Tree models appear in 9 of 40 papers.\n"
            "| C1 | row 3 of 40 | supported | x |\n")
    items = NC.unverified_items(body, allowed_counts(ctx))
    assert [it["number"] for it in items] == ["7 of 40", "5 of 40", "34 of 40", "9 of 40"]   # tables skipped
    assert items[0]["sentence"].startswith("Random forests")                                  # no bullet
    rf = C.propose_claim(ctx, "Random forests appear in 7 of 40 papers.", "prevalence",
                         {"field": "methods", "any_of": ["random forest"], "min_count": 5})
    code = C.propose_claim(ctx, "Code is shared by 3 of 40 papers.", "prevalence",
                           {"field": "code_or_data_available", "any_of": ["yes"], "max_share": 0.1})
    ft = C.propose_claim(ctx, "Random forests appear in 7 of 10 full-text papers.", "prevalence",
                         {"field": "methods", "any_of": ["random forest"], "min_count": 5,
                          "where": [{"field": "read", "any_of": ["fulltext"]}]})
    fixes = [{"item": "U1", "action": "measure", "claim_id": rf["claim_id"]},
             {"item": "U2", "action": "measure", "claim_id": code["claim_id"]},
             {"item": "U3", "action": "drop", "sentence": "Method families cover most papers across several groups."},
             {"item": "U4", "action": "measure", "claim_id": ft["claim_id"]}]
    out, stats = NC.apply_fixes(ctx, body, items, fixes)
    assert f"7 of 40 papers (18%) [PMC1] [C{rf['claim_id']}]." in out    # backed: same number, now cited
    assert f"Code is shared by 3 of 40 papers (8%) [C{code['claim_id']}]." in out   # corrected, % follows
    assert "- Method families cover 34 of 40" in out                    # 'most' is an uncounted quantity: refused
    assert "9 of 40" in out                                              # claim counted out of 10: refused
    assert (stats["backed"], stats["corrected"], stats["removed"], stats["refused"]) == (1, 1, 0, 2)
    # a faithful drop is accepted, and the bullet stays
    out2, stats2 = NC.apply_fixes(ctx, body, items, [
        {"item": "U3", "action": "drop", "sentence": "Method families span several groups (not counted per paper)."}])
    assert "- Method families span several groups (not counted per paper)." in out2 and stats2["removed"] == 1
    # dropping may not lose a citation or add numbers
    assert NC._drop_ok("X in 7 of 40 papers [PMC1].", "X in some papers.", "7 of 40")
    assert NC._drop_ok("X in 7 of 40 papers.", "X in 12 papers.", "7 of 40")


def test_recheck_confirms_uses_the_extraction_missed(ctx, monkeypatch):
    from research_agent.tools import claims as C
    from research_agent.tools import extraction as X
    from research_agent.tools import recheck as R

    ids = [r["paper_id"] for r in ctx.pg.execute(
        "SELECT paper_id FROM papers WHERE tsv @@ plainto_tsquery('english', 'malaria rainfall') LIMIT 8").fetchall()]
    base = [{"paper_id": pid, "source": "abstract", "corpus": "arxiv", "year": 2020, "title": "t",
             "data": {"data_modalities": ["surveillance counts"]}} for pid in ids]

    def rows(_ctx):   # the real merge of confirmed re-checks, on these rows
        out = [{**r, "data": dict(r["data"])} for r in base]
        X._merge_rechecks(_ctx, out)
        return out
    monkeypatch.setattr(C, "_rows", rows)
    before = C.evaluate_prevalence(ctx, {"field": "data_modalities", "any_of": ["rainfall"], "max_count": 0})
    assert before["n_matching"] == 0 and before.get("absence_problem")
    # proposing "no paper uses rainfall" triggers the re-check, which finds every paper does
    res = C.propose_claim(ctx, "No paper uses rainfall data (0 of %d)." % len(ids), "prevalence",
                          {"field": "data_modalities", "any_of": ["rainfall"], "max_count": 0})
    assert f"{len(ids)} of {len(ids)}" in (res.get("measured") or res.get("measured_now"))
    if "claim_id" in res:     # accepted as pending, it fails verification: its bound (none) is not met
        C.verify_claims(ctx, [res["claim_id"]])
        assert ctx.pg.execute("SELECT status FROM claims WHERE id=%s", (res["claim_id"],)).fetchone()["status"] \
            == "unsupported"
    stored = ctx.pg.execute("SELECT verdict, quote FROM rechecks WHERE run_id=%s", (ctx.run_id,)).fetchall()
    assert len(stored) == len(ids) and all(s["verdict"] == "yes" and "rainfall" in s["quote"] for s in stored)
    after = C.evaluate_prevalence(ctx, {"field": "data_modalities", "any_of": ["rainfall"], "min_count": 1})
    assert after["n_matching"] == len(ids) and after["supported"] and "absence_problem" not in after
    assert all("rainfall" in r["data"]["_rechecked"]["data_modalities"] for r in rows(ctx))
    # a second call re-checks nobody: answers are cached per paper and concept
    assert R.recheck(ctx, "data_modalities", ["rainfall"])["checked_now"] == 0
    assert "list fields" in R.recheck(ctx, "validation_level", ["external"])["error"]


def test_recheck_answers_need_a_real_quote():
    from research_agent.llm.base import LLMResponse, ToolCall
    from research_agent.tools.recheck import _ask

    text = ("Title: Forecasting malaria\n\nAbstract: Earlier studies used transformer models. We fit an "
            "ARIMA model to monthly case counts.")

    class Stub:
        def __init__(self, args):
            self.args = args

        def chat(self, *a, **k):
            tc = ToolCall("1", "record_check", self.args)
            return LLMResponse("", [tc], [], "tool_use", {})

    terms = ["transformer"]
    invented = _ask(Stub({"uses": "yes", "which": "transformer", "quote": "we train a transformer network"}),
                    "transformers", terms, {}, text)
    assert invented["verdict"] == "unclear" and not invented["quote"]      # quote is not in the paper
    real = _ask(Stub({"uses": "yes", "which": "transformer",
                      "quote": "Earlier studies used transformer models."}), "transformers", terms, {}, text)
    assert real["verdict"] == "yes"
    assert _ask(Stub({"uses": "maybe"}), "transformers", terms, {}, text)["verdict"] == "unclear"


def test_coverage_probe_finds_what_discovery_missed(ctx):
    from research_agent.tools.search import add_to_shortlist, coverage_probe

    malaria = [r["paper_id"] for r in ctx.pg.execute(
        "SELECT paper_id FROM papers WHERE tsv @@ plainto_tsquery('english', 'malaria') ORDER BY paper_id").fetchall()]
    add_to_shortlist(ctx, malaria[:1], "seed")
    res = coverage_probe(ctx, "malaria", ["rainfall", "(sepsis)"])
    rain, sepsis = res["probes"]
    assert res["topic_papers_in_corpus"] >= rain["corpus_papers"] > 1
    assert rain["on_shortlist"] <= 1 and rain["under_covered"] == (rain["corpus_papers"] >= 5)
    assert all(p["paper_id"] not in malaria[:1] for p in rain["not_shortlisted_examples"])
    assert sepsis["corpus_papers"] == 0 and not sepsis["under_covered"]
    # the corpus counts are recorded, so a report may quote them
    assert ctx.pg.execute("SELECT count(*) n FROM run_events WHERE run_id=%s AND agent='coverage_probe' "
                          "AND kind='count'", (ctx.run_id,)).fetchone()["n"] == 2


def test_trimming_is_done_once_and_kept():
    from research_agent.agents.base import _msg_chars, compact

    msgs = [{"role": "user", "content": [{"type": "text", "text": "q"}]}]
    for i in range(12):
        msgs.append({"role": "assistant", "content": [{"type": "text", "text": "call"}]})
        msgs.append({"role": "user", "content": [{"type": "tool_result", "tool_use_id": str(i), "content": "x" * 3000}]})
    budget = 20000
    trimmed = compact(msgs, int(budget * 0.6))
    assert _msg_chars(trimmed) <= budget * 0.6
    # adding one more turn keeps it under budget, so the trimmed start (the cached prefix) is not touched again
    trimmed.append({"role": "assistant", "content": [{"type": "text", "text": "next"}]})
    assert _msg_chars(trimmed) <= budget and compact(trimmed, budget) is trimmed


def test_usage_is_broken_down_by_step(loaded_db, monkeypatch):
    import dataclasses

    from research_agent import runstate
    from research_agent.agents.orchestrator import run_research
    from research_agent.cli import format_usage

    monkeypatch.setattr(runstate, "settings", dataclasses.replace(
        runstate.settings, price_input_per_m=1.0, price_cached_input_per_m=0.1, price_output_per_m=2.0))
    res = run_research("What ML methods are used for malaria forecasting?", mode="pipeline",
                       llm_factory=lambda strong=False: FakeLLM())
    steps = res["usage_by_step"]
    assert {"discovery", "extraction", "synthesis", "protocol_fields"} <= set(steps)
    assert sum(s["llm_calls"] for s in steps.values()) == res["usage"]["llm_calls"]
    assert abs(sum(s["cost_usd"] for s in steps.values()) - res["usage"]["cost_usd"]) < 0.01
    assert "cost $" in format_usage(steps, res["usage"])


def _reading_stub(script):
    from research_agent.llm.base import LLMResponse, ToolCall

    class Stub:
        model = "stub"
        calls = []

        def chat(self, system, messages, tools=None, force_tool=None, max_tokens=0, temperature=0.0):
            Stub.calls.append(force_tool)
            args = script(force_tool, messages[0]["content"][0]["text"], tools)
            return LLMResponse("", [ToolCall(str(len(Stub.calls)), force_tool, args)], [], "tool_use", {})
    return Stub


def test_one_read_per_paper_and_batched_abstracts(ctx, monkeypatch):
    import dataclasses

    from research_agent.tools import reading as RD
    from tests.fake_llm import fake_reading

    protocol = {"fields": [{"name": "q_probabilistic", "type": "enum", "values": ["yes", "no", "not_stated"],
                            "definition": "Gives intervals.", "desirable": ["yes"]}]}
    papers = [{"paper_id": f"f{i}", "title": f"Paper f{i}", "abstract": "We forecast monthly malaria cases using "
               "an LSTM driven by rainfall.", "fulltext": "## Methods\nWe used surveillance counts from Kenya."}
              for i in range(3)]
    papers += [{"paper_id": f"a{i}", "title": f"Paper a{i}", "abstract": "We forecast monthly malaria cases "
                "using ARIMA in Uganda.", "fulltext": None} for i in range(7)]

    def script(tool, text, tools):
        if tool == "record_readings":
            blocks = re.split(r"^=== Paper (\S+) ===\n", text, flags=re.M)[1:]
            return {"records": [{"paper_id": pid, **fake_reading(body, tools)} for pid, body in zip(blocks[::2], blocks[1::2])]}
        return fake_reading(text, tools)
    Stub = _reading_stub(script)
    monkeypatch.setattr(RD, "settings", dataclasses.replace(RD.settings, abstract_batch=5))
    ctx.llm_factory = lambda strong=False: Stub()
    results, failed, _ = RD.read_papers(ctx, papers, protocol)
    # 3 full-text papers alone + 7 abstracts in batches of 5 and 2 = 5 calls, each covering BOTH sets of fields
    assert sorted(Stub.calls) == ["record_reading"] * 3 + ["record_readings"] * 2 and not failed
    assert set(results) == {p["paper_id"] for p in papers}
    assert all(r["base"] is not None and r["protocol"] is not None for r in results.values())
    assert results["f0"]["source"] == "fulltext" and results["a0"]["source"] == "abstract"
    assert "Kenya" in results["f0"]["base"]["geography"]            # read from the full-text section
    assert "Uganda" in results["a3"]["base"]["geography"]


def test_batched_reads_cannot_borrow_quotes_and_recover_missing_papers(ctx):
    import re as _re

    from research_agent.tools import reading as RD

    papers = [{"paper_id": "a1", "title": "One", "abstract": "We used ARIMA models on clinic data from Ghana.",
               "fulltext": None},
              {"paper_id": "a2", "title": "Two", "abstract": "A survey of bed net use in Malawi households.",
               "fulltext": None},
              {"paper_id": "a3", "title": "Three", "abstract": "Rainfall predicted cases in Zambia districts.",
               "fulltext": None}]

    def script(tool, text, tools):
        if tool == "record_readings":
            # a2's record borrows a1's sentence as evidence; a3 is left out altogether
            return {"records": [
                {"paper_id": "a1", "methods": ["ARIMA"], "geography": ["Ghana"],
                 "evidence": {"methods": ["We used ARIMA models on clinic data"], "geography": ["clinic data from Ghana"]}},
                {"paper_id": "a2", "methods": ["ARIMA"], "geography": ["Malawi"],
                 "evidence": {"methods": ["We used ARIMA models on clinic data"],
                              "geography": ["bed net use in Malawi households"]}}]}
        assert "Zambia" in text and "Ghana" not in text                  # the fallback reads a3 alone
        return {"geography": ["Zambia"], "evidence": {"geography": ["Rainfall predicted cases in Zambia districts"]}}
    Stub = _reading_stub(script)
    ctx.llm_factory = lambda strong=False: Stub()
    results, failed, _ = RD.read_papers(ctx, papers, None)
    assert results["a1"]["base"]["methods"] == ["ARIMA"]
    assert results["a2"]["base"]["methods"] == [] and results["a2"]["base"]["geography"] == ["Malawi"]
    assert results["a3"]["base"]["geography"] == ["Zambia"] and Stub.calls.count("record_reading") == 1
    _ = _re


def test_cut_off_batch_keeps_finished_records(ctx):
    import json as _json

    from research_agent.tools import reading as RD

    papers = [{"paper_id": f"b{i}", "title": f"T{i}", "abstract": f"Study {i} used ARIMA in Ghana.", "fulltext": None}
              for i in range(3)]
    done = {"paper_id": "b0", "methods": ["ARIMA"], "evidence": {"methods": ["Study 0 used ARIMA in Ghana"]}}

    def script(tool, text, tools):
        if tool == "record_readings":
            raw = _json.dumps({"records": [done]})[:-2] + ', {"paper_id": "b1", "meth'
            return {"_raw_arguments": raw}
        return {"methods": [], "evidence": {}}
    Stub = _reading_stub(script)
    ctx.llm_factory = lambda strong=False: Stub()
    results, failed, _ = RD.read_papers(ctx, papers, None)
    assert results["b0"]["base"]["methods"] == ["ARIMA"]                 # kept from the cut-off batch
    assert Stub.calls.count("record_reading") == 2 and set(results) == {"b0", "b1", "b2"}


def test_followup_answers_are_checked_and_leave_the_report_alone(loaded_db):
    from research_agent.agents import followup
    from research_agent.agents.orchestrator import run_research
    from research_agent.db import get_conn

    res = run_research("What ML methods are used for malaria forecasting?", mode="pipeline",
                       llm_factory=lambda strong=False: FakeLLM())
    rid = res["run_id"]
    with get_conn() as pg:
        before = pg.execute("SELECT status, report_md, finished_at FROM runs WHERE run_id=%s", (rid,)).fetchone()
        cid = pg.execute("SELECT id FROM claims WHERE run_id=%s AND status='supported' ORDER BY id LIMIT 1",
                         (rid,)).fetchone()["id"]
    out = followup.answer(rid, f"Which papers are behind C{cid}?", f"C{cid}", llm_factory=lambda strong=False: FakeLLM())
    ans = out["answer"]
    new = out["new_claims"][0]
    assert f"[{new}]" in ans and "0000.00000" not in ans and "citation removed" in ans   # invented paper removed
    assert "7 of 9 [unverified]" in ans                                                   # untraced number tagged
    with get_conn() as pg:
        after = pg.execute("SELECT status, report_md, finished_at FROM runs WHERE run_id=%s", (rid,)).fetchone()
        added = pg.execute("SELECT agent, status FROM claims WHERE id=%s", (int(new[1:]),)).fetchone()
    assert after == before                                  # the run and its report are untouched
    assert added["agent"] == "followup" and added["status"] == "supported"
    # the conversation is stored, and the next question sees it
    second = followup.answer(rid, "And overall?", None, llm_factory=lambda strong=False: FakeLLM())
    assert second["answer"].startswith("As discussed before")
    msgs = followup.history(rid)
    assert [m["role"] for m in msgs] == ["user", "assistant", "user", "assistant"]
    assert all(m["status"] == "done" for m in msgs)


def test_followup_item_lookup_and_paper_reading(ctx, monkeypatch):
    from research_agent.agents import followup
    from research_agent.tools import claims as C

    assert "error" in followup.get_item(ctx, "C999999")
    assert "error" in followup.get_item(ctx, "G7")
    assert "not one of the papers analysed" in followup.read_paper(ctx, "0000.00000")["error"]


def test_followup_api_and_cli(loaded_db, monkeypatch, capsys):
    import time

    from fastapi.testclient import TestClient

    from research_agent import cli
    from research_agent.agents import followup
    from research_agent.agents.orchestrator import run_research
    from research_agent.api import main as api

    from research_agent import jobs
    monkeypatch.setattr(jobs, "_llm_factory", lambda strong=False: FakeLLM())
    real = followup.answer
    monkeypatch.setattr(followup, "answer", lambda *a, **k: real(*a, **{**k, "llm_factory": lambda strong=False: FakeLLM()}))
    rid = run_research("What ML methods are used for malaria forecasting?", mode="pipeline",
                       llm_factory=lambda strong=False: FakeLLM())["run_id"]
    client = TestClient(api.app)
    assert client.post(f"/runs/{rid}/followups", json={"question": "Why?", "about": "C1; DROP"}).status_code == 422
    r = client.post(f"/runs/{rid}/followups", json={"question": "What is behind the main claim?"})
    assert r.status_code == 202 and r.json()["job_id"]
    assert jobs.Worker(name="test").run_once()
    for _ in range(60):
        msgs = client.get(f"/runs/{rid}/followups").json()["messages"]
        if msgs and msgs[-1]["status"] != "pending":
            break
        time.sleep(0.3)
    assert msgs[-1]["role"] == "assistant" and msgs[-1]["status"] == "done" and msgs[-1]["new_claims"]
    assert client.post("/runs/00000000-0000-0000-0000-000000000000/followups",
                       json={"question": "anything?"}).status_code == 404
    # terminal: one question, then the answer is printed
    assert cli.chat(rid, "Does it hold for PMC papers only?") == 0
    assert "Surveillance counts appear in" in capsys.readouterr().out


def test_followups_work_on_a_database_from_before_followups(loaded_db):
    from fastapi.testclient import TestClient

    from research_agent.agents import followup
    from research_agent.api import main as api
    from research_agent.db import get_conn

    with get_conn() as pg:
        pg.execute("DROP TABLE IF EXISTS followups")
        rid = str(pg.execute("SELECT run_id FROM runs WHERE status='done' LIMIT 1").fetchone()["run_id"])
    followup._schema_ready = False
    client = TestClient(api.app, raise_server_exceptions=False)
    assert client.get(f"/runs/{rid}/followups").status_code == 200       # the table is created on demand
    with get_conn() as pg:
        pg.execute("DROP TABLE IF EXISTS followups")
    followup._schema_ready = False
    with TestClient(api.app) as started:                                  # and when the server starts
        with get_conn() as pg:
            assert pg.execute("SELECT to_regclass('followups') AS t").fetchone()["t"] is not None
        assert started.get(f"/runs/{rid}/followups").status_code == 200


def test_followup_reads_abstract_only_papers_in_full_and_recounts(loaded_db, monkeypatch):
    from research_agent.agents import followup
    from research_agent.agents.orchestrator import run_research
    from research_agent.db import get_conn
    from research_agent.runstate import RunContext
    from research_agent.tools import resolve as RS

    # this is about the follow-up tool itself, so leave the papers as the run read them: the pre-report
    # resolution pass would otherwise have already read the abstract-only ones it is about to be asked for
    monkeypatch.setattr(RS, "resolve", lambda *a, **k: {"attempted": 0})
    rid = run_research("What ML methods are used for malaria forecasting?", mode="pipeline",
                       llm_factory=lambda strong=False: FakeLLM())["run_id"]
    ctx = RunContext.attach(rid, llm_factory=lambda strong=False: FakeLLM())
    try:
        c = ctx.pg.execute("""SELECT id, result FROM claims WHERE run_id=%s AND claim_type='prevalence'
                              AND predicate->'where' IS NULL ORDER BY id LIMIT 1""", (rid,)).fetchone()
        # Earlier runs in this database may already have read every retrievable paper in full, so put the
        # ones whose full text IS available back to abstract-only. read_in_full reads them again, which
        # leaves the shared extraction records exactly as they were.
        with get_conn() as pg:
            pg.execute("""UPDATE extractions SET source='abstract' WHERE schema_version=%s AND paper_id IN
                          (SELECT rp.paper_id FROM run_papers rp JOIN paper_fulltext f USING (paper_id)
                           WHERE rp.run_id=%s AND f.status='ok')""",
                       (settings.extraction_schema_version, rid))
        abstract_before = ctx.pg.execute(
            """SELECT count(*) n FROM run_papers rp JOIN extractions e USING (paper_id)
               WHERE rp.run_id=%s AND e.source='abstract'""", (rid,)).fetchone()["n"]
        out = followup.read_in_full(ctx, f"C{c['id']}", limit=20)
        # papers whose full text cannot be downloaded here (PMC needs the network) are reported, not fatal
        assert out["read_in_full_now"] + len(out["no_full_text_available"]) + len(out["failed"]) == \
            min(abstract_before, 20) and out["read_in_full_now"] > 0
        abstract_after = ctx.pg.execute(
            """SELECT count(*) n FROM run_papers rp JOIN extractions e USING (paper_id)
               WHERE rp.run_id=%s AND e.source='abstract'""", (rid,)).fetchone()["n"]
        assert abstract_after == abstract_before - out["read_in_full_now"]
        # the original claim keeps its result; the re-count is a separate follow-up claim
        orig = ctx.pg.execute("SELECT result FROM claims WHERE id=%s", (c["id"],)).fetchone()["result"]
        new = ctx.pg.execute("SELECT agent, status, result FROM claims WHERE id=%s",
                             (int(out["recount_claim"][1:]),)).fetchone()
        assert orig == c["result"] and new["agent"] == "followup" and new["status"] in ("supported", "unsupported")
        assert out["after"] == f"{new['result']['n_matching']} of {new['result']['denominator']}"
        # nothing left to read the second time
        again = followup.read_in_full(ctx, f"C{c['id']}")
        assert again.get("note") or again["read_in_full_now"] == 0
        assert "error" in followup.read_in_full(ctx, "C999999")
    finally:
        ctx.close()
    with get_conn() as pg:
        assert pg.execute("SELECT report_md IS NOT NULL AS r FROM runs WHERE run_id=%s", (rid,)).fetchone()["r"]


def test_model_routing_by_step_and_tier(monkeypatch):
    import dataclasses

    from research_agent import llm as L
    from research_agent.runstate import _routing_factory

    assert L.parse_routes("cheap=deepseek:deepseek-flash; synthesis = anthropic:claude-x;bad") == {
        "cheap": ("deepseek", "deepseek-flash"), "synthesis": ("anthropic", "claude-x")}
    monkeypatch.setattr(L, "settings", dataclasses.replace(
        L.settings, model_routes="cheap=deepseek:m-cheap; strong=deepseek:m-strong; synthesis=deepseek:m-write"))
    assert L.route_for("extraction") == ("deepseek", "m-cheap")        # tier
    assert L.route_for("gap_reasoning") == ("deepseek", "m-strong")    # tier
    assert L.route_for("synthesis") == ("deepseek", "m-write")         # the step's own route wins
    assert L.route_for("methods") is None                              # not routed: default provider
    assert L.route_for("methods", strong=True) == ("deepseek", "m-strong")
    monkeypatch.setenv("DEEPSEEK_API_KEY", "x")
    assert L.get_llm(step="extraction").model == "m-cheap" and L.get_llm(step="synthesis").model == "m-write"
    # factories that only take `strong` (tests, custom code) keep working, and clients are labelled by step
    clients = []
    f = _routing_factory(lambda strong=False: FakeLLM(), clients)
    c = f(step="recheck")
    assert c._step == "recheck" and clients == [c]


def test_reported_results_and_associations_need_their_own_quotes():
    from research_agent.tools.results import verify_structured

    text = ("The LSTM reached an RMSE of 12.4 cases on the test set, against 18.9 for SARIMA. "
            "Rainfall at a two-month lag was positively associated with incidence.")
    kept, dropped = verify_structured({
        "reported_results": [
            {"metric": "RMSE", "value": "12.4", "model": "LSTM", "split": "test_or_holdout",
             "quote": "The LSTM reached an RMSE of 12.4 cases on the test set"},
            {"metric": "RMSE", "value": "18.9", "model": "SARIMA", "is_baseline": True,
             "quote": "The LSTM reached an RMSE of 12.4 cases on the test set"},        # number not in quote
            {"metric": "AUC", "value": "0.93", "model": "LSTM", "quote": "AUC was 0.93 overall"}],  # not in text
        "reported_associations": [
            {"driver": "rainfall", "direction": "positive", "lag": "2 months",
             "quote": "Rainfall at a two-month lag was positively associated with incidence"},
            {"driver": "temperature", "direction": "negative",
             "quote": "Rainfall at a two-month lag was positively associated with incidence"}]}, text)   # driver absent
    assert [(r["metric"], r["value_num"], r["higher_is_better"]) for r in kept["reported_results"]] == [("rmse", 12.4, False)]
    assert [a["driver"] for a in kept["reported_associations"]] == ["rainfall"]
    assert dropped == {"reported_results": 2, "reported_associations": 1}


def _results_rows():
    rows = []
    def res(model, v, metric="RMSE", baseline=False):
        return {"metric": metric.lower(), "value": str(v), "value_num": v, "higher_is_better": False,
                "model": model, "is_baseline": baseline, "split": "test_or_holdout", "setting": "", "horizon": "",
                "quote": f"{model} {v}"}
    def assoc(driver, direction):
        return {"driver": driver, "direction": direction, "lag": "", "significant": "yes", "outcome": "cases",
                "quote": f"{driver} {direction}"}
    for i in range(8):   # deep learning beats ARIMA in 6 of 8 papers, loses in 2
        dl_better = i < 6
        rows.append({"paper_id": f"p{i}", "source": "fulltext", "corpus": "pmc", "year": 2021, "title": "t",
                     "data": {"reported_results": [res("LSTM", 10 if dl_better else 20), res("ARIMA", 15, baseline=True)],
                              "reported_associations": [assoc("temperature", "positive" if i < 4 else "negative")],
                              "geography": ["Ethiopia"] if i < 4 else ["Kenya"], "methods": ["LSTM"]}})
    return rows


def test_fair_method_comparison_and_contradictions(ctx, monkeypatch):
    from research_agent.tools import claims as C
    from research_agent.tools import results as RS

    rows = _results_rows()
    monkeypatch.setattr(C, "_rows", lambda _ctx: rows)
    cmp = RS.method_comparison(ctx, ["lstm"], ["arima"], "rmse")
    assert (cmp["papers_with_head_to_head"], cmp["A_better"], cmp["B_better"]) == (8, 6, 2)
    assert 0.2 < cmp["sign_test_p"] < 0.4                      # 6 vs 2 is not convincing on its own
    table = RS.results_table(ctx, metric="root mean squared error")
    assert table["summary"]["rmse"]["papers"] == 8 and table["n_results"] == 16
    co = RS.contradictions(ctx)["contradictions"]
    assert [c["driver"] for c in co] == ["temperature"]
    assert len(co[0]["positive"]) == 4 and len(co[0]["negative"]) == 4
    top = co[0]["separating_attributes"][0]
    assert top["attribute"] in ("place: Ethiopia", "place: Kenya") and abs(top["share_positive_side"] - top["share_negative_side"]) == 1
    md = "\n".join(RS.results_markdown(ctx))
    assert "## Reported performance (computed)" in md and "## Where studies disagree (computed)" in md
    assert "temperature" in md and "place: Ethiopia" in md


def test_extraction_records_results_and_report_appends_them(loaded_db):
    from research_agent.agents.orchestrator import run_research
    from research_agent.db import get_conn

    res = run_research("How well do models detect pneumonia and forecast malaria?", mode="pipeline",
                       llm_factory=lambda strong=False: FakeLLM())
    with get_conn() as pg:
        rows = pg.execute("""SELECT e.data FROM run_papers rp JOIN extractions e USING (paper_id)
                             WHERE rp.run_id=%s AND e.schema_version=%s""",
                          (res["run_id"], settings.extraction_schema_version)).fetchall()
    results = [it for r in rows for it in r["data"].get("reported_results") or []]
    assocs = [it for r in rows for it in r["data"].get("reported_associations") or []]
    assert rows and (results or assocs)
    assert all(it["value"] in it["quote"] for it in results)


def test_screening_log_and_prisma_flow(ctx):
    from research_agent.tools.review import prisma_counts, prisma_markdown, protocol_markdown, screening_rows
    from research_agent.tools.search import add_to_shortlist, hybrid_search, remove_from_shortlist

    found = hybrid_search(ctx, "malaria forecasting", limit=10)["results"]
    ids = [r["paper_id"] for r in found]
    add_to_shortlist(ctx, ids[:4], "forecasts malaria incidence")
    remove_from_shortlist(ctx, ids[3:4], "no forecasting model")
    c = prisma_counts(ctx)
    assert c["identified"] == len(ids) and c["included"] == 3
    assert c["screened"] == c["identified"] - c["duplicates_removed"]
    assert c["excluded_at_screening"] == c["screened"] - c["included"]
    assert c["exclusion_reasons"] == {"no forecasting model": 1}
    stages = {r["paper_id"]: r["stage"] for r in screening_rows(ctx)}
    assert stages[ids[0]] == "included" and stages[ids[3]] == "excluded" and stages[ids[-1]] == "identified"
    md = "\n".join(prisma_markdown(ctx))
    assert "## Review record (PRISMA 2020)" in md and "```mermaid" in md and f"Included: {c['included']}" in md
    assert protocol_markdown(ctx).startswith("# Review protocol") and "## Sources" in protocol_markdown(ctx)


def _openalex_mock(paper_ids):
    """A tiny OpenAlex: work W<i> for each fixture paper; each paper cites the one before it, and paper 0 is
    cited by an outside work. Also serves reference lookups and 'cites' queries."""
    import json as _json

    import httpx

    doi = {pid: f"10.48550/arxiv.{pid.lower()}" for pid in paper_ids}
    wid = {pid: f"W{i + 1}" for i, pid in enumerate(paper_ids)}
    work = {pid: {"id": f"https://openalex.org/{wid[pid]}", "doi": f"https://doi.org/{doi[pid]}",
                  "cited_by_count": 10 * (len(paper_ids) - i),
                  "referenced_works": [f"https://openalex.org/{wid[paper_ids[i - 1]]}"] if i else []}
            for i, pid in enumerate(paper_ids)}
    calls = []

    def handler(request):
        f = request.url.params.get("filter", "")
        calls.append(f)
        if f.startswith("doi:"):
            want = set(f[4:].split("|"))
            res = [w for pid, w in work.items() if doi[pid] in want]
        elif f.startswith("openalex:"):
            want = set(f[9:].split("|"))
            res = [w for pid, w in work.items() if wid[pid] in want]
        elif f.startswith("cites:"):
            target = f[6:]
            res = [w for w in work.values() if f"https://openalex.org/{target}" in w["referenced_works"]]
        else:
            return httpx.Response(400, text="bad filter")
        return httpx.Response(200, text=_json.dumps({"results": res}))
    return httpx.MockTransport(handler), calls


def test_citation_graph_and_snowball(ctx, monkeypatch):
    from research_agent.tools import citations as CT
    from research_agent.tools.search import add_to_shortlist

    ids = [r["paper_id"] for r in ctx.pg.execute(
        "SELECT paper_id FROM papers WHERE source='arxiv' AND tsv @@ plainto_tsquery('english','malaria') "
        "ORDER BY paper_id LIMIT 6").fetchall()]
    transport, calls = _openalex_mock(ids)
    monkeypatch.setattr(CT, "_transport", transport)
    ctx.pg.execute("DELETE FROM openalex_works WHERE paper_id = ANY(%s)", (ids,))
    add_to_shortlist(ctx, ids[:4], "seed")
    g = CT.citation_graph(ctx)
    assert g["matched_in_openalex"] == 4 and g["links_between_papers"] == 3        # a chain 1<-2<-3<-4
    assert calls == [f for f in calls if f.startswith("doi:")] and len(calls) == 1   # one request for 4 papers
    top = g["most_cited_within_set"][0]
    assert top["cited_by_papers_in_set"] == 1 and top["cited_by_all_literature"] > 0
    # snowballing from paper 3 finds paper 4 (cites it) in the corpus; paper 2 is already shortlisted
    found = CT.snowball(ctx, [ids[3]], "forward")
    assert [p["paper_id"] for p in found["found_in_corpus"]] == [ids[4]]
    back = CT.snowball(ctx, [ids[4]], "backward")
    assert back["found_in_corpus"] == []                      # it cites paper 3, which is on the shortlist
    md = "\n".join(CT.citation_markdown(ctx))
    assert "## Citation structure (computed)" in md and "cite each other 3 times" in md
    # an unreachable OpenAlex leaves the citation parts out instead of failing
    import httpx
    monkeypatch.setattr(CT, "_transport", httpx.MockTransport(lambda r: httpx.Response(503)))
    monkeypatch.setattr(CT.time, "sleep", lambda s: None)
    ctx.pg.execute("DELETE FROM openalex_works")
    assert "unavailable" in CT.ensure(ctx)


def test_research_versus_burden(ctx, monkeypatch, tmp_path):
    import json as _json

    import httpx

    from research_agent.tools import burden as B
    from research_agent.tools import claims as C

    rows = [{"paper_id": f"p{i}", "source": "abstract", "corpus": "pmc", "year": 2021, "title": "t",
             "data": {"geography": g, "health_domains": ["malaria"]}} for i, g in enumerate(
        [["Kenya"], ["western Kenya"], ["Kenya", "Uganda"], ["Amhara region"], ["sub-Saharan Africa"], []])]
    monkeypatch.setattr(C, "_rows", lambda _ctx: rows)
    ctx.pg.execute("DELETE FROM burden")
    # WHO GHO download (mocked): Nigeria carries most cases and has no studies here
    data = {"MALARIA_EST_CASES": [("NGA", 2022, 68e6), ("KEN", 2022, 3e6), ("UGA", 2022, 12e6), ("ETH", 2022, 1e6),
                                  ("NGA", 2021, 65e6)],
            "MALARIA_EST_DEATHS": [("NGA", 2022, 180e3), ("KEN", 2022, 10e3)]}

    def handler(request):
        code = request.url.path.rsplit("/", 1)[-1]
        return httpx.Response(200, text=_json.dumps({"value": [
            {"SpatialDim": i, "SpatialDimType": "COUNTRY", "TimeDim": y, "NumericValue": v} for i, y, v in data[code]]}))
    monkeypatch.setattr(B, "_transport", httpx.MockTransport(handler))
    assert B.fetch_who(ctx.pg)["country_years"] == 5
    rb = B.research_vs_burden(ctx)
    by = {x["iso3"]: x for x in rb["countries"]}
    assert (by["KEN"]["papers"], by["UGA"]["papers"], by["ETH"]["papers"], by["NGA"]["papers"]) == (3, 1, 1, 0)
    assert rb["burden_year"] == 2022 and abs(by["NGA"]["burden_share"] - 68 / 84) < 0.001
    assert by["KEN"]["ratio"] > 1 and "NGA" in [x["iso3"] for x in rb["under_researched"]]
    assert rb["regions"] == {"Sub-Saharan Africa": 1} and rb["papers_without_place"] == 1
    md = "\n".join(B.burden_markdown(ctx))
    assert "## Research versus burden (computed)" in md and "Nigeria" in md
    assert B.burden_chart(ctx, str(tmp_path / "burden.png")) and (tmp_path / "burden.png").stat().st_size > 1000
    # the same places in a run about something else: no malaria burden, just where the studies come from
    for r in rows:
        r["data"]["health_domains"] = ["antimicrobial resistance"]
    other = B.research_vs_burden(ctx)
    assert not other["burden_applies"] and all(x["burden_share"] is None for x in other["countries"])
    md = "\n".join(B.burden_markdown(ctx))
    assert "## Where the studies come from (computed)" in md and "malaria" not in md and "Kenya" in md
    assert B.burden_chart(ctx, str(tmp_path / "burden2.png")) is None
    # a CSV from another source, with country names instead of codes
    f = tmp_path / "map.csv"
    f.write_text("Country,Year,Cases\nCôte d'Ivoire,2022,\"7,000,000\"\nAtlantis,2022,5\n")
    res = B.import_csv(ctx.pg, str(f), "MAP")
    assert res["rows"] == 1 and res["skipped"] == ["Atlantis"]


def test_exports_in_every_format(loaded_db, tmp_path):
    import io
    import zipfile

    from fastapi.testclient import TestClient
    from openpyxl import load_workbook

    from research_agent import cli
    from research_agent.agents.orchestrator import run_research
    from research_agent.api import main as api
    from research_agent.exports import FORMATS, blocks, export

    rid = run_research("What ML methods are used for malaria forecasting?", mode="pipeline",
                       llm_factory=lambda strong=False: FakeLLM())["run_id"]
    got = {}
    for fmt in FORMATS:
        if fmt == "burden_chart":
            continue   # needs burden data; covered in the burden test
        data, name, media = export(rid, fmt)
        assert data and name.startswith(rid[:8]), fmt
        got[fmt] = data
    assert zipfile.is_zipfile(io.BytesIO(got["docx"])) and zipfile.is_zipfile(io.BytesIO(got["xlsx"]))
    wb = load_workbook(io.BytesIO(got["xlsx"]))
    assert {"Papers", "Reported results", "Associations", "Claims", "Screening", "PRISMA"} <= set(wb.sheetnames)
    assert wb["Papers"].max_row > 1 and wb["Claims"].max_row > 1
    bib = got["bib"].decode()
    from research_agent.db import get_conn
    with get_conn() as pg:
        n_papers = pg.execute("SELECT count(*) n FROM run_papers WHERE run_id=%s", (rid,)).fetchone()["n"]
    assert bib.count("\n@") + bib.startswith("@") == n_papers
    assert "archivePrefix = {arXiv}" in bib and got["ris"].decode().count("ER  - ") == bib.count("@")
    assert "<table>" in got["html"].decode() and 'class="mermaid"' in got["html"].decode()
    assert got["protocol"].decode().startswith("# Review protocol")
    kinds = [b[0] for b in blocks("## A\n\nText **bold**.\n\n- one\n- two\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n```mermaid\nx\n```")]
    assert kinds == ["h", "p", "ul", "table", "code"]
    # API and terminal
    client = TestClient(api.app)
    r = client.get(f"/runs/{rid}/export/docx")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    assert client.get(f"/runs/{rid}/export/exe").status_code == 404
    out = tmp_path / "r.bib"
    assert cli.main(["export", rid, "-f", "bib", "-o", str(out)]) in (0, None) and out.read_text().startswith("@")


def test_job_queue_claims_retries_and_recovers(loaded_db, monkeypatch):
    import uuid

    from research_agent import jobs
    from research_agent.db import get_conn

    with get_conn() as pg:
        pg.execute("UPDATE jobs SET status='done' WHERE status IN ('queued','running')")
    monkeypatch.setattr(jobs, "_llm_factory", lambda strong=False: FakeLLM())
    monkeypatch.setattr(jobs, "RETRY_DELAY_SECONDS", 0)
    # two workers never get the same job
    a, b = jobs.enqueue("noop", {}), jobs.enqueue("noop", {})
    c1, c2 = jobs.connect(), jobs.connect()
    try:
        j1, j2 = jobs.Worker("w1").claim(c1), jobs.Worker("w2").claim(c2)
        assert {j1["id"], j2["id"]} == {a, b} and j1["id"] != j2["id"]
        assert jobs.Worker("w3").claim(c1) is None
    finally:
        c1.close(); c2.close()
    # a failing job is retried once, then marked failed along with its follow-up
    run_id = str(uuid.uuid4())
    with get_conn() as pg:
        pg.execute("INSERT INTO runs (run_id, question, status) VALUES (%s,'q','done')", (run_id,))
        ans = pg.execute("INSERT INTO followups (run_id, role, status) VALUES (%s,'assistant','pending') RETURNING id",
                         (run_id,)).fetchone()["id"]
    bad = jobs.enqueue("followup", {"run_id": run_id, "question": "x", "answer_id": ans})
    monkeypatch.setattr(jobs, "execute", lambda kind, payload, attempt: (_ for _ in ()).throw(RuntimeError("boom")))
    w = jobs.Worker("w4")
    with get_conn() as pg:
        pg.execute("UPDATE jobs SET status='done' WHERE id IN (%s,%s)", (a, b))
    assert w.run_once() and w.run_once() and not w.run_once()
    with get_conn() as pg:
        job = pg.execute("SELECT status, attempts, error FROM jobs WHERE id=%s", (bad,)).fetchone()
        f = pg.execute("SELECT status FROM followups WHERE id=%s", (ans,)).fetchone()
    assert (job["status"], job["attempts"]) == ("failed", 2) and "boom" in job["error"] and f["status"] == "failed"
    # a worker that stops sending heartbeats: its job goes back to the queue
    stale = jobs.enqueue("run", {"run_id": run_id, "question": "q"})
    with get_conn() as pg:
        pg.execute("UPDATE jobs SET status='running', attempts=1, heartbeat_at=now() - interval '1 hour' WHERE id=%s", (stale,))
        assert jobs.requeue_stale(pg) == 1
        assert pg.execute("SELECT status FROM jobs WHERE id=%s", (stale,)).fetchone()["status"] == "queued"
        pg.execute("UPDATE jobs SET status='done' WHERE id=%s", (stale,))


def test_queued_run_is_worked_and_a_retry_resumes(loaded_db, monkeypatch):
    import uuid

    from research_agent import jobs
    from research_agent.db import get_conn

    with get_conn() as pg:
        pg.execute("UPDATE jobs SET status='done' WHERE status IN ('queued','running')")
    monkeypatch.setattr(jobs, "_llm_factory", lambda strong=False: FakeLLM())
    run_id = str(uuid.uuid4())
    with get_conn() as pg:
        pg.execute("INSERT INTO runs (run_id, question, status) VALUES (%s,%s,'queued')",
                   (run_id, "What ML methods are used for malaria forecasting?"))
    jobs.enqueue("run", {"run_id": run_id, "question": "What ML methods are used for malaria forecasting?",
                         "mode": "pipeline"})
    assert jobs.Worker("w").run_once()
    with get_conn() as pg:
        run = pg.execute("SELECT status, report_md FROM runs WHERE run_id=%s", (run_id,)).fetchone()
    assert run["status"] == "done" and run["report_md"]
    # attempt 2 of the same run resumes it rather than starting a new one
    calls = []
    import research_agent.agents.orchestrator as O
    monkeypatch.setattr(O, "run_research", lambda *a, **k: calls.append(k))
    jobs.execute("run", {"run_id": run_id, "question": "q", "mode": "pipeline"}, attempt=2)
    assert calls and calls[0]["resume"] == run_id


def test_drafts_a_proposal_and_a_review_manuscript(loaded_db, monkeypatch, tmp_path):
    import dataclasses
    import io
    import zipfile

    from fastapi.testclient import TestClient

    from research_agent import cli, jobs
    from research_agent.agents import draft as D
    from research_agent.agents.orchestrator import run_research
    from research_agent.api import main as api
    from research_agent.db import get_conn
    from research_agent.exports import export_draft

    fake = lambda strong=False: FakeLLM()   # noqa: E731
    rid = run_research("What ML methods are used for malaria forecasting?", mode="pipeline", llm_factory=fake)["run_id"]
    with get_conn() as pg:
        before = pg.execute("SELECT status, report_md, finished_at FROM runs WHERE run_id=%s", (rid,)).fetchone()

    # proposal, with the number check off so the untraced number stays visible as [unverified]
    monkeypatch.setattr(D, "settings", dataclasses.replace(D.settings, number_check=False))
    did = D.start(rid, "proposal", "External validation of malaria forecasts in the settings that use them")
    md = D.run_draft(did, llm_factory=fake)["content_md"]
    assert md.startswith("# Validating malaria forecasts")
    order = [md.index(f"## {h}") for h in ("Summary", "Background and rationale", "The evidence gap",
                                           "Proposed methods", "Workplan and timeline", "Budget")]
    assert order == sorted(order)                                      # summary first, written last
    assert "—" not in md and "0000.00000" not in md and "citation removed" in md
    assert "7 of 9 [unverified]" in md and "[to be confirmed: number of districts]" in md
    assert "applicant's proposal, not findings" in md
    # studies are cited as readers expect: author-year in the text, a full reference list, no internal ids left
    assert re.search(r"\(Author and Author, (19|20)\d\d[a-z]?(; Author and Author, (19|20)\d\d[a-z]?)*\)", md)
    assert "[arXiv:" not in md and "(arXiv:" not in md                       # "(arXiv:ID)" was normalised too
    assert "9.99 [unverified]" in md                                          # a number the cited study lacks
    assert "## Appendix A. Counts behind the numbers" in md and "## References (" in md
    refs = md.split("## References (", 1)[1]
    assert "https://arxiv.org/abs/" in refs and "[preprint]" in refs
    d = D.get_draft(did)
    assert d["status"] == "done" and d["meta"]["new_claims"] and d["meta"]["cited_papers"]
    assert len(d["meta"]["cited_in_text"]) >= 2 and d["meta"]["audit"]["unverified_attributions"] >= 1
    with get_conn() as pg:
        agent = pg.execute("SELECT agent FROM claims WHERE id=%s", (int(d["meta"]["new_claims"][0][1:]),)).fetchone()
        after = pg.execute("SELECT status, report_md, finished_at FROM runs WHERE run_id=%s", (rid,)).fetchone()
    assert agent["agent"] == "draft" and after == before              # the run and its report are untouched

    # review manuscript, with the number check on
    monkeypatch.setattr(D, "settings", dataclasses.replace(D.settings, number_check=True))
    rev = D.run_draft(D.start(rid, "review", "", citation_style="numbered"), llm_factory=fake)["content_md"]
    assert re.search(r"\[\d+(, \d+|-\d+)*\]", rev) and "\n1. Author A, Author B." in rev   # Vancouver
    for h in ("Abstract", "Introduction", "Methods", "Results: study selection and characteristics", "Discussion",
              "Conclusions", "Declarations"):
        assert f"## {h}" in rev, h
    assert "Note on automation" in rev and rev.index("## Abstract") < rev.index("## Introduction")

    # exports: Word opens, references hold exactly the cited papers
    data, name, _ = export_draft(did, "docx")
    assert zipfile.is_zipfile(io.BytesIO(data)) and name.startswith(f"proposal_{did}_")
    bib = export_draft(did, "bib")[0].decode()
    assert bib.count("@") == len(d["meta"]["cited_papers"])
    assert "<h1>" in export_draft(did, "html")[0].decode()

    # web: directions, a queued draft worked by the queue, downloads
    monkeypatch.setattr(jobs, "_llm_factory", fake)
    client = TestClient(api.app)
    assert client.get(f"/runs/{rid}/directions").status_code == 200
    assert client.post(f"/runs/{rid}/drafts", json={"kind": "grant"}).status_code == 422
    assert client.post(f"/runs/{rid}/drafts", json={"kind": "proposal", "about": ["X; DROP"]}).status_code == 422
    assert client.post("/runs/00000000-0000-0000-0000-000000000000/drafts", json={"kind": "review"}).status_code == 404
    r = client.post(f"/runs/{rid}/drafts", json={"kind": "proposal", "direction": "Probabilistic forecasts"})
    assert r.status_code == 202
    assert jobs.Worker(name="test").run_once()
    got = client.get(f"/drafts/{r.json()['draft_id']}").json()["draft"]
    assert got["status"] == "done" and got["content_md"].startswith("# ")
    assert [x["id"] for x in client.get(f"/runs/{rid}/drafts").json()["drafts"]][0] == got["id"]
    assert client.get(f"/drafts/{got['id']}/export/ris").status_code == 200
    assert client.get(f"/drafts/{got['id']}/export/pdf").status_code == 404
    # terminal
    assert cli.main(["draft", rid, "--directions"]) == 0
    out = tmp_path / "p.md"
    assert cli.main(["draft", "--export", str(did), "-f", "md", "-o", str(out)]) == 0
    assert out.read_text().startswith("# Validating")


def test_run_graph_and_mind_map(loaded_db):
    from fastapi.testclient import TestClient

    from research_agent.agents.orchestrator import run_research
    from research_agent.api import main as api
    from research_agent.runstate import RunContext
    from research_agent.tools.graph import build_graph

    rid = run_research("What ML methods are used for malaria forecasting?", mode="pipeline",
                       llm_factory=lambda strong=False: FakeLLM())["run_id"]
    ctx = RunContext.attach(rid, llm_factory=lambda strong=False: None)
    try:
        g = build_graph(ctx)
        n_papers = len(ctx.shortlist_ids())
    finally:
        ctx.close()
    by = {n["id"]: n for n in g["nodes"]}
    types = {n["type"] for n in g["nodes"]}
    assert {"question", "category", "concept", "paper", "claim"} <= types
    assert sum(1 for n in g["nodes"] if n["type"] == "paper") == n_papers
    assert all(l["source"] in by and l["target"] in by for l in g["links"])       # no dangling links
    # a paper links to the concepts it uses, and each concept's count is the papers linked to it
    for c in (n for n in g["nodes"] if n["type"] == "concept"):
        users = {l["source"] for l in g["links"] if l["target"] == c["id"] and l["type"] == "uses"}
        assert len(users) == c["papers"], c["label"]
    # a supported claim links to the papers it counted
    sup = next(n for n in g["nodes"] if n["type"] == "claim" and n["status"] == "supported")
    assert any(l["source"] == sup["id"] and l["type"] == "counted" for l in g["links"])
    # the mind map: question -> categories -> concepts -> papers, every ref a real node
    tree = g["tree"]
    refs = []

    def walk(t):
        if t.get("ref"):
            refs.append(t["ref"])
        for c in t.get("children", []):
            walk(c)
    walk(tree)
    assert tree["ref"] == "question" and all(r in by for r in refs)
    labels = [c["label"] for c in tree["children"]]
    assert labels[0] == "Methods" and "Evidence" in labels and "Research gaps" in labels
    assert labels.index("Evidence") < labels.index("Research gaps")          # gaps come last, after the evidence
    assert tree["subtitle"].startswith("What ML methods") and tree["label"] != tree["subtitle"]   # short title
    concept = next(c for s in tree["children"] if s["label"] == "Methods" for c in s["children"])
    assert concept["count"].endswith(f"/{n_papers} papers")                  # "3/16 papers", not "(+3)"
    assert concept["children"][0]["type"] == "paper" and concept["children"][0]["children"]   # paper -> attributes
    evidence = next(s for s in tree["children"] if s["label"] == "Evidence")
    sup_claim = next(c for g in evidence["children"] if g["label"] == "Supported claims" for c in g["children"])
    assert any(k["label"] == "Papers counted" for k in sup_claim["children"])
    client = TestClient(api.app)
    assert client.get(f"/runs/{rid}/graph").json()["stats"]["papers"] == n_papers
    assert client.get("/runs/00000000-0000-0000-0000-000000000000/graph").status_code == 404
    # papers are told apart by a short tag, not a cut-off title; each lists where else it sits in the map
    tags = [n["label"] for n in g["nodes"] if n["type"] == "paper"]
    assert len(set(tags)) == len(tags) and all(" · " in t for t in tags)
    first_paper = concept["children"][0]
    assert first_paper["count"].startswith("also in") and all(k["type"] in ("attr", "more") for k in first_paper["children"])


def test_graph_links_concepts_used_together_and_common_pairs_never_combined(loaded_db, monkeypatch):
    """Co-use links rank by lift (a concept nearly every paper has links to nothing), and two common concepts no
    paper combines get an "untried" link with the number expected by chance."""
    import random

    from research_agent.agents.orchestrator import run_research
    from research_agent.runstate import RunContext
    from research_agent.tools import extraction as E
    from research_agent.tools.graph import build_graph

    rng = random.Random(5)
    rows = []
    for i in range(60):
        lstm = i % 2 == 0
        rows.append({"paper_id": f"PMC8{i:05d}", "source": "abstract", "corpus": "pmc", "year": 2015 + i % 8,
                     "title": "Forecasting malaria incidence with machine learning",
                     "data": {"methods": ["LSTM"] if lstm else ["ARIMA"],
                              # LSTM goes with satellite data; ARIMA never does, though both are common
                              "data_modalities": ["satellite"] if lstm and rng.random() < .9 else ["surveillance counts"],
                              "geography": [rng.choice(["Kenya", "Uganda", "Malawi"])],
                              "study_designs": ["modelling study"]}})          # every paper: no co-use links
    monkeypatch.setattr(E, "_rows", lambda ctx: rows)
    rid = run_research("What ML methods are used for malaria forecasting?", mode="pipeline",
                       llm_factory=lambda strong=False: FakeLLM())["run_id"]
    ctx = RunContext.attach(rid, llm_factory=lambda strong=False: None)
    try:
        g = build_graph(ctx)
    finally:
        ctx.close()
    by = {n["id"]: n for n in g["nodes"]}
    pair = lambda l: {by[l["source"]]["label"], by[l["target"]]["label"]}
    co = [l for l in g["links"] if l["type"] == "co"]
    untried = [l for l in g["links"] if l["type"] == "untried"]
    assert any(pair(l) == {"LSTM", "satellite"} for l in co)
    assert not any("modelling study" in pair(l) for l in co)              # in every paper: links to nothing
    arima_sat = [l for l in untried if pair(l) == {"ARIMA", "satellite"}]
    assert arima_sat and arima_sat[0]["n"] == 0 and arima_sat[0]["expected"] >= 1.5
    assert not any(pair(l) == {"LSTM", "satellite"} for l in untried)


def test_old_runs_keep_the_extraction_version_they_were_read_with(loaded_db):
    """A schema change must not make an old run look unread: follow-ups, maps, drafts and exports read the
    records the run was built from."""
    from research_agent.agents.orchestrator import run_research
    from research_agent.agents.report import run_facts
    from research_agent.db import get_conn
    from research_agent.runstate import RunContext

    rid = run_research("What ML methods are used for malaria forecasting?", mode="pipeline",
                       llm_factory=lambda strong=False: FakeLLM())["run_id"]
    with get_conn() as pg:
        assert pg.execute("SELECT extraction_version FROM runs WHERE run_id=%s", (rid,)).fetchone()[
            "extraction_version"] == settings.extraction_schema_version
        ids = [r["paper_id"] for r in pg.execute("SELECT paper_id FROM run_papers WHERE run_id=%s", (rid,)).fetchall()]
        # the run's papers exist under an older version, all read in full; the current version has fewer rows
        pg.execute("""INSERT INTO extractions (paper_id, schema_version, source, data, model)
                      SELECT paper_id, 'health-old', 'fulltext', data, model FROM extractions
                      WHERE schema_version=%s AND paper_id = ANY(%s)""", (settings.extraction_schema_version, ids))
        pg.execute("DELETE FROM extractions WHERE schema_version=%s AND paper_id = %s",
                   (settings.extraction_schema_version, ids[0]))
        pg.execute("UPDATE runs SET extraction_version=NULL WHERE run_id=%s", (rid,))    # a run from before
    try:
        ctx = RunContext.attach(rid, llm_factory=lambda strong=False: FakeLLM())
        try:
            assert ctx.extraction_version == "health-old"
            facts = run_facts(ctx)
            assert facts["fulltext"] == facts["extracted"] == len(ids)
            assert ctx.child().extraction_version == "health-old"
        finally:
            ctx.close()
        with get_conn() as pg:   # pinned from now on
            assert pg.execute("SELECT extraction_version FROM runs WHERE run_id=%s", (rid,)).fetchone()[
                "extraction_version"] == "health-old"
    finally:
        with get_conn() as pg:
            pg.execute("DELETE FROM extractions WHERE schema_version='health-old'")


def test_details_credited_to_a_paper_must_be_in_it(loaded_db):
    from research_agent.agents.report import _N_OF_M, _int, audit_numbers
    from research_agent.db import get_conn
    from research_agent.tools.citing import check_attributions

    with get_conn() as pg:
        a = pg.execute("SELECT paper_id FROM papers WHERE title ILIKE '%%malawi%%random forest%%' LIMIT 1").fetchone()["paper_id"]
        b = pg.execute("SELECT paper_id FROM papers WHERE title ILIKE '%%kenya%%lstm%%' LIMIT 1").fetchone()["paper_id"]
        text = (f"A random-forest model in Malawi [arXiv:{a}] and a gradient-boosting study in Ethiopia [arXiv:{b}]. "
                f"A Kenyan study [arXiv:{b}]; an LSTM in Uganda [arXiv:{a}]. Counted [C12], among them a Malawian "
                f"one [arXiv:{a}], reporting an error of 9.99 [arXiv:{a}]. We will work in [to be confirmed: Ethiopia].")
        out, n = check_attributions(pg, text)
    assert "gradient-boosting [unverified]" in out and "Ethiopia [unverified]" in out       # wrong for paper b
    assert "LSTM [unverified]" in out and "Uganda [unverified]" in out                      # wrong for paper a
    assert "Kenyan study [arXiv" in out and "Malawian one [arXiv" in out                    # right: untouched
    assert "9.99 [unverified]" in out and "[to be confirmed: Ethiopia]" in out and n == 5
    # counts written in words, and "of the"
    t = "one of 49 papers, two of 39 report, 12 of the 49 were read, none of the 49 use it, one of the most common"
    assert [(_int(m.group(1)), _int(m.group(2))) for m in _N_OF_M.finditer(t)] == [(1, 49), (2, 39), (12, 49), (0, 49)]
    out, flagged = audit_numbers(t, {(1, 49), (0, 49)})
    assert flagged == 2 and "two of 39 [unverified]" in out and "12 of the 49 [unverified]" in out


def test_counts_must_be_measured_where_the_thing_is_recorded(ctx):
    from research_agent.tools.claims import PMC_TREND_CAVEAT, evaluate_trend, field_misfit, propose_claim

    assert field_misfit("evaluation_metrics", ["out-of-sample", "cross-validation"])
    assert field_misfit("datasets", ["NDVI"]) and not field_misfit("evaluation_metrics", ["RMSE", "AUC"])
    res = propose_claim(ctx, "Out-of-sample testing is rare.", "prevalence",
                        {"field": "evaluation_metrics", "any_of": ["out-of-sample"], "max_share": 0.1}, _agent="test")
    assert "error" in res and "metric names" in res["error"]
    r = evaluate_trend(ctx, {"keywords": "malaria", "early": [2016, 2018], "late": [2021, 2023],
                             "direction": "increase", "source": "pmc"})
    assert PMC_TREND_CAVEAT in r["caveat"]
    r = evaluate_trend(ctx, {"keywords": "malaria", "early": [2016, 2018], "late": [2021, 2023],
                             "direction": "increase", "source": "arxiv"})
    assert PMC_TREND_CAVEAT not in (r["caveat"] or "")


def test_drafts_speak_to_outside_readers():
    from research_agent.agents.draft import _dedupe_claims, plain_words
    from research_agent.agents.report import depth_warning

    t = plain_words("In the 49-paper shortlist, shortlisted papers use it; the shortlist is small. The rise was "
                    "not seen (claim C264, not supported). The extraction is thin, but the extraction records it.")
    assert "shortlist" not in t and "claim C264" not in t and "49 included studies" in t
    assert "the set of included studies is small" in t and "(not supported by the counts)" in t
    kept = _dedupe_claims([{"claim": "C1", "status": "supported"}, {"claim": "C2", "status": "supported"},
                           {"claim": "C3", "status": "supported"}],
                          {"C1": list("abcdefghijklmn"), "C2": list("abcdefghijklmno"), "C3": list("xyz")})
    assert [c["claim"] for c in kept] == ["C1", "C3"] and kept[0]["same_as"] == ["C2"]      # 14 vs 15 of 49
    assert depth_warning({"shortlist": 49, "fulltext": 0}).startswith("Only 0 of the 49 studies were read in full")
    assert depth_warning({"shortlist": 49, "fulltext": 20}) is None


# ---------------------------------------------------------------- open-ended researcher
def _synthetic_literature(seed: int, n: int = 240):
    """Studies with a real pattern (satellite data -> external validation, in every place and corpus), a
    confounded one (LSTM studies are mostly Kenyan, and Kenyan studies validate externally more) and noise."""
    import random

    rng, rows = random.Random(seed), []
    for i in range(n):
        place = rng.choice(["Kenya", "Uganda", "Malawi"])
        sat = rng.random() < 0.5
        lstm = rng.random() < (0.7 if place == "Kenya" else 0.1)
        p_ext = (0.55 if sat else 0.15) + (0.25 if place == "Kenya" else 0.0)
        rows.append({"paper_id": f"s{i}", "source": "fulltext" if i % 3 else "abstract",
                     "corpus": "pmc" if i % 2 else "arxiv", "year": 2012 + i % 12, "title": "t",
                     "data": {"data_modalities": ["satellite"] if sat else ["surveillance counts"],
                              "methods": ["LSTM"] if lstm else ["ARIMA"], "geography": [place],
                              "validation_level": "external" if rng.random() < p_ext else "internal",
                              "noise": [rng.choice(["alpha", "beta"])]}})
    return rows


def test_pattern_statistics():
    from research_agent.research.patterns import (binom_tail, fisher_exact, mantel_haenszel, signature, validate,
                                                  wilson)

    assert round(fisher_exact(8, 2, 1, 5), 4) == 0.035 and round(fisher_exact(3, 1, 1, 3), 4) == 0.4857
    assert [round(x, 3) for x in wilson(0, 49)] == [0.0, 0.073] and round(binom_tail(0, 49, 0.1, False), 4) == 0.0057
    # a difference that exists only because of the strata disappears once they are held apart
    confounded = mantel_haenszel([[30, 10, 15, 5], [2, 8, 10, 40]])          # within each stratum: equal rates
    assert abs(confounded["adjusted_difference"]) < 0.01 and confounded["p"] > 0.5
    real = mantel_haenszel([[30, 10, 5, 15], [15, 5, 3, 17]])
    assert real["adjusted_difference"] > 0.4 and real["p"] < 0.001
    a = {"kind": "difference", "outcome": [{"field": "methods", "any_of": ["LSTM", "GRU"]}],
         "group_a": [{"field": "geography", "any_of": ["Kenya"]}]}
    b = {"group_a": [{"any_of": ["kenya"], "field": "geography"}], "kind": "difference",
         "outcome": [{"field": "methods", "any_of": ["gru", "lstm"]}]}
    assert signature(a) == signature(b)                                       # same pattern, any order or case
    lists = ["methods", "geography", "evaluation_metrics"]
    assert validate(a, lists, {}) is None and "kind" in validate({"kind": "vibes"}, lists, {})
    assert "metric names" in validate({"kind": "prevalence", "max_share": 0.1, "outcome": [
        {"field": "evaluation_metrics", "any_of": ["out-of-sample"]}]}, lists, {})


def test_gauntlet_confirms_real_patterns_and_stops_confounded_ones():
    from collections import Counter

    from research_agent.research import gauntlet as GT
    from research_agent.research import patterns as P

    enums = {"validation_level": ["internal", "external", "not_stated"]}
    outcome = [{"field": "validation_level", "any_of": ["external"]}]
    rivals = [{"explanation": "place", "kind": "stratify", "field": "geography", "any_of": ["Kenya"]},
              {"explanation": "corpus", "kind": "stratify", "field": "corpus", "any_of": ["arxiv"]}]
    grades = {k: Counter() for k in ("real", "confounded", "noise")}
    for seed in range(12):
        disc, hold = P.split_rows(_synthetic_literature(seed), f"k{seed}")
        for name, group in (("real", {"field": "data_modalities", "any_of": ["satellite"]}),
                            ("confounded", {"field": "methods", "any_of": ["LSTM"]}),
                            ("noise", {"field": "noise", "any_of": ["alpha"]})):
            spec = {"kind": "difference", "outcome": outcome, "group_a": [group]}
            d = P.evaluate(spec, disc, enums)
            g = GT.grade(d, GT.subgroup_tests(spec, disc, enums, d["direction"]), GT.traps(spec, disc, d, enums),
                         [GT.test_alternative(spec, a, disc, enums, d) for a in rivals],
                         GT.holdout_test(spec, hold, enums, d["direction"], 0.05))
            grades[name][g["grade"]] += 1
    assert set(grades["real"]) <= {"strong", "moderate"} and grades["real"]["strong"] >= 10
    assert not (grades["confounded"]["strong"] or grades["confounded"]["moderate"])   # never confirmed
    assert set(grades["noise"]) == {"rejected"}
    # the rival that is the real cause is named as such
    disc, hold = P.split_rows(_synthetic_literature(7), "s1")
    spec = {"kind": "difference", "outcome": outcome, "group_a": [{"field": "methods", "any_of": ["LSTM"]}]}
    d = P.evaluate(spec, disc, enums)
    assert GT.test_alternative(spec, rivals[0], disc, enums, d)["verdict"] == "explains it"
    # traps: a corpus comparison is refused outright; statements carry no numbers
    corpus_spec = {"kind": "difference", "outcome": outcome, "group_a": [{"field": "corpus", "any_of": ["pmc"]}]}
    cd = P.evaluate(corpus_spec, disc, enums)
    assert any(t["status"] == "fail" for t in GT.traps(corpus_spec, disc, cd, enums))
    assert GT.plain_statement_problem("Satellite studies validate externally in 42 of 58 cases.")


def test_researcher_cycles_with_safeguards(loaded_db, monkeypatch):
    import dataclasses

    from research_agent.agents.orchestrator import run_research
    from research_agent.db import get_conn
    from research_agent.research import researcher as R

    monkeypatch.setattr(R, "settings", dataclasses.replace(R.settings, research_scope_min=0.0,
                                                           research_checkpoint_every=3, research_question_cycles=2))
    fake = lambda strong=False: FakeLLM()   # noqa: E731
    run_id = run_research("What ML methods are used for malaria forecasting?", mode="pipeline", llm_factory=fake)["run_id"]
    rid = R.create(run_id, "Understand how malaria forecasting studies choose methods and validation",
                   scope="malaria forecasting", out_of_bounds=["clinical treatment"], max_cycles=6)
    seeded = R.run_cycle(rid, llm_factory=fake)
    assert seeded["seeded"] == 3
    g = R.get(rid)
    by_q = {a["question"][:30]: a for a in g["agenda"]}
    assert by_q["How should clinical treatment "]["status"] == "needs_approval"          # out of bounds
    assert sum(a["status"] == "open" for a in g["agenda"]) == 2
    out = R.run_cycle(rid, llm_factory=fake)
    assert out["agenda_id"] and out["new_findings"] == 1
    g = R.get(rid)
    assert g["tests"] == 1                                    # re-testing the same pattern did not add a test
    finding = g["findings"][0]
    assert finding["grade"] in ("rejected", "provisional", "moderate", "strong") and finding["evidence"]["holdout"]
    child = [a for a in g["agenda"] if a["parent_id"]]
    assert child and child[0]["parent_id"] == out["agenda_id"]                           # sub-questions hang under
    # the researcher never sees the hidden half: reading one of its studies is refused
    ctx, _ = R._open(rid, fake)
    try:
        hidden = next(iter({r["paper_id"] for r in ctx._research.hold}), None)
        if hidden:
            assert "held-out" in R.read_study(ctx, hidden)["error"]
        # proposing the same tested pattern twice is refused; the verdict hides the held-out counts
        again = R.propose_finding(ctx, finding["test_id"], "Random forest studies cluster in Malawi again, surely.")
        assert "already judged" in again["error"]
    finally:
        ctx.close()
    # a checkpoint runs on cycle 3; questions that used their cycles are parked with the reason
    R.run_cycle(rid, llm_factory=fake)
    g = R.get(rid)
    assert any(n["kind"] == "checkpoint" for n in g["notebook"])
    R.run_cycle(rid, llm_factory=fake)
    R.run_cycle(rid, llm_factory=fake)        # parking is decided at the start of the cycle after the limit
    g = R.get(rid)
    assert any(a["status"] == "parked" and "cycles" in (a["status_note"] or "") for a in g["agenda"])
    # pause stops cycles; approval opens a waiting question
    R.set_status(rid, "paused", "test")
    assert R.run_cycle(rid, llm_factory=fake) == {"skipped": "paused"}
    waiting = next(a for a in g["agenda"] if a["status"] == "needs_approval")
    R.decide_question(rid, waiting["id"], True)
    assert next(a for a in R.get(rid)["agenda"] if a["id"] == waiting["id"])["status"] == "open"
    # budgets: the daily cap makes the next cycle wait; the total cap finishes the researcher
    with get_conn() as pg:
        pg.execute("UPDATE researchers SET status='active', daily_tokens=1 WHERE id=%s", (rid,))
    assert "waiting" in R.run_cycle(rid, llm_factory=fake)
    with get_conn() as pg:
        pg.execute("UPDATE researchers SET daily_tokens=100000000, total_tokens=1 WHERE id=%s", (rid,))
    assert R.run_cycle(rid, llm_factory=fake) == {"finished": "total budget used"}
    assert R.get(rid)["researcher"]["status"] == "finished"


def test_researcher_scope_check_and_web(loaded_db, monkeypatch):
    import dataclasses

    from fastapi.testclient import TestClient

    from research_agent import cli, jobs
    from research_agent.agents.orchestrator import run_research
    from research_agent.api import main as api
    from research_agent.db import get_conn
    from research_agent.research import researcher as R

    monkeypatch.setattr(R, "settings", dataclasses.replace(R.settings, research_scope_min=1.01))  # nothing is close
    monkeypatch.setattr(jobs, "_llm_factory", lambda strong=False: FakeLLM())
    run_id = run_research("What ML methods are used for malaria forecasting?", mode="pipeline",
                          llm_factory=lambda strong=False: FakeLLM())["run_id"]
    client = TestClient(api.app)
    assert client.post(f"/runs/{run_id}/researchers", json={"goal": "short"}).status_code == 422
    r = client.post(f"/runs/{run_id}/researchers", json={"goal": "Find promising directions in malaria forecasting",
                                                          "max_cycles": 3})
    assert r.status_code == 202
    rid = r.json()["researcher_id"]
    assert jobs.Worker(name="t").run_once()                  # the first cycle seeds the agenda...
    g = client.get(f"/researchers/{rid}").json()
    assert g["agenda"] and all(a["status"] == "needs_approval" for a in g["agenda"])   # ...all far from the charter
    with get_conn() as pg:                                    # ...and the next cycle was queued, after a pause
        nxt = pg.execute("SELECT run_at > now() AS later FROM jobs WHERE kind='research_cycle' AND status='queued' "
                         "AND payload->>'researcher_id' = %s", (str(rid),)).fetchone()
    assert nxt and nxt["later"]
    assert client.post(f"/researchers/{rid}/status", json={"action": "pause"}).json()["ok"]
    assert client.get(f"/researchers/{rid}").json()["researcher"]["status"] == "paused"
    aid = g["agenda"][0]["id"]
    assert client.post(f"/researchers/{rid}/agenda/{aid}", json={"approve": True}).json()["ok"]
    assert client.get(f"/runs/{run_id}/researchers").json()["researchers"][0]["id"] == rid
    assert client.post(f"/researchers/{rid}/status", json={"action": "stop"}).json()["ok"]
    assert client.post(f"/researchers/{rid}/status", json={"action": "resume"}).status_code == 409
    assert cli.main(["research", "status", str(rid)]) == 0


def test_wider_extraction_form_covers_pharmacology_and_amr():
    """health-v4: clinical, lab and pharmacology papers get their own fields, checked by quotes like the rest;
    runs read under v3 keep the v3 form; plasma AUC is never ranked like a ROC AUC; malaria burden stays out
    of other topics."""
    from types import SimpleNamespace

    from research_agent.tools.burden import is_malaria_run
    from research_agent.tools.extraction import GENERAL_FIELDS, fields_for, known_fields
    from research_agent.tools.reading import _verify
    from research_agent.tools.results import metric_key

    text = ("Title: Sertraline exposure in adolescents\n\nAbstract: In this randomised controlled trial of 84 "
            "adolescents with major depressive disorder, sertraline was given once daily. Sertraline acts by "
            "selective serotonin reuptake inhibition at the serotonin transporter. The AUC was 450 ng h/mL and "
            "the odds ratio for remission was 2.1 compared with placebo. Isolates of Klebsiella pneumoniae "
            "carried blaNDM-1.")
    args = {"study_designs": ["randomised controlled trial"], "populations": ["adolescents with MDD"],
            "organisms": ["Klebsiella pneumoniae"], "interventions": ["sertraline"],
            "mechanisms": ["serotonin reuptake inhibition", "HPA axis dysregulation"],
            "targets": ["serotonin transporter", "blaNDM-1"], "outcomes": ["remission"],
            "reported_results": [
                {"metric": "AUC", "value": "450", "model": "sertraline", "quote": "The AUC was 450 ng h/mL"},
                {"metric": "odds ratio", "value": "2.1", "model": "sertraline",
                 "quote": "the odds ratio for remission was 2.1 compared with placebo"}],
            "evidence": {"study_designs": ["In this randomised controlled trial of 84 adolescents"],
                         "populations": ["84 adolescents with major depressive disorder, sertraline was given"],
                         "organisms": ["Isolates of Klebsiella pneumoniae carried blaNDM-1"],
                         "interventions": ["sertraline was given once daily"],
                         "mechanisms": ["the drug suppresses cortisol through the HPA axis"],   # not in the text
                         "targets": ["selective serotonin reuptake inhibition at the serotonin transporter"]}}
    base, _ = _verify(args, text, None, True)
    assert base["interventions"] == ["sertraline"] and base["organisms"] == ["Klebsiella pneumoniae"]
    assert base["study_designs"] == ["randomised controlled trial"] and base["outcomes"] == ["remission"]
    assert base["mechanisms"] == [] and "mechanisms" in base["_unverified"]      # quote not found: dropped
    res = {r["metric"]: r for r in base["reported_results"]}
    assert res["plasma auc"]["higher_is_better"] is None and res["odds ratio"]["higher_is_better"] is None
    assert metric_key("AUROC") == ("auc", True) and metric_key("AUC0-inf")[0] == "plasma auc"
    assert metric_key("MIC90")[0] == "mic" and metric_key("hazard ratio")[0] == "hazard ratio"

    v3_lists, v3_ev, _ = fields_for("health-v3")
    assert not set(GENERAL_FIELDS) & set(v3_lists) and not set(GENERAL_FIELDS) & set(v3_ev)
    old = SimpleNamespace(extraction_version="health-v3", notes=lambda: {})
    new = SimpleNamespace(extraction_version="health-v4", notes=lambda: {})
    assert "interventions" not in known_fields(old)[0] and "interventions" in known_fields(new)[0]

    amr = SimpleNamespace(question="How common is carbapenem resistance in Klebsiella in East Africa?")
    rows = [{"data": {"health_domains": ["antimicrobial resistance"]}}] * 3
    assert not is_malaria_run(amr, rows)
    assert is_malaria_run(SimpleNamespace(question="Can ML forecast malaria?"), [])
    assert is_malaria_run(amr, [{"data": {"health_domains": ["malaria"]}}] * 2 + rows[:1])


def test_protocol_enum_values_state_one_condition_each():
    """A category joining two conditions with "and" leaves a study meeting only the first with nowhere to sit,
    so it gets coded down and the field under-counts. The missing rung is added; a genuine "both" is left."""
    from research_agent.opportunity.protocol import validate_protocol

    clean, problems = validate_protocol({"fields": [
        {"name": "mobile_element_context_reported", "type": "enum",
         "values": ["none", "sequence_context_only", "plasmid_characterised_and_transfer_tested"],
         "definition": "What the study reports about the genetic context of the resistance gene."},
        {"name": "isolate_source_setting", "type": "enum", "values": ["hospital", "community", "hospital_and_community"],
         "definition": "Where the isolates came from."},
    ]})
    ctx, setting = clean["fields"]
    assert ctx["values"] == ["none", "sequence_context_only", "plasmid_characterised",
                             "plasmid_characterised_and_transfer_tested", "not_stated"]
    assert any("plasmid_characterised" in p and "nowhere to sit" in p for p in problems)
    # "both" is a real category when each half is a category of its own: left exactly as the agent wrote it
    assert setting["values"] == ["hospital", "community", "hospital_and_community", "not_stated"]


def test_reread_fills_fields_that_an_abstract_could_not_quote(loaded_db, monkeypatch):
    """A field left empty because the abstract had no sentence to quote is filled once the paper is read in
    full; the run's report and claims are untouched."""
    from research_agent.agents.orchestrator import run_research
    from research_agent.db import get_conn
    from research_agent.runstate import RunContext
    from research_agent.tools.extraction import extract_papers, extraction_coverage

    rid = run_research("What ML methods are used for malaria forecasting?", mode="pipeline",
                       llm_factory=lambda strong=False: FakeLLM())["run_id"]
    with get_conn() as pg:
        report = pg.execute("SELECT report_md FROM runs WHERE run_id=%s", (rid,)).fetchone()["report_md"]
        claims = pg.execute("SELECT count(*) n FROM claims WHERE run_id=%s", (rid,)).fetchone()["n"]
        # put the run back as if every paper had been read from its abstract only
        pg.execute("""UPDATE extractions SET source='abstract' WHERE schema_version=%s
                      AND paper_id IN (SELECT paper_id FROM run_papers WHERE run_id=%s)""",
                   (settings.extraction_schema_version, rid))
    ctx = RunContext.attach(rid, llm_factory=lambda strong=False, step=None: FakeLLM())
    try:
        assert extraction_coverage(ctx)["by_source"]["fulltext"] == 0
        res = extract_papers(ctx, depth="fulltext", limit=3)
        read_in_full = extraction_coverage(ctx)["by_source"]["fulltext"]
        assert 0 < read_in_full <= res["newly_extracted"] == 3      # a paper with no full text stays on its abstract
        # papers already read in full are skipped; only the ones still on their abstract are tried again
        assert extract_papers(ctx, depth="fulltext", limit=3)["newly_extracted"] == 3 - read_in_full
        # force reads them all again, for a re-read after the protocol fields or the instructions changed
        assert extract_papers(ctx, depth="fulltext", limit=3, force=True)["newly_extracted"] == 3
    finally:
        ctx.close()
    with get_conn() as pg:
        assert pg.execute("SELECT report_md FROM runs WHERE run_id=%s", (rid,)).fetchone()["report_md"] == report
        assert pg.execute("SELECT count(*) n FROM claims WHERE run_id=%s", (rid,)).fetchone()["n"] == claims


def test_claim_states_separate_what_is_known_from_what_is_merely_unreported():
    """A count of zero is "not reported in this corpus", never "absent"; a zero resting on abstracts, or on a
    field papers mention without recording, is not even that."""
    from types import SimpleNamespace

    from research_agent.tools import epistemics as E

    def ctx(full=20, abstract=4):
        c = SimpleNamespace(run_id="r", pg=None)
        c._cov = {"by_source": {"fulltext": full, "abstract": abstract}, "extracted": full + abstract}
        return c

    import research_agent.tools.extraction as X
    real, X.extraction_coverage = X.extraction_coverage, lambda c: c._cov
    try:
        st = lambda res, **kw: E.state_of(ctx(**kw), "prevalence", res)                     # noqa: E731
        # a healthy measurement
        assert st({"supported": True, "n_matching": 9, "denominator": 24})["state"] == "supported"
        # supported, but most papers were read from their abstracts: the caveat is part of the state
        weak = st({"supported": True, "n_matching": 9, "denominator": 24}, full=4, abstract=20)
        assert weak["state"] == "partially_supported" and "abstract" in weak["reasons"][0]
        # nothing reports it: an observation about the corpus, and the state says the absence is not established
        zero = st({"supported": True, "n_matching": 0, "denominator": 24})
        assert zero["state"] == "not_reported" and "NOT established" in zero["reasons"][0]
        # the same zero, but the papers were mostly abstracts: the measurement cannot carry anything
        assert st({"supported": True, "n_matching": 0, "denominator": 24}, full=4, abstract=20
                  )["state"] == "not_searched_enough"
        # papers mention it in their text without it reaching the field
        assert st({"supported": False, "n_matching": 0, "denominator": 24,
                   "absence_problem": "11 papers mention it"})["state"] == "not_searched_enough"
        # too few papers to decide either way, whichever way the bound went
        assert st({"supported": True, "n_matching": 2, "denominator": 6})["state"] == "uncertain"
        assert st({"supported": False, "n_matching": 2, "denominator": 6})["state"] == "uncertain"
        # measured against the claim, on a base big enough to mean it
        assert st({"supported": False, "n_matching": 18, "denominator": 24})["state"] == "contradicted"
        # a trend needs ten matching papers before it is a trend at all
        assert E.state_of(ctx(), "trend", {"supported": True, "total_matching_papers": 4})["state"] == "uncertain"
        assert E.state_of(ctx(), "trend", {"supported": True, "total_matching_papers": 40})["state"] == "supported"
        caveat = E.state_of(ctx(), "trend", {"supported": True, "total_matching_papers": 40, "caveat": "PMC slice"})
        assert caveat["state"] == "partially_supported"
        assert set(E.PLAIN) == set(E.STATES)
    finally:
        X.extraction_coverage = real


def test_opportunities_are_rows_with_what_supports_and_weakens_them(loaded_db):
    """Gaps, untried combinations and designs become rows carrying their evidence and provenance, so the
    report is no longer the only place they exist."""
    from research_agent.db import get_conn
    from research_agent.opportunity.pipeline import run_map
    from research_agent.runstate import RunContext
    from research_agent.tools import opportunities as O

    rid = run_map("Can machine learning improve 1-6 month malaria forecasting?",
                  llm_factory=lambda strong=False: FakeLLM())["run_id"]
    ctx = RunContext.attach(rid, llm_factory=lambda strong=False: None)
    try:
        rows = O.listing(ctx)
        by_id = {r["item_id"]: r for r in rows}
        assert rows and {r["kind"] for r in rows} & {"gap", "combination", "design"}
        for r in rows:
            assert r["state"] == "open" and r["label"] and r["provenance"].get("agent")
            assert "supporting_claims" in r["evidence"] and "weakening_claims" in r["evidence"]
        # a claim the measurement could not settle is listed as weakening, never as supporting
        with get_conn() as pg:
            pg.execute("UPDATE claims SET state='not_searched_enough' WHERE run_id=%s AND id = "
                       "(SELECT min(id) FROM claims WHERE run_id=%s AND status <> 'rejected')", (rid, rid))
        O.record(ctx)
        ev = [r["evidence"] for r in O.listing(ctx)]
        assert not any(c in (e.get("supporting_claims") or []) for e in ev
                       for c in (e.get("weakening_claims") or []))
        # recording again keeps one row per item rather than duplicating it
        before = len(O.listing(ctx))
        O.record(ctx)
        assert len(O.listing(ctx)) == before
        # a design says what it would take, carried over from what the design agent proposed and code checked
        design = next((r for r in O.listing(ctx) if r["kind"] == "design"), None)
        if design:
            assert design["evidence"].get("addresses")
        # a researcher's finding carries its evidence across under the keys the finding really uses
        ev = {"pattern": "p", "rivals": [{"explanation": "method family", "verdict": "inconclusive"},
                                         {"explanation": "read depth", "verdict": "does not explain it"}],
              "examples": ["PMC1"], "discovery": {"p": 0.01}, "holdout": {"verdict": "replicated"}}
        with get_conn() as pg:
            rr = pg.execute("INSERT INTO researchers (run_id, charter, split) VALUES (%s,'{}'::jsonb,'{}'::jsonb) "
                            "RETURNING id", (rid,)).fetchone()["id"]
            pg.execute("INSERT INTO research_findings (researcher_id, statement, grade, status, reasons, evidence) "
                       "VALUES (%s,'Validation is under-reported','moderate','provisional','[]'::jsonb,%s::jsonb)",
                       (rr, __import__("json").dumps(ev)))
        O.record(ctx)
        fnd = next(r for r in O.listing(ctx) if r["kind"] == "finding")
        assert fnd["evidence"]["unresolved_questions"] == ["method family"]      # the rival it could not rule out
        assert len(fnd["evidence"]["alternative_explanations"]) == 2
        assert fnd["evidence"]["supporting_papers"] == ["PMC1"]

        assert O.set_state(ctx, rows[0]["item_id"], "dismissed", "already done")["state"] == "dismissed"
        assert "error" in O.set_state(ctx, rows[0]["item_id"], "nonsense")
        assert "error" in O.set_state(ctx, "ZZ9", "dismissed")
        md = "\n".join(O.markdown(ctx))
        assert "## Research opportunities (computed)" in md and by_id and list(by_id)[0] in md
    finally:
        ctx.close()


def test_field_pass_fills_one_field_from_full_text_and_only_with_a_quote(loaded_db, monkeypatch):
    """A field left blank by the first reading is filled by reading for it alone; a value whose quote is not
    in the paper is discarded, and nothing else in the record is touched."""
    from research_agent.agents.orchestrator import run_research
    from research_agent.db import get_conn
    from research_agent.runstate import RunContext
    from research_agent.tools import fieldpass as FP
    from research_agent.tools.extraction import _rows

    rid = run_research("What ML methods are used for malaria forecasting?", mode="pipeline",
                       llm_factory=lambda strong=False: FakeLLM())["run_id"]
    ctx = RunContext.attach(rid, llm_factory=lambda strong=False, step=None: FakeLLM())
    try:
        with get_conn() as pg:   # start from a run where nobody recorded how models were validated
            pg.execute("""UPDATE extractions SET data = jsonb_set(data, '{validation_level}', '"not_stated"')
                          WHERE schema_version=%s AND paper_id IN
                                (SELECT paper_id FROM run_papers WHERE run_id=%s)""",
                       (settings.extraction_schema_version, rid))
        before = _rows(ctx)
        assert all(r["data"]["validation_level"] == "not_stated" for r in before)
        other = {r["paper_id"]: r["data"].get("methods") for r in before}

        res = FP.fill_field(ctx, "validation_level", limit=6)
        assert res["blank_before"] == len(before) and res["read_now"] <= 6
        after = {r["paper_id"]: r["data"] for r in _rows(ctx)}
        filled = [p for p, d in after.items() if d["validation_level"] != "not_stated"]
        assert len(filled) == res["filled"]
        for pid in filled:                       # every filled value carries a quote, and only it changed
            assert (after[pid].get("evidence") or {}).get("validation_level")
            assert after[pid].get("_filled_by_pass", {}).get("validation_level")
            assert after[pid].get("methods") == other[pid]
        # a paper the pass could not answer for stays blank rather than being guessed at
        assert res["still_not_stated"] == res["read_now"] - res["filled"]

        # an answer whose quote is not in the paper is dropped, however confident it looks
        monkeypatch.setattr(FP, "_ask", lambda llm, fd, text: {"value": "external", "quote": "not in any paper",
                                                               "filled": False, "dropped_without_quote": True})
        with get_conn() as pg:
            pg.execute("""UPDATE extractions SET data = jsonb_set(data, '{validation_level}', '"not_stated"')
                          WHERE schema_version=%s AND paper_id IN
                                (SELECT paper_id FROM run_papers WHERE run_id=%s)""",
                       (settings.extraction_schema_version, rid))
        res2 = FP.fill_field(ctx, "validation_level", limit=4)
        assert res2["filled"] == 0 and res2["dropped_without_a_quote"] == res2["read_now"]
        assert all(r["data"]["validation_level"] == "not_stated" for r in _rows(ctx))
        assert "error" in FP.fill_field(ctx, "no_such_field")
    finally:
        ctx.close()


def test_a_field_pass_can_define_a_finer_field_when_the_old_one_merges_two_answers(loaded_db, monkeypatch):
    """A category that joins two things the question separates ("naive_or_seasonal_naive_baseline") cannot be
    un-merged by reading harder. A pass defines the finer field and codes every paper into it, and the old
    field keeps every value it had."""
    from research_agent.agents.orchestrator import run_research
    from research_agent.runstate import RunContext
    from research_agent.tools import fieldpass as FP
    from research_agent.tools.extraction import _rows, known_fields

    rid = run_research("What ML methods are used for malaria forecasting?", mode="pipeline",
                       llm_factory=lambda strong=False: FakeLLM())["run_id"]
    ctx = RunContext.attach(rid, llm_factory=lambda strong=False, step=None: FakeLLM())
    try:
        ctx.save_note("protocol", {"protocol": {"fields": [
            {"name": "q_baseline_comparison", "type": "enum", "definition": "What it was compared against.",
             "values": ["naive_or_seasonal_naive_baseline", "other_ml_models_only", "not_stated"],
             "desirable": [], "search": {}, "groups": [], "role": ""}]}})
        before = {r["paper_id"]: r["data"].get("q_baseline_comparison") for r in _rows(ctx)}

        monkeypatch.setattr(FP, "_ask", lambda llm, fd, text: {
            "value": "seasonal_naive_or_historical_expectance", "quote": text.split(".")[0], "filled": True})
        res = FP.field_pass(ctx, "baseline_reference", limit=4,
                            values=["no_reference", "other_models_only", "naive_persistence",
                                    "seasonal_naive_or_historical_expectance"],
                            definition="The reference forecast its accuracy is compared against.")
        assert res["added_field"]["field"] == "q_baseline_reference"
        assert res["added_field"]["values"][-1] == "not_stated" and res["filled"] >= 1

        lists, enums = known_fields(ctx)
        assert "q_baseline_reference" in enums and "naive_persistence" in enums["q_baseline_reference"]
        after = {r["paper_id"]: r["data"] for r in _rows(ctx)}
        assert any(d.get("q_baseline_reference") == "seasonal_naive_or_historical_expectance"
                   for d in after.values())
        for pid, old in before.items():      # the field it replaces is untouched: nothing already counted moves
            assert after[pid].get("q_baseline_comparison") == old
        # the same name twice is refused rather than silently redefining a field papers are already coded by
        assert "error" in FP.add_field(ctx, "baseline_reference", ["a", "b"], "again")
        assert "error" in FP.add_field(ctx, "one_sided", ["only"], "needs two categories")
    finally:
        ctx.close()


def test_more_budget_restarts_a_researcher_that_ran_out_of_room_but_not_one_that_ran_out_of_questions(loaded_db):
    """Raising the limit continues the work of a researcher that stopped for want of cycles or tokens. One
    that finished because its agenda was answered is left alone: it is not waiting for anything."""
    from research_agent.agents.orchestrator import run_research
    from research_agent.db import get_conn
    from research_agent.research import researcher as R

    rid = run_research("What ML methods are used for malaria forecasting?", mode="pipeline",
                       llm_factory=lambda strong=False: FakeLLM())["run_id"]
    r1 = R.create(rid, "Patterns in how models are validated", max_cycles=2, total_tokens=1000)
    with get_conn() as pg:
        pg.execute("UPDATE researchers SET cycles_done=2, tokens_used=400, status='finished', "
                   "status_note='cycle limit reached' WHERE id=%s", (r1,))
    out = R.set_budget(r1, add_cycles=5)
    assert out["max_cycles"] == 7 and out["resumed"] and R.get(r1)["researcher"]["status"] == "active"

    # out of tokens is the same case
    with get_conn() as pg:
        pg.execute("UPDATE researchers SET tokens_used=1000, status='finished', "
                   "status_note='total budget used' WHERE id=%s", (r1,))
    assert R.set_budget(r1, add_tokens=5000)["resumed"]

    # a limit below what it has already used is refused rather than silently ignored
    assert "error" in R.set_budget(r1, max_cycles=1)
    assert "error" in R.set_budget(r1, total_tokens=10)
    assert "error" in R.set_budget(999999, add_cycles=1)

    # finished because its questions ran out: more room changes nothing, and it says so
    r2 = R.create(rid, "Another charter", max_cycles=5)
    with get_conn() as pg:
        pg.execute("UPDATE researchers SET cycles_done=5, status='finished', "
                   "status_note='the agenda is empty' WHERE id=%s", (r2,))
    out2 = R.set_budget(r2, add_cycles=5)
    assert out2["max_cycles"] == 10 and not out2["resumed"] and "agenda" in out2["note"]
    assert R.get(r2)["researcher"]["status"] == "finished"


def test_cohort_scopes_every_count_and_keeps_silent_papers_visible(loaded_db, monkeypatch):
    """A run's scope is stored once and applied wherever papers are counted, so no test has to restate it.
    A paper that states an out-of-scope value is dropped; a paper that says nothing is kept and counted
    separately, because silence is about how it was read."""
    from research_agent.agents.orchestrator import run_research
    from research_agent.db import get_conn
    from research_agent.runstate import RunContext
    from research_agent.tools import claims as C
    from research_agent.tools import cohort as CO
    from research_agent.tools.extraction import _rows, value_counts

    rid = run_research("Which methods forecast malaria?", mode="pipeline",
                       llm_factory=lambda strong=False: FakeLLM())["run_id"]
    ctx = RunContext.attach(rid, llm_factory=lambda strong=False: None)
    try:
        # one paper read from its abstract never named a place: exactly the case a hard filter would drop
        with get_conn() as pg:
            pid = pg.execute("""SELECT e.paper_id FROM run_papers rp JOIN extractions e USING (paper_id)
                                WHERE rp.run_id=%s AND e.schema_version=%s
                                  AND e.data->'geography' <> '[]'::jsonb ORDER BY e.paper_id LIMIT 1""",
                             (rid, settings.extraction_schema_version)).fetchone()["paper_id"]
            pg.execute("UPDATE extractions SET data = jsonb_set(data, '{geography}', '[]'::jsonb) "
                       "WHERE paper_id=%s AND schema_version=%s", (pid, settings.extraction_schema_version))
        rows = _rows(ctx)
        places = {p.lower() for r in rows for p in r["data"].get("geography") or []}
        assert "kenya" in places and len(rows) >= 6
        in_africa = [r for r in rows if {"kenya", "uganda", "malawi", "tanzania"}
                     & {p.lower() for p in r["data"].get("geography") or []}]
        silent = [r for r in rows if not (r["data"].get("geography") or [])]
        assert in_africa and silent == [r for r in rows if r["paper_id"] == pid]

        # an unknown field is refused rather than silently counting nothing
        assert "error" in CO.set_cohort(ctx, include=[{"field": "continent", "any_of": ["africa"]}])
        assert "error" in CO.set_cohort(ctx, include=[{"field": "geography"}])
        assert "error" in CO.set_cohort(ctx, unstated="maybe", include=[{"field": "geography", "any_of": ["kenya"]}])

        res = CO.set_cohort(ctx, include=[{"field": "geography", "any_of": ["kenya", "uganda", "malawi",
                                                                           "tanzania"]}],
                            note="East African data only.")
        assert res["counted"] == len(in_africa) + len(silent)
        assert res["excluded_out_of_scope"] == len(rows) - res["counted"]
        assert res["unstated_on_a_cohort_field"] == len(silent) and res["unstated_are"] == "counted"
        assert "field pass" in res["warning"]            # silence is flagged, not quietly trusted

        # every count now counts the cohort: the row builder, value counts and a claim's denominator
        scoped = _rows(ctx)
        assert len(scoped) == res["counted"]
        assert not any("brazil" in p.lower() or "india" in p.lower()
                       for r in scoped for p in r["data"].get("geography") or [])
        assert value_counts(ctx, "geography")["n_papers"] == res["counted"]
        ev = C.evaluate_prevalence(ctx, {"field": "methods", "any_of": ["LSTM"], "min_count": 1})
        assert ev["denominator"] == res["counted"]

        # excluding silence drops the same papers, and says so in the description
        tight = CO.set_cohort(ctx, include=[{"field": "geography", "any_of": ["kenya"]}], unstated="exclude")
        assert tight["unstated_are"] == "excluded" and tight["counted"] < res["counted"]
        assert "silent on these fields are excluded" in CO.describe(CO.of(ctx))

        # a paper that states an out-of-scope value is out of scope even when it is silent on another field
        out = CO.set_cohort(ctx, exclude=[{"field": "geography", "any_of": ["brazil", "india", "usa"]}])
        assert out["counted"] == len(rows) - out["excluded_out_of_scope"] > len(in_africa)

        md = "\n".join(CO.markdown(ctx))
        assert "What the numbers are counted over" in md and "excluding geography in brazil" in md
        assert CO.clear(ctx)["cohort"] is None and len(_rows(ctx)) == len(rows)
        assert CO.markdown(ctx) == [] and CO.describe(None) == "every shortlisted paper"
    finally:
        ctx.close()


def test_precedent_check_grades_each_opportunity_against_the_whole_corpus(loaded_db):
    """An opportunity is searched for across the corpus, not just the shortlist, and graded by what was
    actually matched. An opportunity written as prose can never be graded a direct precedent."""
    from research_agent.opportunity.pipeline import run_map
    from research_agent.runstate import RunContext
    from research_agent.tools import opportunities as O
    from research_agent.tools import precedent as PR

    rid = run_map("Can machine learning improve 1-6 month malaria forecasting?",
                  llm_factory=lambda strong=False: FakeLLM())["run_id"]
    ctx = RunContext.attach(rid, llm_factory=lambda strong=False: None)
    try:
        assert "error" in PR.check(ctx, item_id="ZZ9")
        res = PR.check(ctx, limit=25)
        assert res["checked"] == len(O.listing(ctx)) and res["by_verdict"]
        assert set(res["by_verdict"]) <= set(PR.LEVELS)

        rows = {r["item_id"]: r for r in O.listing(ctx)}
        for item_id, row in rows.items():
            p = row["evidence"]["precedent"]
            assert p["verdict"] in PR.LEVELS and p["means"]
            # the verdict says what it was judged on, so a reader can disagree with it
            assert "term_groups" in p["judged_on"] and "topic_terms" in p["judged_on"]
            assert p["searched"]["corpus"].startswith("whole corpus")
            # a prose-only opportunity is capped: a sentence cannot show a paper did the same thing
            if not p["judged_on"]["from_structured_field"]:
                assert p["verdict"] in ("adjacent", "none") and "cannot" in p["means"]
            for level in ("direct", "partial", "adjacent"):
                for hit in p.get(level) or []:
                    assert hit["paper_id"] and "analysed_in_this_run" in hit

        # a gap on a field and a value IS structured, so it can reach direct, and its terms come from the value
        gap = next((r for r in rows.values() if r["kind"] == "gap" and r["value"]), None)
        if gap:
            req = PR.requirements(ctx, gap)
            assert req["structured"] and req["groups"]
            assert all(t not in PR._STOP for g in req["groups"] for t in g)
        # an untried combination asks for BOTH sides, carried as fields rather than parsed out of the label
        combo = next((r for r in rows.values() if r["kind"] == "combination"), None)
        if combo:
            assert len(PR.requirements(ctx, combo)["groups"]) == 2
            assert combo["evidence"].get("second_value")

        # the check does not change what the run says it screened
        from research_agent.tools.review import screening_rows

        before = len(screening_rows(ctx))
        PR.check(ctx, limit=25)
        assert len(screening_rows(ctx)) == before

        md = "\n".join(PR.markdown(ctx))
        assert "Has it been done already?" in md and list(rows)[0] in md
        # short terms are matched as whole words, so 'mic' does not count 'microscopy' as precedent
        assert PR._mentions("broth microdilution mic of 2", "mic")
        assert not PR._mentions("light microscopy of slides", "mic")

        # a value that only says which way a field points carries no meaning, so the field's words are used:
        # a gap on 'code_or_data_available = yes' must not grade every paper containing "yes" as precedent
        assert PR._field_and_value("code_or_data_available", "yes") == ["code", "data", "available"]
        assert PR._field_and_value("q_validation_split", "spatial_holdout") == ["spatial", "holdout"]
        thin = PR.requirements(ctx, {"kind": "gap", "field": "x", "value": "yes", "label": "l", "evidence": {}})
        assert not thin["structured"] and "any paper that" in thin["cap_reason"]

        # the grades themselves: both groups present is direct, one is partial, neither is adjacent
        req = {"groups": [["spatial", "holdout"], ["prevalence"]], "topic": ["malaria"], "structured": True,
               "cap_reason": ""}
        both = {"title": "Malaria prevalence mapping", "abstract": "spatial holdout validation across villages"}
        one = {"title": "Malaria prevalence survey", "abstract": "random cross-validation"}
        far = {"title": "Malaria case counts", "abstract": "an ARIMA baseline"}
        off = {"title": "Sepsis triage", "abstract": "spatial holdout prevalence"}
        assert PR._grade(both, req, None)[0] == "direct"
        assert PR._grade(one, req, None)[0] == "partial"
        assert PR._grade(far, req, None)[0] == "adjacent"
        assert PR._grade(off, req, None)[0] == "none"          # not about the run's subject at all
        # a paper the run analysed is judged on its checked extraction too, not only on its abstract
        assert PR._grade(one, req, {"q_validation_split": "spatial_holdout"})[0] == "direct"
        assert PR._grade(both, {**req, "structured": False}, None)[0] == "adjacent"
    finally:
        ctx.close()


def test_evidence_package_is_one_archive_that_explains_itself(loaded_db):
    """Everything needed to check a run, in one file, with a manifest that says what each part is and which
    questions the evidence cannot answer. A part that cannot be built is named as missing, not skipped."""
    import io
    import zipfile

    from research_agent.agents.orchestrator import run_research
    from research_agent.exports import PACKAGE, export
    from research_agent.runstate import RunContext
    from research_agent.tools import cohort as CO

    rid = run_research("Which methods forecast malaria?", mode="pipeline",
                       llm_factory=lambda strong=False: FakeLLM())["run_id"]
    ctx = RunContext.attach(rid, llm_factory=lambda strong=False: None)
    try:
        CO.set_cohort(ctx, include=[{"field": "geography", "any_of": ["kenya", "uganda", "malawi", "tanzania"]}])
    finally:
        ctx.close()

    data, name, media = export(rid, "package")
    assert name.endswith("evidence_package.zip") and media == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(data))
    assert not z.testzip()
    names = set(z.namelist())
    assert "MANIFEST.md" in names
    # the parts that must always be there, each non-empty
    for part in ("report.md", "claims.csv", "extractions.csv", "references.bib", "opportunities.csv"):
        assert part in names and len(z.read(part)) > 0, part

    manifest = z.read("MANIFEST.md").decode()
    assert ctx.question[:30] in manifest or "Question:" in manifest
    assert "geography in kenya" in manifest                  # the denominator is stated, not left implicit
    assert "What this package cannot tell you" in manifest
    assert "count of zero" in manifest and "reading depth" in manifest and "topic slices" in manifest
    # every file in the archive is described, and every described file is in the archive or named missing
    for part in names - {"MANIFEST.md"}:
        assert f"`{part}`" in manifest, part
    for part, _fmt, _what in PACKAGE:
        assert f"`{part}`" in manifest, part
    if any(p for p, _f, _w in PACKAGE if p not in names):
        assert "Not in here" in manifest

    # the report's own text is the report, not a rebuild of it
    with_report = z.read("report.md").decode()
    ctx = RunContext.attach(rid, llm_factory=lambda strong=False: None)
    try:
        stored = ctx.pg.execute("SELECT report_md FROM runs WHERE run_id=%s", (rid,)).fetchone()["report_md"]
        assert with_report == stored
    finally:
        ctx.close()


def test_thin_evidence_is_sent_back_to_the_papers_before_the_report_is_written(loaded_db, monkeypatch):
    """A claim whose state rests on reading depth has its abstract-only papers read in full and is measured
    again, before synthesis. A claim that is thin because the literature is thin is left alone and said to
    be. The attempt is recorded either way, including when it makes a claim worse."""
    from research_agent.agents.orchestrator import run_research
    from research_agent.db import get_conn
    from research_agent.runstate import RunContext
    from research_agent.tools import resolve as RS

    rid = run_research("Which methods forecast malaria?", mode="pipeline",
                       llm_factory=lambda strong=False: FakeLLM())["run_id"]
    ctx = RunContext.attach(rid, llm_factory=lambda strong=False: FakeLLM())
    try:
        # reading depth is what limits it: send it back to the papers
        assert RS._depth_limited({"papers_extracted": 9, "abstract_only": 6})
        assert RS._depth_limited({"papers_extracted": 9, "abstract_only": 0,
                                  "mentioned_but_not_recorded": "7 papers mention it"})
        # mostly read in full already: depth is not the problem, so do not spend reads on it
        assert not RS._depth_limited({"papers_extracted": 9, "abstract_only": 1})

        # a move is judged by code, and a settled answer against the claim is not called a failure
        assert RS._movement("not_searched_enough", "supported") == "better"
        assert RS._movement("supported", "supported") == "unchanged"
        assert RS._movement("partially_supported", "contradicted") == "settled against the claim"
        assert RS._movement("supported", "partially_supported") == "worse"

        cands = RS.candidates(ctx)
        assert cands, "the fixture run should leave some claim short of 'supported'"
        # too few papers to decide is NOT a reading problem: it is named and left alone
        with get_conn() as pg:
            cid = pg.execute("SELECT min(id) i FROM claims WHERE run_id=%s AND claim_type='prevalence' "
                             "AND status <> 'rejected'", (rid,)).fetchone()["i"]
            pg.execute("UPDATE claims SET state='uncertain' WHERE id=%s", (cid,))
        parked = next(c for c in RS.candidates(ctx) if c["id"] == cid)
        assert parked["skip"] and "cannot add papers" in parked["skip"]

        # Put this run's own claims in the state an abstract-heavy run produces. Only the claims rows are
        # touched: extraction records are shared by paper across runs, so rewriting those would change what
        # every other run in the database has read.
        with get_conn() as pg:
            pg.execute("""UPDATE claims SET state='not_searched_enough',
                              state_facts = coalesce(state_facts, '{}'::jsonb)
                                  || '{"papers_extracted": 12, "abstract_only": 9, "read_in_full": 3}'::jsonb
                          WHERE run_id=%s AND claim_type='prevalence' AND status <> 'rejected' AND id <> %s""",
                       (rid, cid))
        depth = [c for c in RS.candidates(ctx) if not c["skip"]]
        assert depth, "an abstract-only run should leave claims that a deeper read could settle"

        rec = RS.resolve(ctx, max_claims=2, per_claim=3)
        assert 0 < rec["attempted"] <= 2 and "resolution_rate" in rec
        assert any(s["claim"] == f"C{cid}" for s in rec["skipped"])
        for a in rec["attempts"]:
            assert a["claim"].startswith("C")
            assert "state_before" in a and ("state_after" in a or "failed" in a)
            # an attempt says why it was worth making, so the spend is auditable
            assert a.get("attempted_because") or a.get("note") or a.get("failed")
        # the original claim's text and predicate are never edited: only the records underneath it change
        with get_conn() as pg:
            after = pg.execute("SELECT text, predicate FROM claims WHERE id=%s", (depth[0]["id"],)).fetchone()
        assert after["text"] == depth[0]["text"] and after["predicate"] == depth[0]["predicate"]
        # whatever happened, it is on the run as a note and in the report's own words
        assert ctx.notes().get("evidence_resolution") == rec
        md = "\n".join(RS.markdown(ctx))
        assert "What a deeper read settled" in md and "Left alone" in md
        assert "same predicate was re-run" in md

        # the re-count is not left in the report as a second claim counting the same papers
        with get_conn() as pg:
            dupes = pg.execute("SELECT review_note FROM claims WHERE run_id=%s AND agent='followup' "
                               "AND status <> 'rejected'", (rid,)).fetchall()
        assert not dupes
    finally:
        ctx.close()


def test_an_automatically_cached_prefix_is_recorded_so_cost_is_not_overstated():
    """Groq reports a cached prefix on some models without any explicit cache control. Recording it keeps the
    run's cost honest, and makes a step whose prefix is NOT stable visible as a near-zero hit rate."""
    from types import SimpleNamespace

    from research_agent.llm import groq_client as GC
    from research_agent.runstate import _hit_rate

    def create(**_kw):
        msg = SimpleNamespace(content="ok", tool_calls=[])
        return SimpleNamespace(
            choices=[SimpleNamespace(message=msg, finish_reason="stop")],
            usage=SimpleNamespace(prompt_tokens=1000, completion_tokens=10,
                                  prompt_tokens_details=SimpleNamespace(cached_tokens=750)))

    c = GC.GroqClient.__new__(GC.GroqClient)
    c._client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    c.model, c.usage = "openai/gpt-oss-120b", GC.Usage()
    c.chat("sys", [{"role": "user", "content": "hi"}], max_tokens=64)
    assert c.usage.cached_input_tokens == 750 and c.usage.input_tokens == 1000

    # a provider that reports no such field must not break, and must not invent a hit
    def bare(**_kw):
        msg = SimpleNamespace(content="ok", tool_calls=[])
        return SimpleNamespace(choices=[SimpleNamespace(message=msg, finish_reason="stop")],
                               usage=SimpleNamespace(prompt_tokens=100, completion_tokens=2))
    c2 = GC.GroqClient.__new__(GC.GroqClient)
    c2._client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=bare)))
    c2.model, c2.usage = "m", GC.Usage()
    c2.chat("sys", [{"role": "user", "content": "hi"}], max_tokens=8)
    assert c2.usage.cached_input_tokens == 0

    assert _hit_rate({"input_tokens": 1000, "cached_input_tokens": 750}) == 0.75
    assert _hit_rate({"input_tokens": 0, "cached_input_tokens": 0}) is None


def test_novelty_needs_a_stratum_where_both_components_could_have_met(ctx, monkeypatch):
    """A corpus that splits into two literatures that never meet makes every cross-split pair look untried.
    The expectation is computed inside each kind of study, so a pair any one field already accounts for is
    dropped, and one that is genuinely absent where both components live survives."""
    from research_agent.opportunity import compute as C

    # 20 papers in two disjoint halves: 10 animal experiments, 10 human trials. Within the human half,
    # tissue sampling and questionnaires both occur and never together: that one IS an untried combination.
    rows = []
    for i in range(20):
        animal = i < 10
        data = {"q_population": "animal_model" if animal else "women",
                "q_design": "animal_experiment" if animal else "randomised_trial",
                "data_modalities": (["histopathology"] if animal else
                                    (["tissue"] if i < 15 else ["questionnaires"]))}
        rows.append({"paper_id": f"p{i}", "source": "fulltext", "corpus": "pmc", "year": 2015,
                     "title": "t", "data": data})
    protocol = {"fields": [
        {"name": "q_population", "type": "enum", "role": "setting", "definition": "x", "desirable": [],
         "values": ["animal_model", "women"]},
        {"name": "q_design", "type": "enum", "role": "evaluation", "definition": "y", "desirable": [],
         "values": ["animal_experiment", "randomised_trial"]}]}
    monkeypatch.setattr(C, "_rows", lambda _ctx: rows)
    monkeypatch.setattr(C, "known_fields", lambda _ctx: (["data_modalities"], {
        "q_population": ["animal_model", "women"], "q_design": ["animal_experiment", "randomised_trial"]}))
    monkeypatch.setattr(C, "_corpus_check", lambda *a: None)
    m = C.compute_map(ctx, protocol)
    pairs = [{n["a"]["value"], n["b"]["value"]} for n in m["novelty"]]

    # the animal/human split explains every one of these, so none is an opportunity
    assert {"women", "animal_experiment"} not in pairs
    assert {"animal_model", "randomised_trial"} not in pairs
    assert {"animal_model", "questionnaires"} not in pairs, "rats do not fill in questionnaires"
    assert {"women", "histopathology"} not in pairs

    # the stratified expectation is what did it, and it is zero for a cross-split pair
    ids = {v: {r["paper_id"] for r in rows
               if v in str(r["data"].values())} for v in ("animal_model", "women")}
    groups = dict(C.strata(rows, protocol, 4))["q_population"]
    assert C.stratified_expected(ids["animal_model"], ids["women"], groups) == 0
    assert C.explained_by_a_stratum(ids["animal_model"], ids["women"], C.strata(rows, protocol, 4), 1.5) \
        == "q_population"
    # a pair that both strata could have produced is not explained away
    everywhere = {f"p{i}" for i in range(0, 20, 2)}
    assert C.stratified_expected(everywhere, everywhere, groups) > 1.5

    # mirrors: the same fact in two fields needs a shared word, not just the same papers
    animal_pop = {"value": "animal_model"}
    animal_design = {"value": "animal_experiment"}
    same = {f"p{i}" for i in range(10)}
    assert C.mirrors(animal_pop, animal_design, same, same, 0.8)
    assert not C.mirrors({"value": "one_month_horizon"}, {"value": "satellite_data"}, same, same, 0.8)
    # and it folds them into one established item rather than reporting one fact twice
    est = {e["label"] for e in m["established"]}
    folded = [e for e in m["established"] if e.get("also_recorded_as")]
    assert folded and any("animal" in e["label"] for e in folded)
    assert len([x for x in est if "animal" in x]) == 1

    # a watch list where nothing leans the right way is not printed
    assert all(w["trend"]["p"] < C.THRESHOLDS["watch_max_p"] for w in m.get("watch", []))


def test_reported_values_are_pooled_only_within_one_unit(loaded_db):
    """A number without its unit is not a measurement. ng/dL and nmol/L are the same testosterone level, so
    they are converted before pooling; a unit that cannot be read keeps its quote but never enters a median;
    a value outside what its metric can be is a reading error and is dropped."""
    from research_agent.tools.results import (_CANONICAL, in_bounds, metric_key, normalise_unit,
                                              verify_structured)

    # total and free testosterone differ by about fifty, so they are never one bucket
    assert metric_key("serum testosterone")[0] == "total testosterone"
    assert metric_key("mean total testosterone")[0] == "total testosterone"
    assert metric_key("free testosterone")[0] == "free testosterone"
    assert metric_key("free androgen index")[0] == "free testosterone"

    # 350 ng/dL and 12.1 nmol/L are the same level and must land on the same scale
    v1, u1 = normalise_unit("total testosterone", "ng/dL", 350.0)
    v2, u2 = normalise_unit("total testosterone", "nmol/L", 12.1)
    assert u1 == u2 == "nmol/L" and abs(v1 - v2) < 0.2
    assert normalise_unit("total testosterone", "ng/mL", 3.5)[0] == pytest.approx(12.13, abs=0.01)
    assert normalise_unit("body weight", "g", 342.0) == (0.342, "kg")       # a rat, in kilograms
    # a unit the table does not know cannot be converted, and says so rather than guessing
    assert normalise_unit("total testosterone", "arbitrary units", 5.0) == (None, None)
    # a metric with no canonical unit keeps whatever was written, so like is still pooled with like
    assert normalise_unit("duration of attack", "s", 18.1) == (18.1, "s")

    # an impossible value is a reading error
    assert in_bounds("r2", 0.4) and not in_bounds("r2", 2.0)
    assert in_bounds("auc", 1.0) and not in_bounds("auc", 1.4)
    assert in_bounds("body weight", 342.0)                                  # unbounded metrics pass anything

    text = ("Mean total testosterone rose from 318 ng/dL at baseline, and the model reached an R2 of 2 in "
            "table 3, with free testosterone of 65 pg/mL reported separately.")
    args = {"reported_results": [
        {"metric": "mean total testosterone", "value": "318", "unit": "ng/dL", "model": "TRT arm",
         "quote": "Mean total testosterone rose from 318 ng/dL at baseline"},
        {"metric": "R2", "value": "2", "unit": "", "model": "regression",
         "quote": "the model reached an R2 of 2 in table 3"},
        {"metric": "free testosterone", "value": "65", "unit": "pg/mL", "model": "TRT arm",
         "quote": "free testosterone of 65 pg/mL reported separately"}]}
    kept, _dropped = verify_structured(args, text)
    got = {r["metric"]: r for r in kept["reported_results"]}
    assert "r2" not in got, "an R2 of 2 is arithmetically impossible and is not a result"
    assert got["total testosterone"]["unit"] == "nmol/L"
    assert got["total testosterone"]["value_canonical"] == pytest.approx(11.03, abs=0.01)
    assert got["total testosterone"]["value"] == "318"          # what the paper wrote is kept as written
    assert got["free testosterone"]["unit"] == "pmol/L"
    assert "total testosterone" in _CANONICAL and "free testosterone" in _CANONICAL


def test_a_design_field_never_makes_an_animal_trial_choose_what_it_was(loaded_db):
    """What was studied and how it was allocated are two axes. A design enum that mixes them is split, so an
    animal experiment randomised to treatment or vehicle keeps both facts."""
    from research_agent.opportunity.protocol import SYSTEM_FIELD, validate_protocol

    raw = {"fields": [
        {"name": "study_design_for_causal_inference", "type": "enum", "role": "evaluation",
         "definition": "The design used to estimate the effect.",
         "values": ["case_report_or_case_series", "animal_or_in_vitro_experiment",
                    "observational_cohort_or_registry", "randomised_placebo_controlled"],
         "desirable": ["randomised_placebo_controlled"]},
        {"name": "follow_up_duration", "type": "enum", "definition": "How long.",
         "values": ["up_to_3_months", "3_to_12_months", "more_than_12_months"], "desirable": []}]}
    clean, problems = validate_protocol(raw)
    f = {x["name"]: x for x in clean["fields"]}
    design = f["q_study_design_for_causal_inference"]
    assert "animal_or_in_vitro_experiment" not in design["values"]
    assert "randomised_placebo_controlled" in design["values"] and design["desirable"]
    assert "whatever was studied" in design["definition"]
    system = f[SYSTEM_FIELD]
    assert system["values"][0] == "human" and "not_stated" in system["values"]
    # one value naming two systems gives both a home
    assert system["values"] == ["human", "animal", "in_vitro_or_ex_vivo", "not_stated"]
    assert any("moved animal_or_in_vitro_experiment" in p for p in problems)
    # a field that is only about time, or only about allocation, is left exactly as it was
    assert f["q_follow_up_duration"]["values"][:3] == ["up_to_3_months", "3_to_12_months", "more_than_12_months"]

    # separate system values map to separate kinds, and an existing system field is reused, not duplicated
    raw2 = {"fields": [
        {"name": "study_system", "type": "enum", "definition": "x", "values": ["human", "rodent"]},
        {"name": "design", "type": "enum", "definition": "y",
         "values": ["cell_culture", "randomised_trial", "cohort_study"]}]}
    clean2, _ = validate_protocol(raw2)
    names = [x["name"] for x in clean2["fields"]]
    assert names.count(SYSTEM_FIELD) == 1
    assert "cell_culture" not in next(x for x in clean2["fields"] if x["name"] == "q_design")["values"]

    # with only one allocation value left there is no allocation axis to protect, so nothing moves
    raw3 = {"fields": [{"name": "kind", "type": "enum", "definition": "z",
                        "values": ["animal_study", "human_trial", "review"]}]}
    clean3, problems3 = validate_protocol(raw3)
    assert [x["name"] for x in clean3["fields"]] == ["q_kind"] and not any("moved" in p for p in problems3)


def test_papers_back_each_other_up_counted_in_independent_sources(ctx, monkeypatch):
    """Every reported association is set beside every other about the same driver and outcome. Wordings of
    one driver meet; an endogenous level and a treatment do not. Papers sharing a dataset are one source, and
    support from another kind of study is counted apart."""
    from research_agent.tools import corroboration as CB
    from research_agent.tools import extraction as E
    from research_agent.tools.results import associations

    def row(pid, assocs, organisms=None, datasets=None):
        return {"paper_id": pid, "source": "fulltext", "corpus": "pmc", "year": 2015, "title": pid,
                "data": {"organisms": organisms or [], "populations": [], "study_designs": [],
                         "datasets": datasets or [], "geography": [], "methods": [],
                         "reported_associations": [
                             {"driver": d.lower(), "driver_as_written": d, "outcome": o, "direction": dr,
                              "significant": sig, "lag": "", "quote": f"{pid}: {d} and {o}"}
                             for d, o, dr, sig in assocs]}}
    rows = [
        # haematocrit: three human papers (two from one registry) and one rat study, all the same way
        row("p1", [("testosterone replacement therapy", "haematocrit", "positive", "yes")], datasets=["REG-A"]),
        row("p2", [("TRT", "serum haematocrit levels", "positive", "yes")]),
        row("p3", [("testosterone therapy", "haematocrit", "positive", "yes")], datasets=["REG-A"]),
        row("p4", [("testosterone therapy", "haematocrit", "positive", "yes")], organisms=["Wistar rats"]),
        # PSA: two independent rises against one paper finding no change
        row("p5", [("testosterone therapy", "PSA", "positive", "yes")]),
        row("p6", [("TRT", "prostate specific antigen", "none", "no")]),
        row("p7", [("testosterone therapy", "PSA", "positive", "yes")]),
        # an endogenous level is a different driver from a treatment, and only one paper reports it
        row("p8", [("testosterone", "mortality", "negative", "yes")]),
        # agreement only between a human study and an animal one
        row("p11", [("testosterone therapy", "bone density", "positive", "yes")]),
        row("p12", [("testosterone therapy", "bone density", "positive", "yes")], organisms=["mice"]),
        # agreement only between two papers from the same registry
        row("p13", [("testosterone therapy", "waist circumference", "negative", "yes")], datasets=["REG-B"]),
        row("p14", [("testosterone therapy", "waist circumference", "negative", "yes")], datasets=["REG-B"]),
        # one paper each way, no replication on either side
        row("p15", [("testosterone therapy", "blood pressure", "positive", "yes")]),
        row("p16", [("testosterone therapy", "blood pressure", "negative", "yes")]),
        # every statement nonlinear or mixed
        row("p9", [("testosterone therapy", "depressive symptoms", "mixed", "not_stated")]),
        row("p10", [("testosterone therapy", "depressive symptoms", "nonlinear", "not_stated")]),
    ]
    from research_agent.tools import claims as CL

    monkeypatch.setattr(E, "_rows", lambda _ctx: rows)
    monkeypatch.setattr(CL, "_rows", lambda _ctx: rows)      # results reads papers through the claims module

    # grouping: what meets and what must not
    g = CB.group_phrases(["testosterone replacement therapy", "TRT", "testosterone therapy", "testosterone",
                          "serum testosterone levels"])
    assert g["TRT"]["group"] == g["testosterone therapy"]["group"] == g["testosterone replacement therapy"]["group"]
    assert g["testosterone"]["group"] == g["serum testosterone levels"]["group"]     # how it was measured is noise
    assert g["testosterone"]["group"] != g["testosterone therapy"]["group"], "a level is not a treatment"
    assert any("acronym" in r for r in g["TRT"]["why_merged"])
    assert not CB._acronym_of("age", "androgen gel exposure"), "only capitals are read as an acronym"
    # a qualified exposure is a different exposure, even where the malaria driver list would merge them
    rain = CB.group_phrases(["rainfall", "precipitation", "heavy rainfall", "excess rainfall"],
                            canonical={"rainfall": "rainfall", "precipitation": "rainfall",
                                       "heavy rainfall": "rainfall", "excess rainfall": "rainfall"})
    assert rain["rainfall"]["group"] == rain["precipitation"]["group"] == "rainfall"
    assert rain["heavy rainfall"]["group"] != "rainfall" and rain["excess rainfall"]["group"] != "rainfall"
    # for an outcome, the measure is noise: cases and incidence of one disease are the same outcome
    out = CB.group_phrases(["malaria incidence", "malaria cases", "malaria prevalence", "malaria mortality"],
                           outcomes=True)
    assert out["malaria incidence"]["group"] == out["malaria cases"]["group"] == out["malaria prevalence"]["group"]
    assert out["malaria mortality"]["group"] != out["malaria cases"]["group"], "death is not a measure of how much"

    # relations come from directions alone
    assert CB.relation("positive", "positive") == "supports"
    assert CB.relation("positive", "negative") == "contradicts"
    assert CB.relation("positive", "none") == "disputes_existence"
    assert CB.relation("positive", "mixed") == "qualifies"
    assert CB.relation("positive", "positive", "yes", "no") == "agrees_in_direction"

    res = CB.corroborate(ctx)
    by = {(f["driver"], f["outcome"]): f for f in res["findings"]}
    # the group's label is its most used wording in THIS run, and every wording is listed with it
    therapy = next(f["driver"] for f in res["findings"] if "TRT" in f["driver_wordings"])
    assert therapy == "testosterone therapy"
    assert g["TRT"]["group"] == "testosterone replacement therapy", "on a tie the spelled-out wording wins"

    hct = next(f for (d, o), f in by.items() if d == therapy and "haematocrit" in o)
    assert hct["verdict"] == "corroborated" and hct["papers"] == 4
    assert hct["independent_sources"] == 3, "p1 and p3 share a registry, so they are one source"
    p1 = next(s for s in hct["statements"] if s["paper_id"] == "p1")
    assert p1["independent_support"] == 2 and p1["independent_support_same_system"] == 1
    assert "p3" in p1["backed_by"] and any("REG-A" in x.lower() or "reg-a" in x for x in p1["linked_by"])
    assert set(hct["driver_wordings"]) >= {"TRT", "testosterone therapy", "testosterone replacement therapy"}

    psa = next(f for (d, o), f in by.items() if d == therapy and o in ("PSA", "prostate specific antigen"))
    assert psa["verdict"] == "contested" and psa["by_direction"]["positive"]["independent_sources"] == 2
    assert "p6" in next(s for s in psa["statements"] if s["paper_id"] == "p5")["against"]

    alone = next(f for (d, o), f in by.items() if o == "mortality")
    assert alone["verdict"] == "not_addressed_elsewhere" and alone["driver"] != therapy
    assert "not evidence against" in alone["means"]

    assert next(f for (d, o), f in by.items() if o == "bone density")["verdict"] == "corroborated_across_systems_only"
    assert next(f for (d, o), f in by.items() if o == "waist circumference")["verdict"] == \
        "repeated_by_related_papers_only"
    assert next(f for (d, o), f in by.items() if o == "blood pressure")["verdict"] == "contradicted"
    assert next(f for (d, o), f in by.items() if o == "depressive symptoms")["verdict"] == "qualified"

    # the run keeps it, the report says it in plain terms, and nothing internal leaks into the note
    note = ctx.notes()["corroboration"]
    assert note["findings"] and all("_row" not in s for f in note["findings"] for s in f["statements"])
    md = "\n".join(CB.markdown(ctx))
    assert "Do the papers back each other up?" in md and "independent sources" in md
    assert "not that they are unsupported" in md and "TRT" in md

    # the old contradictions code now meets these wordings too, instead of only the malaria list
    drivers = associations(ctx)["drivers"]
    assert therapy in drivers and "TRT" not in drivers

    # the follow-up tool sees the same thing, trimmed for a model
    tool = CB.corroborate_tool(ctx, driver="TRT")
    assert tool["findings"] and all(f["driver"] == therapy for f in tool["findings"])


def test_corpus_check_reads_full_texts_and_names_what_it_did_not_read(loaded_db):
    """The corpus check goes past titles and abstracts: it searches the full texts of the topic's papers and
    returns each match with its passage, because a mention is not a finding. Papers that mention it but were
    never analysed are named, and gap reasoning may not call an absence corpus-wide when it counted the
    analysed papers."""
    from research_agent.opportunity import compute as C
    from research_agent.opportunity.agents import scope_wording
    from research_agent.opportunity.pipeline import run_map
    from research_agent.opportunity.render import _corpus_lines
    from research_agent.runstate import RunContext
    from research_agent.tools.fulltext_check import search_fulltexts, search_terms

    rid = run_map("Can machine learning improve 1-6 month malaria forecasting?",
                  llm_factory=lambda strong=False: FakeLLM())["run_id"]
    ctx = RunContext.attach(rid, llm_factory=lambda strong=False: None)
    try:
        # the fixture papers report calibration in their results sections, never in an abstract
        c = C._corpus_check(ctx, "malaria", "calibration")
        assert c["matching"] == 0, "the abstracts do not say it"
        ft = c["fulltext"]
        assert ft["full_texts_read"] > 0 and ft["matching"] > 0, "the full texts do"
        cand = ft["candidates"][0]
        assert "calibration" in cand["passage"].lower(), "every match comes with the passage it came from"
        assert "analysed_in_this_run" in cand and ft["topic_papers_checked"] >= ft["full_texts_read"]
        # absence is only ever stated over texts that were actually searched
        none = search_fulltexts(ctx, "malaria", '"mosquito larval habitat mapping"')
        assert none["matching"] == 0 and none["full_texts_read"] == ft["full_texts_read"]
        # the topic's texts are fetched once per run, then reused by every gap
        assert ctx._topic_texts["malaria"]["readable"]
        # 0 turns the full-text tier off
        assert search_fulltexts(ctx, "malaria", "calibration", limit=0) is None

        # papers that mention it and were not analysed are counted and named
        assert "not_analysed" in c and isinstance(c["not_analysed_examples"], list)

        # and the grade reflects what the full texts say
        s = {"n_fulltext": 0, "N_fulltext": 20, "stated_rate": 0.9, "n": 0, "ci95": [0, 0.1]}
        found = C._gap_confidence(s, {**c, "share": 0.0, "fulltext": {**ft, "full_texts_read": 40, "matching": 3}}, {})
        clean = C._gap_confidence(s, {**c, "share": 0.0, "fulltext": {**ft, "full_texts_read": 40, "matching": 0}}, {})
        assert clean["points"] > found["points"]
        assert any("full text" in r for r in found["reasons"]) and any("absent from all 40" in r for r in clean["reasons"])
        unread = C._gap_confidence(s, {**c, "share": 0.0, "not_analysed": 6}, {})
        assert any("were not read" in r for r in unread["reasons"])

        lines = "\n".join(_corpus_lines({**c, "not_analysed": 2}))
        assert "Full texts:" in lines and "A mention is not a finding" in lines
        assert "not analysed in this run" in lines or "analysed" in lines
    finally:
        ctx.close()

    assert search_terms('malaria ("interaction" OR "effect modification" OR synergy) rainfall -review') == \
        ["interaction", "effect modification", "malaria", "synergy", "rainfall"]
    assert search_terms('rainfall NOT drought -"case report"') == ["rainfall"], "exclusions are never looked for"
    assert search_terms("long-term follow-up") == ["long-term", "follow-up"], "a hyphen inside a word is kept"

    # the sentence that reached a report: absence counted over 91 papers, written as the whole corpus
    said = ("No paper in the corpus reports a rainfall×temperature interaction or joint term (0/91; 0/32 "
            "full-text, H2). The corpus check found 10 papers in the corpus that mention interaction terms.")
    fixed, note = scope_wording(said, {"not_analysed": 10, "fulltext": {"matching": 4, "full_texts_read": 210}})
    assert fixed.startswith("No analysed paper reports a rainfall×temperature interaction")
    assert "found 10 papers in the corpus" in fixed, "a sentence that is not about absence is left alone"
    assert "10 further topic papers" in fixed and "4 of 210 topic papers mention it in their full text" in fixed
    assert note and "narrowed" in note
    # wording that already says what it counted is not touched
    ok = "None of the 91 analysed papers reports one; 10 further topic papers were not analysed."
    assert scope_wording(ok, {"not_analysed": 10})[0] == ok


def test_systematic_search_finds_screens_and_analyses_every_eligible_paper(loaded_db):
    """Systematic mode: the protocol's fixed query identifies papers, every hit is screened against the criteria
    with a recorded reason, every eligible paper is analysed (a seeded random sample past the cap), and no agent
    adds to the set afterwards."""
    from research_agent.agents.orchestrator import run_research, run_specialist
    from research_agent.opportunity.pipeline import run_map
    from research_agent.runstate import RunContext
    from research_agent.tools import systematic as SY
    from research_agent.tools.review import prisma_counts, screening_rows

    rid = run_map("Can machine learning improve 1-6 month malaria forecasting?",
                  llm_factory=lambda strong=False: FakeLLM(), search="systematic")["run_id"]
    ctx = RunContext.attach(rid, llm_factory=lambda strong=False: FakeLLM())
    try:
        n = ctx.notes()["search"]
        assert n["mode"] == "systematic" and not n["query_was_fallback"]
        assert n["query"] == "malaria AND (forecast OR forecasting OR prediction OR incidence)"
        assert n["identified"] >= n["screened"] >= n["eligible"] >= n["analysed"] > 0
        assert n["analysed"] == n["eligible"] and not n["analysed_is_sample"], "every eligible paper is analysed"
        assert set(ctx.shortlist_ids()) == set(SY.identify(ctx, n["query"])) & set(ctx.shortlist_ids())

        added = ctx.pg.execute("SELECT DISTINCT added_by FROM run_papers WHERE run_id=%s", (rid,)).fetchall()
        assert [r["added_by"] for r in added] == ["systematic"], "no paper was chosen by an agent"
        # every decision is on the record, and exclusions say why
        log = screening_rows(ctx)
        assert all(r["reason"] for r in log if r["stage"] == "excluded")
        assert {r["stage"] for r in log} >= {"included"}
        pc = prisma_counts(ctx)
        assert pc["included"] == n["analysed"] and pc["identified"] >= n["identified"]

        # the discovery agent cannot add to a systematic set
        before = set(ctx.shortlist_ids())
        res = run_specialist(ctx, "discovery", "add more papers")
        assert "systematic search" in res["note"] and set(ctx.shortlist_ids()) == before

        md = ctx.pg.execute("SELECT report_md FROM runs WHERE run_id=%s", (rid,)).fetchone()["report_md"]
        assert "## How the papers were found (computed)" in md and n["query"] in md
        assert "not of the literature" in md
        assert SY.is_systematic(ctx)
    finally:
        ctx.close()

    # past the cap, a seeded random sample is analysed, and the rest are logged as eligible but not sampled
    ctx2 = RunContext.create("malaria forecasting sample", llm_factory=lambda strong=False, step=None: FakeLLM())
    try:
        protocol = {"search_query": "malaria AND (forecast OR forecasting OR prediction OR incidence)",
                    "inclusion": ["malaria studies"], "exclusion": ["not malaria"], "fields": []}
        s2 = SY.systematic_discovery(ctx2, protocol, max_analysed=3)
        assert s2["analysed"] == 3 < s2["eligible"] and s2["analysed_is_sample"]
        assert len(ctx2.shortlist_ids()) == 3
        unsampled = [r for r in screening_rows(ctx2) if r["stage"] == "over_limit"]
        assert len(unsampled) == s2["eligible"] - 3 and all(f"seed {s2['seed']}" in r["reason"] for r in unsampled)
        # the same run draws the same sample: the seed comes from the run, not from the clock
        assert SY._seed(ctx2) == s2["seed"]
        # screening past its own cap samples the identified papers in the same way
        ctx3 = RunContext.create("malaria screening sample", llm_factory=lambda strong=False, step=None: FakeLLM())
        s3 = SY.systematic_discovery(ctx3, protocol, max_screened=4)
        assert s3["screened"] == 4 and s3["screened_is_sample"] and s3["identified"] > 4
        ctx3.close()
        # without a specific query the broad topic query is used, and the report says so
        ctx4 = RunContext.create("malaria fallback", llm_factory=lambda strong=False, step=None: FakeLLM())
        s4 = SY.systematic_discovery(ctx4, {"topic_query": "malaria forecast", "fields": []})
        assert s4["query_was_fallback"] and s4["query"] == "malaria forecast"
        assert "broad topic query" in "\n".join(SY.markdown(ctx4))
        ctx4.close()
    finally:
        ctx2.close()

    # a report run can be systematic too
    out = run_research("Which methods forecast malaria?", mode="pipeline",
                       llm_factory=lambda strong=False: FakeLLM(), search="systematic")
    assert "## How the papers were found (computed)" in out["report"]
